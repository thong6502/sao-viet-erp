"""Bảng khấu hao xuất Excel — không có cột hạch toán, có cột Diễn giải.

Cột ghi chú hạch toán đã gỡ (mg 0278). File này vẫn là thứ kế toán đọc rồi tự gõ định khoản sang
phần mềm kế toán bên ngoài, nên bốn cột tiền phải ra đủ, đúng thứ tự màn hình, và cột Diễn giải
(08/09/2026) nói vì sao nguyên giá tháng này khác tháng trước.
"""
from io import BytesIO

from openpyxl import load_workbook

from app.services.tai_san.excel import xuat_bang_ky


def _rows():
    return [{
        "ma": "TS-0001", "ten": "May in Komori 4 mau", "loai": "tscd", "so_luong": 1,
        "bo_phan_ten": "To In", "nguyen_gia": 3_300_000_000, "muc_trich": 27_500_000,
        "luy_ke": 157_016_129, "con_lai": 3_142_983_871, "dien_giai": None,
    }, {
        "ma": "CC-0001", "ten": "Tam cao su", "loai": "ccdc", "so_luong": 12,
        "bo_phan_ten": "To In", "nguyen_gia": 32_400_000, "muc_trich": 1_350_000,
        "luy_ke": 3_750_000, "con_lai": 28_650_000,
        "dien_giai": "Sửa chữa lớn +3.600.000 ngày 15/08: nguyên giá 28.800.000 → 32.400.000, "
                     "mức tháng 1.200.000 → 1.350.000",
    }]


def test_file_excel_co_dung_cot_va_dong_tong():
    wb = load_workbook(BytesIO(xuat_bang_ky(_rows(), nam=2026, thang=9)))
    ws = wb.active
    tieu_de = [c.value for c in ws[3]]
    assert tieu_de[0] == "Mã"
    assert tieu_de[-1] == "Diễn giải"
    assert "hạch toán" not in " ".join(str(x).lower() for x in tieu_de)
    assert ws.cell(row=4, column=1).value == "TS-0001"
    assert ws.cell(row=4, column=8).value == 3_142_983_871
    assert ws.cell(row=5, column=3).value == 12
    assert "28.800.000 → 32.400.000" in ws.cell(row=5, column=9).value
    # dòng cuối là TỔNG mức trích (cột 6 = "Trích tháng này")
    assert ws.cell(row=6, column=1).value == "TỔNG"
    assert ws.cell(row=6, column=6).value == 28_850_000


def test_bang_rong_van_xuat_duoc():
    wb = load_workbook(BytesIO(xuat_bang_ky([], nam=2026, thang=8)))
    ws = wb.active
    assert ws.cell(row=4, column=1).value == "TỔNG"
    assert ws.cell(row=4, column=6).value == 0
