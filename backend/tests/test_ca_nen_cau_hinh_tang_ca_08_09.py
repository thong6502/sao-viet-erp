"""Bước 6–7 bản rà liên thông (08/09/2026): ca nền / cấu hình đổi sau chốt (A3, B8, B6, C16) và
tăng ca (D1, D2, D5, D8, D9/C4, D11).
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO

from app.db import SessionLocal
from app.repositories.attendance_repo import AttendanceRepository
from app.repositories.overtime_repo import OvertimeRepository
from app.services.attendance_service import VN_TZ
from tests.test_chan_ghi_khi_ky_cong_da_chot import _chot_cong
from tests.test_ra_soat_nhan_su_luong_07_09 import (
    NAM, _bam_du_thang, _ca, _dept_id, _gen, _h, _khai_luong, _line, _nv,
)


def _bam(eid: int, y: int, m: int, d: int, hh: int, mm: int, kieu: str) -> None:
    db = SessionLocal()
    try:
        AttendanceRepository(db).create_log(
            employee_id=eid, check_type=kieu, within_range=True,
            checked_at=datetime(y, m, d, hh, mm, tzinfo=VN_TZ).astimezone(timezone.utc))
    finally:
        db.close()


def _shift_json(client, h, sid: int) -> dict:
    return next(s for s in client.get("/api/attendance/shifts", headers=h).json()["items"] if s["id"] == sid)


# --- 6.1 A3 --------------------------------------------------------------------------------------


def test_khong_go_duoc_moc_ca_nen_da_qua(client):
    h = _h(client)
    e = _nv(client, h, "A3", "Kinh doanh")
    sid = _ca(client, h)
    assert client.put(f"/api/employees/{e}/shift", json={"default_shift_id": sid, "effective_from": "2020-01-01"},
                      headers=h).status_code == 200
    mai = date.today().replace(day=1)
    mai = date(mai.year + (mai.month // 12), mai.month % 12 + 1, 1)     # mùng 1 tháng sau
    assert client.put(f"/api/employees/{e}/shift", json={"default_shift_id": sid, "effective_from": mai.isoformat()},
                      headers=h).status_code == 200
    hist = client.get(f"/api/employees/{e}/shift-history", headers=h).json()["items"]
    qua = next(a for a in hist if a["effective_from"] == "2020-01-01")
    toi = next(a for a in hist if a["effective_from"] == mai.isoformat())
    # Chưa chốt kỳ công nào: gỡ mốc cũ = sửa sai bình thường (kịch bản "lỡ tay bỏ gán ca" của chủ 24/08).
    _khai_luong(client, h, e)
    _bam_du_thang(client, h, e, den=9)
    _chot_cong(client, h, nam=NAM, thang=6)
    # Đã chốt công 06/2026 ⇒ mốc 2020 chạm kỳ đã chốt ⇒ không gỡ được; mốc tương lai thì vẫn gỡ được.
    r = client.delete(f"/api/employees/{e}/shift-history/{qua['id']}", headers=h)
    assert r.status_code == 400 and "ĐÃ CHỐT" in r.json()["detail"], r.text
    assert client.delete(f"/api/employees/{e}/shift-history/{toi['id']}", headers=h).status_code == 204


# --- 6.2 B8 + 6.3 B6 -----------------------------------------------------------------------------


def test_doi_ca_nen_sau_chot_cong_co_co_va_chan_chot_luong(client):
    h = _h(client)
    e = _nv(client, h, "B8", "Kinh doanh")
    _khai_luong(client, h, e)
    _bam_du_thang(client, h, e, thang=8, den=31)
    _chot_cong(client, h, nam=NAM, thang=8)
    _gen(client, h, month=8)
    # Đổi ca nền hiệu lực trong tháng đã chốt (API vẫn mở theo chốt cũ) ⇒ cờ + chặn chốt lương.
    r = client.post("/api/attendance/shifts", json={"name": "HC 9-17 (rà)", "start_time": "09:00",
                                                    "end_time": "17:00"}, headers=h)
    assert r.status_code == 201, r.text
    assert client.put(f"/api/employees/{e}/shift", json={"default_shift_id": r.json()["id"],
                      "effective_from": f"{NAM}-08-01"}, headers=h).status_code == 200
    st = client.get(f"/api/attendance/period?year={NAM}&month=8", headers=h).json()
    assert st["doi_ca_nen_sau_chot"] == 1, st
    r = client.post("/api/luong/lock", json={"year": NAM, "month": 8}, headers=h)
    assert r.status_code == 400 and "đổi ca" in r.json()["detail"], r.text


def test_muc_com_ca_dong_bang_luc_chot_cong(client):
    h = _h(client)
    e = _nv(client, h, "B6", "Kinh doanh")
    _khai_luong(client, h, e)
    sid = _ca(client, h)
    ca = _shift_json(client, h, sid)
    body = {"name": ca["name"], "start_time": "08:00", "end_time": "16:00", "meal_allowance": 25_000,
            "shift_allowance": 0}
    assert client.put(f"/api/attendance/shifts/{sid}", json=body, headers=h).status_code == 200
    _bam_du_thang(client, h, e, thang=8, den=31)
    _chot_cong(client, h, nam=NAM, thang=8)
    truoc = _line(_gen(client, h, month=8), e)["meal_allowance_pay"]
    assert truoc == 26 * 25_000
    # Sửa mức cơm ca SAU chốt công: kỳ 8 (đã chốt công) giữ mức cũ khi Tính lại; kỳ 9 (chưa chốt) theo mức mới.
    assert client.put(f"/api/attendance/shifts/{sid}", json={**body, "meal_allowance": 50_000}, headers=h).status_code == 200
    assert _line(_gen(client, h, month=8), e)["meal_allowance_pay"] == 26 * 25_000
    _bam_du_thang(client, h, e, thang=9, den=5)
    assert _line(_gen(client, h, month=9), e)["meal_allowance_pay"] == 5 * 50_000


# --- 6.4 C16 --------------------------------------------------------------------------------------


def test_doi_co_co_luong_loai_nghi_bao_so_don_bi_anh_huong(client):
    h = _h(client)
    e = _nv(client, h, "C16", "Kinh doanh")
    r = client.post("/api/leaves/types", json={"name": "Nghỉ thử C16", "is_paid": True, "annual_quota": 0}, headers=h)
    assert r.status_code == 201, r.text
    tid = r.json()["id"]
    from app.repositories.leave_repo import LeaveRepository
    hom_nay = date.today()
    db = SessionLocal()
    try:
        LeaveRepository(db).create_request(employee_id=e, leave_type_id=tid, start_date=hom_nay, end_date=hom_nay,
                                           days=1, reason="t", status="approved", created_by=None)
    finally:
        db.close()
    r = client.put(f"/api/leaves/types/{tid}", json={"name": "Nghỉ thử C16", "is_paid": False, "annual_quota": 0,
                                                     "is_active": True}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["canh_bao"] and "1 đơn" in r.json()["canh_bao"], r.json()
    # Không đổi cờ thì không cảnh báo.
    r = client.put(f"/api/leaves/types/{tid}", json={"name": "Nghỉ thử C16 đổi tên", "is_paid": False,
                                                     "annual_quota": 0, "is_active": True}, headers=h)
    assert r.status_code == 200 and r.json()["canh_bao"] is None


# --- 7.1 D1/D2 ------------------------------------------------------------------------------------


def _tat_tang_ca(client, h, phong: str) -> None:
    r = client.put(f"/api/luong/dept-components/{_dept_id(phong)}",
                   json={"items": [{"component_key": "tang_ca", "is_enabled": False}]}, headers=h)
    assert r.status_code == 200, r.text


def test_to_tat_tang_ca_khong_tra_phu_cap_tang_ca_dem_nhung_van_tra_cong_off1x(client):
    h = _h(client)
    _tat_tang_ca(client, h, "Kinh doanh")
    e = _nv(client, h, "D1", "Kinh doanh")
    _khai_luong(client, h, e)
    _bam_du_thang(client, h, e, thang=8, den=31)
    # Ngày 10/08 (T2): phiếu TC 16:30–23:30 duyệt + bấm TC 16:30 → 23:30 (đêm từ 22:00).
    r = client.post("/api/overtime", json={"employee_id": e, "work_date": f"{NAM}-08-10", "from_minute": 990,
                    "to_minute": 1410, "reason": "x"}, headers=h)
    assert r.status_code == 201, r.text
    _bam(e, NAM, 8, 10, 16, 30, "in")
    _bam(e, NAM, 8, 10, 23, 30, "out")
    # Ngày nghỉ 1× (off1x) 15/08 (T7 làm việc) có đi làm.
    r = client.post("/api/calendar/special-days", json={"day": f"{NAM}-08-15", "kind": "off1x",
                    "name": "Nghỉ 1×", "is_paid": False}, headers=h)
    assert r.status_code == 201, r.text
    ln = _line(_gen(client, h, month=8), e)
    # Tắt tăng ca: không hệ số OT, không phụ cấp TC đêm; `ot_pay` chỉ còn đúng 1× lương ngày off1x (D1/D2).
    assert ln["night_premium_pay"] == 0, ln
    # Tháng 8 có 1 ngày off1x (không lương khi nghỉ) ⇒ công chuẩn 25 ⇒ đơn giá ngày 416.000.
    assert ln["ot_pay"] == 416_000, ln
    assert ln["ot_minutes"] > 0                             # phút tăng ca vẫn ghi nhận, chỉ không ra tiền


# --- 7.2 D5 ---------------------------------------------------------------------------------------


def test_ngay_chi_co_tang_ca_van_ra_phut(client):
    h = _h(client)
    e = _nv(client, h, "D5", "Kinh doanh")
    _khai_luong(client, h, e)
    sid = _ca(client, h)
    client.put(f"/api/employees/{e}/shift", json={"default_shift_id": sid, "effective_from": "2020-01-01"}, headers=h)
    # CN 09/08: phiếu 16:30–18:30 duyệt, bấm đúng khung — trước: cong 0, OT 0 "vào trễ quá dung sai".
    r = client.post("/api/overtime", json={"employee_id": e, "work_date": f"{NAM}-08-09", "from_minute": 990,
                    "to_minute": 1110, "reason": "x"}, headers=h)
    assert r.status_code == 201, r.text
    _bam(e, NAM, 8, 9, 16, 30, "in")
    _bam(e, NAM, 8, 9, 18, 30, "out")
    ts = client.get(f"/api/attendance/timesheet?year={NAM}&month=8", headers=h).json()
    row = next(r for r in ts["rows"] if r["employee_id"] == e)
    o = row["days"]["9"]
    assert o["ot_minutes"] == 120 and (o["cong"] or 0) == 0 and not o.get("late"), o


# --- 7.3 D8 ---------------------------------------------------------------------------------------


def test_huy_ho_phieu_da_duyet_dem_vao_chua_xem_cua_nv(client):
    h = _h(client)
    e = _nv(client, h, "D8", "Kinh doanh")
    r = client.post("/api/overtime", json={"employee_id": e, "work_date": f"{NAM}-08-10", "from_minute": 1080,
                    "to_minute": 1200, "reason": "x"}, headers=h)
    assert r.status_code == 201 and r.json()["status"] == "approved", r.text
    rid = r.json()["id"]
    assert client.post(f"/api/overtime/{rid}/cancel", headers=h).status_code == 200
    db = SessionLocal()
    try:
        assert OvertimeRepository(db).count_my_unseen(e) == 1
    finally:
        db.close()


# --- 7.4 D9 / C4 ----------------------------------------------------------------------------------


def test_khong_gui_phieu_don_vao_thang_da_chot_cong(client):
    h = _h(client)
    e = _nv(client, h, "D9", "Kinh doanh")
    _khai_luong(client, h, e)
    _bam_du_thang(client, h, e, thang=8, den=31)
    _chot_cong(client, h, nam=NAM, thang=8)
    r = client.post("/api/overtime", json={"employee_id": e, "work_date": f"{NAM}-08-20", "from_minute": 1080,
                    "to_minute": 1200, "reason": "x"}, headers=h)
    assert r.status_code == 400 and "chốt" in r.json()["detail"].lower(), r.text
    # Đơn nghỉ của chính HCNS (hồ sơ NV Admin, đã vào làm từ 2020) cho tháng đã chốt.
    lt = client.get("/api/leaves/types", headers=h).json()["items"]
    tid = lt[0]["id"] if lt else client.post("/api/leaves/types", json={"name": "Phép năm", "is_paid": True,
                                                                        "annual_quota": 12}, headers=h).json()["id"]
    r = client.post("/api/leaves", json={"leave_type_id": tid, "start_date": f"{NAM}-08-20",
                    "end_date": f"{NAM}-08-20"}, headers=h)
    assert r.status_code == 400 and "chốt" in r.json()["detail"].lower(), r.text


# --- 7.5 D11 --------------------------------------------------------------------------------------


def test_excel_co_cot_gio_tang_ca(client):
    from openpyxl import load_workbook
    h = _h(client)
    e = _nv(client, h, "D11", "Kinh doanh")
    _khai_luong(client, h, e)
    _bam_du_thang(client, h, e)
    _gen(client, h)
    r = client.get(f"/api/luong/export.xlsx?year={NAM}&month=6", headers=h)
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    head = [c.value for c in ws[4]]      # khuôn mới 09/09/2026: tiêu đề ở dòng 4
    assert "Giờ tăng ca" in head and "Ngày ca đêm" in head
