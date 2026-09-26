"""File chuyển khoản tạm ứng / lương đợt 1 theo khuôn lô lương BIZ MBBank (chủ gửi mẫu 25/09/2026).

Mẫu `CK LUONG ỨNG T8.2026.xlsx`: sheet `eMB_BulkPayment`, dòng 1 tiêu đề, dòng 2 đầu cột, dữ liệu từ
dòng 3 — STT · Số tài khoản · Tên thụ hưởng (IN HOA không dấu) · Ngân hàng · Số tiền · Chi tiết
("CTY SAO VIET NHAT CHI LUONG T8.2026 DOT 1"). Chủ: *"tiền mặt thì ghi tiền mặt, còn chuyển khoản
thì ghi số tài khoản, tên ngân hàng của nhân viên"*.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.services.tam_ung_excel import khong_dau_hoa

from .test_accounting_api import _headers, _nv_tam_ung, _tam_ung
from .test_tam_ung_mot_luot import _khai_ngan_hang


def _tai(client, h, **params):
    r = client.get("/api/luong/advances/export.xlsx",
                   params={"year": 2026, "month": 8, **params}, headers=h)
    return r


def _doc(r):
    ws = load_workbook(BytesIO(r.content))["eMB_BulkPayment"]
    return ws, [[c.value for c in row] for row in ws.iter_rows(min_row=3, max_col=6)]


def test_chi_tiet_la_ly_do_trong_thi_cau_mac_dinh():
    from app.services.tam_ung_excel import chi_tiet_thanh_toan
    ly_do = "Tạm ứng lương tháng 09/2026"
    assert chi_tiet_thanh_toan("tam_ung", 2026, 9, ly_do) == "TAM UNG LUONG THANG 09/2026"
    assert chi_tiet_thanh_toan("luong_dot_1", 2026, 8, "  ") == "CTY SAO VIET NHAT CHI LUONG T8.2026 DOT 1"
    assert len(chi_tiet_thanh_toan("tam_ung", 2026, 9, "x" * 300)) == 140


def test_khong_dau_hoa():
    assert khong_dau_hoa("Đoàn Thị  Chúc Phấn ") == "DOAN THI CHUC PHAN"
    assert khong_dau_hoa("Vũ Thị Ngọc Diễm") == "VU THI NGOC DIEM"


def test_file_theo_khuon_mb_chuyen_khoan_truoc_tien_mat_sau(client):
    h = _headers(client)
    ck = _nv_tam_ung(client, h, ten="Đoàn Thị Chúc")
    tm = _nv_tam_ung(client, h, ten="Lưu Kim Dung")
    _khai_ngan_hang(ck, " 0908 872122 ", "MB - Ngan hang TMCP Quan Doi")
    a_tm = _tam_ung(client, h, tm, amount=2_000_000)    # hồ sơ không có TK; lý do "việc nhà"
    a_ck = _tam_ung(client, h, ck, amount=2_500_000, kind="luong_dot_1")
    _tam_ung(client, h, ck, amount=999_000, duyet=False)                   # chờ duyệt: KHÔNG vào file

    r = _tai(client, h)
    assert r.status_code == 200, r.text
    assert "ck-luong-ung-2026-08.xlsx" in r.headers["content-disposition"]
    ws, dong = _doc(r)
    assert ws["B1"].value.startswith("DANH SÁCH GIAO DỊCH")
    assert ws["B2"].value.startswith("Số tài khoản") and ws["F2"].value.startswith("Chi tiết")
    assert len(dong) == 2
    # Chuyển khoản lên đầu: số TK là CHỮ (giữ số 0 đầu), bỏ hết dấu cách; tên IN HOA không dấu.
    assert dong[0] == [1, "0908872122", "DOAN THI CHUC", "MB - Ngan hang TMCP Quan Doi", 2_500_000,
                       "VIEC NHA"]
    assert ws["B3"].number_format == "@"
    # Tiền mặt xuống cuối, ghi chữ TIỀN MẶT vào chỗ tài khoản / ngân hàng.
    # Chi tiết thanh toán = LÝ DO của phiếu, in HOA không dấu (luật ký tự của ngân hàng).
    assert dong[1] == [2, "TIỀN MẶT", "LUU KIM DUNG", "TIỀN MẶT", 2_000_000, "VIEC NHA"]

    # Chỉ những phiếu đang tick.
    ws, dong = _doc(_tai(client, h, ids=str(a_tm)))
    assert [d[2] for d in dong] == ["LUU KIM DUNG"]

    # Đã lập phiếu chi TIỀN MẶT cho người có tài khoản ⇒ file theo hình thức THẬT của phiếu chi.
    pc = client.post("/api/accounting/payment-vouchers/from-advances", json={
        "salary_advance_ids": [a_ck], "voucher_type": "cash", "voucher_date": "2026-08-20",
    }, headers=h)
    assert pc.status_code == 201, pc.text
    ws, dong = _doc(_tai(client, h, ids=str(a_ck)))
    assert dong[0][1] == "TIỀN MẶT" and dong[0][4] == 2_500_000


def test_khong_co_phieu_nao_thi_404(client):
    h = _headers(client)
    r = _tai(client, h)
    assert r.status_code == 404 and "Chưa có phiếu" in r.json()["detail"]
