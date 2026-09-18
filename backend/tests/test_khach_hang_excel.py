"""Nhập Excel — Khách hàng (11/09/2026). Thay hẳn đường nhập CSV cũ.

Test đi qua HTTP vì đây là thứ người dùng thật đâm vào: tải mẫu → điền → xem trước → ghi, và cửa
quyền `set_credit_terms` chỉ có nghĩa ở tầng router.

File dùng trong test dựng từ CHÍNH file mẫu hệ xuất ra (`/mau-excel`) rồi điền vào — nếu mẫu đổi
cột mà bộ đọc quên theo thì test đỏ ngay, đúng cái "file mẫu luôn khớp hệ thống" mà chủ yêu cầu.

Bản 2 (17/09/2026, cuối file): xuất → sửa → nhập lại. File dựng từ CHÍNH `/xuat-excel` rồi sửa ô.
Các ca quyền gọi thẳng tầng service với đúng cờ router truyền xuống — dựng vai + phạm vi thật qua
HTTP cho từng tổ hợp là tốn mà không thêm gì.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.services.customer_excel import TIEU_DE

ADMIN = {"username": "admin", "password": "admin123"}


def _token(client) -> str:
    return client.post("/api/auth/login", json=ADMIN).json()["access_token"]


def _h(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _mau(client, token) -> bytes:
    r = client.get("/api/customers/mau-excel", headers=_h(token))
    assert r.status_code == 200, r.text
    return r.content


def _dien(mau: bytes, dong: list[list]) -> bytes:
    """Điền các dòng vào ĐÚNG file mẫu hệ vừa xuất ra — giữ nguyên sheet `_meta`."""
    wb = load_workbook(BytesIO(mau))
    ws = wb[TIEU_DE]
    for hang in dong:
        ws.append(hang)
    ra = BytesIO()
    wb.save(ra)
    return ra.getvalue()


def _nhap(client, token, du_lieu: bytes, mode: str):
    return client.post(
        f"/api/customers/import-excel?mode={mode}",
        files={"file": ("kh.xlsx", du_lieu,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_h(token),
    )


def _tong(client, token) -> int:
    return client.get("/api/customers", headers=_h(token)).json()["total"]


# ================= file mẫu =================


def test_mau_excel_RONG_va_du_cot(client):
    """Mẫu rỗng (chốt 11/09/2026) — chỉ dòng tiêu đề, không kèm khách đang có. Và KHÔNG có cột Mã:
    mã là mã hệ tự cấp, người dùng không gõ."""
    wb = load_workbook(BytesIO(_mau(client, _token(client))))
    ws = wb[TIEU_DE]
    tieu_de = [o.value for o in ws[1]]

    assert ws.max_row == 1                       # rỗng: đúng một dòng tiêu đề
    assert "Mã" not in tieu_de and "Mã KH" not in tieu_de
    assert tieu_de[0] == "Tên khách hàng"
    # Admin có `set_credit_terms` ⇒ mẫu kèm cả khối tài chính.
    assert "Hạn mức công nợ (đ)" in tieu_de and "Markup tối đa (%)" in tieu_de
    assert "_meta" in wb.sheetnames               # dấu nhận diện màn


def test_file_cua_man_khac_bi_chan(client):
    """Sheet `_meta` là cửa chặn nhập nhầm — dùng lại đúng cơ chế của mọi màn danh mục."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "_meta"
    wb.active.append(["khoa", "gia_tri"])
    wb.active.append(["loai", "may_thiet_bi"])
    ra = BytesIO()
    wb.save(ra)

    r = _nhap(client, _token(client), ra.getvalue(), "preview")
    assert r.status_code == 422
    assert "không phải Khách hàng" in r.json()["detail"]


def test_file_khong_phai_xlsx_bi_chan(client):
    r = _nhap(client, _token(client), b"kh,mst\nabc,123\n", "preview")
    assert r.status_code == 422


# ================= luồng chính =================


def test_them_nhieu_khach_ma_TU_SINH_noi_tiep(client):
    token = _token(client)
    truoc = _tong(client, token)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Excel Một", "Công ty", "0101234567", "0911", "e1@x.vn", "HN", "Anh A"],
        ["Cty Excel Hai", "Cá nhân", None, "0912", None, "HCM", None],
    ])

    r = _nhap(client, token, du_lieu, "commit")
    assert r.status_code == 200, r.text
    kq = r.json()
    assert kq["hop_le"] and kq["da_ghi"] and kq["tong_dong"] == 2 and kq["tao_moi"] == 2
    assert _tong(client, token) == truoc + 2

    ds = client.get("/api/customers?size=200", headers=_h(token)).json()["items"]
    hai = next(c for c in ds if c["name"] == "Cty Excel Hai")
    assert hai["customer_kind"] == "ca_nhan"      # nhãn tiếng Việt map đúng
    assert hai["code"].startswith("KH")           # mã do hệ cấp


def test_XEM_TRUOC_khong_ghi_gi_xuong_db(client):
    """`preview` chạy y hệt `commit` rồi rollback — con số là con số THẬT, nhưng không còn dấu vết."""
    token = _token(client)
    truoc = _tong(client, token)
    du_lieu = _dien(_mau(client, token), [["Cty Chỉ Xem Trước", "Công ty", None, None, None, None, None]])

    kq = _nhap(client, token, du_lieu, "preview").json()
    assert kq["hop_le"] and kq["tao_moi"] == 1
    assert kq["da_ghi"] is False
    assert _tong(client, token) == truoc          # KHÔNG một bản ghi nào lọt xuống


def test_dong_TRANG_cuoi_file_khong_tinh_la_loi(client):
    """Ai cũng để thừa vài dòng trắng cuối file — bắt lỗi ở đó là bắt lỗi một việc không có."""
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Có Thật", "Công ty", None, None, None, None, None],
        [None, None, None, None, None, None, None],
        ["", "", "", "", "", "", ""],
    ])
    kq = _nhap(client, token, du_lieu, "preview").json()
    assert kq["hop_le"] and kq["tong_dong"] == 1 and kq["tao_moi"] == 1


# ================= lỗi: cả file cùng sống hoặc cùng chết =================


def test_thieu_ten_thi_CA_FILE_khong_ghi_gi(client):
    token = _token(client)
    truoc = _tong(client, token)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Hợp Lệ Nhưng Vẫn Không Được Ghi", "Công ty", None, None, None, None, None],
        [None, "Công ty", "0100000000", None, None, None, None],        # thiếu tên
    ])

    kq = _nhap(client, token, du_lieu, "commit").json()
    assert kq["hop_le"] is False and kq["da_ghi"] is False
    assert kq["tao_moi"] == 0
    assert [x["dong"] for x in kq["loi"]] == [3]      # dòng THẬT trên sheet (tính cả tiêu đề)
    assert kq["loi"][0]["cot"] == "Tên khách hàng"
    assert _tong(client, token) == truoc              # dòng hợp lệ cũng KHÔNG được ghi


def test_sale_sai_ten_bao_dung_dong(client):
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Sale Lạ", "Công ty", None, None, None, None, None, "Nguyễn Không Có Thật"],
    ])
    kq = _nhap(client, token, du_lieu, "preview").json()
    assert kq["hop_le"] is False
    assert kq["loi"][0]["dong"] == 2 and kq["loi"][0]["cot"] == "Sale phụ trách"
    assert "Nguyễn Không Có Thật" in kq["loi"][0]["ly_do"]


def test_sale_NGOAI_KHOI_KINH_DOANH_bi_chan_va_noi_dung_ly_do(client):
    """Ô "NV phụ trách" trên màn chỉ cho chọn người thuộc khối Kinh doanh — để không lỡ tay gán
    khách cho Thủ kho. Đường Excel phải chặt y hệt, nếu không nó thành cửa sau.

    Và phải nói ĐÚNG lý do: "không thuộc khối Kinh doanh", chứ không phải "không tìm thấy người
    nào tên đó" — hai câu dẫn người dùng đi hai hướng khác hẳn nhau.
    """
    from app.db import SessionLocal
    from app.models.department import Department
    from app.models.user import User
    from app.repositories.user_repo import UserRepository
    from app.security import hash_password

    db = SessionLocal()
    try:
        kho = db.query(Department).filter(Department.name == "Kho").first()
        if kho is None:
            kho = Department(name="Kho")
            db.add(kho)
            db.commit()
            db.refresh(kho)
        u = db.query(User).filter(User.username == "thukho_excel").first()
        if u is None:
            u = UserRepository(db).create(username="thukho_excel", name="Thủ Kho Excel",
                                          password_hash=hash_password("x"))
            UserRepository(db).set_assignment(u, department_id=kho.id, role_id=None,
                                              is_active=True)
        ten_thu_kho = u.name
    finally:
        db.close()

    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Gán Nhầm Thủ Kho", "Công ty", None, None, None, None, None, ten_thu_kho],
    ])
    kq = _nhap(client, token, du_lieu, "preview").json()

    assert kq["hop_le"] is False
    assert kq["loi"][0]["cot"] == "Sale phụ trách"
    assert "không thuộc khối Kinh doanh" in kq["loi"][0]["ly_do"]


def test_o_chon_sale_trong_mau_CHI_liet_ke_nguoi_du_tu_cach(client):
    """Sheet ẩn `_sale` là nguồn ô chọn xổ xuống — nó phải khớp đúng danh sách màn hình cho chọn."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.db import SessionLocal
    from app.models.user import User
    from app.services.customer_service import nguoi_du_tu_cach_nhan_khach

    wb = load_workbook(BytesIO(_mau(client, _token(client))))
    trong_mau = set()
    if "_sale" in wb.sheetnames:
        trong_mau = {r[0] for r in wb["_sale"].iter_rows(values_only=True) if r and r[0]}

    db = SessionLocal()
    try:
        du = nguoi_du_tu_cach_nhan_khach(db)
        mong_doi = {
            (u.name or u.username)
            for u in db.query(User).all() if u.id in du
        }
    finally:
        db.close()
    assert trong_mau == mong_doi


def test_loai_khach_la_bi_chan(client):
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [["Cty Loại Lạ", "Doanh nghiệp", None, None, None, None, None]])
    kq = _nhap(client, token, du_lieu, "preview").json()
    assert kq["hop_le"] is False and kq["loi"][0]["cot"] == "Loại khách"


def test_rao_min_lon_hon_max_bi_chan(client):
    """Luật nằm ở service (`_validate_bounds`) — bộ nhập không chép lại luật, chỉ dịch lỗi ra dòng."""
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Rào Ngược", "Công ty", None, None, None, None, None, None, None, None, 50, 10],
    ])
    kq = _nhap(client, token, du_lieu, "commit").json()
    assert kq["hop_le"] is False and kq["da_ghi"] is False


# ================= trùng = cảnh báo, KHÔNG chặn =================


def test_trung_MST_chi_CANH_BAO_va_VAN_GHI(client):
    """Model ghi rõ MST không unique, trùng là cảnh báo mềm (§34, §41) — form nhập tay cũng chỉ
    cảnh báo. Chặn ở Excel là chặt hơn cả form, tức là sai."""
    token = _token(client)
    client.post("/api/customers", json={"name": "Cty Gốc MST", "tax_code": "0107654321"},
                headers=_h(token))
    truoc = _tong(client, token)

    du_lieu = _dien(_mau(client, token), [
        ["Cty Trùng MST", "Công ty", "0107654321", None, None, None, None],
    ])
    kq = _nhap(client, token, du_lieu, "commit").json()

    assert kq["hop_le"] and kq["da_ghi"] and kq["tao_moi"] == 1
    assert _tong(client, token) == truoc + 1
    assert len(kq["canh_bao"]) == 1
    assert kq["canh_bao"][0]["dong"] == 2 and "MST" in kq["canh_bao"][0]["ly_do"]


def test_trung_NGAY_TRONG_MOT_FILE_cung_canh_bao(client):
    """Hai dòng cùng MST trong CÙNG một file: dòng sau phải thấy dòng trước — nhờ `flush()` trong
    cùng giao dịch, không phải đợi commit."""
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Nội Bộ A", "Công ty", "0109999999", None, None, None, None],
        ["Cty Nội Bộ B", "Công ty", "0109999999", None, None, None, None],
    ])
    kq = _nhap(client, token, du_lieu, "commit").json()
    assert kq["hop_le"] and kq["tao_moi"] == 2
    assert [c["dong"] for c in kq["canh_bao"]] == [3]     # dòng 2 trùng với dòng 1


# ================= khối tài chính theo quyền =================


def test_co_quyen_thi_GHI_duoc_chinh_sach_tai_chinh(client):
    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Có Hạn Mức", "Công ty", None, None, None, None, None, None,
         5000000, 30, 0, 10, 5, 20],
    ])
    assert _nhap(client, token, du_lieu, "commit").json()["da_ghi"] is True

    ds = client.get("/api/customers?size=200", headers=_h(token)).json()["items"]
    kh = next(c for c in ds if c["name"] == "Cty Có Hạn Mức")
    chi_tiet = client.get(f"/api/customers/{kh['id']}", headers=_h(token)).json()["customer"]
    assert chi_tiet["credit_limit"] == 5000000
    assert chi_tiet["payment_term_days"] == 30
    assert chi_tiet["markup_max_pct"] == 20


def test_khong_quyen_thi_BO_QUA_cot_tai_chinh_va_NOI_RA(client):
    """Nhánh dễ làm ẩu nhất: ghi lén là vượt quyền, chặn cả file là mượn mẫu của đồng nghiệp rồi
    không nhập nổi gì. Bỏ qua ĐÚNG mấy cột đó, ghi phần còn lại, và nói thành lời.

    Gọi thẳng tầng service với `co_tai_chinh=False` — đúng giá trị router truyền xuống khi người
    dùng thiếu `set_credit_terms`.
    """
    from app.db import SessionLocal
    from app.repositories.audit_repo import AuditLogRepository
    from app.repositories.customer_repo import CustomerRepository
    from app.repositories.user_repo import UserRepository
    from app.services import customer_excel
    from app.services.customer_service import CustomerService

    token = _token(client)
    du_lieu = _dien(_mau(client, token), [
        ["Cty Không Được Đặt Hạn Mức", "Công ty", None, None, None, None, None, None,
         9000000, 45, None, None, None, None],
    ])

    db = SessionLocal()
    try:
        admin = UserRepository(db).get_by_username("admin")
        svc = CustomerService(CustomerRepository(db), AuditLogRepository(db))
        kq = customer_excel.nhap(db, svc, du_lieu, actor=admin, scope="all",
                                 co_tai_chinh=False, co_quyen_tao=True, co_quyen_sua=True,
                                 co_quyen_dieu_chuyen=True, ghi=True)
        assert kq.hop_le and kq.da_ghi and kq.tao_moi == 1
        assert kq.bo_qua_tai_chinh is True          # có nói ra, không im lặng
        kh = CustomerRepository(db).find_by_name_exact("Cty Không Được Đặt Hạn Mức") \
            if hasattr(CustomerRepository(db), "find_by_name_exact") else None
        if kh is None:
            from app.models.customer import Customer
            from sqlalchemy import select
            kh = db.execute(
                select(Customer).where(Customer.name == "Cty Không Được Đặt Hạn Mức")
            ).scalars().first()
        assert kh is not None
        assert kh.credit_limit == 0                 # KHÔNG ghi lén
        assert kh.payment_term_days is None
    finally:
        db.close()


# ================= BẢN 2 (17/09/2026): xuất → sửa → nhập lại =================


def _xuat(client, token) -> bytes:
    r = client.get("/api/customers/xuat-excel", headers=_h(token))
    assert r.status_code == 200, r.text
    return r.content


def _sua(du_lieu: bytes, fn) -> bytes:
    """Mở file xuất, gọi `fn(wb, ws, cot, dong)` — `cot`: nhãn → số cột, `dong`: Mã KH → số dòng."""
    wb = load_workbook(BytesIO(du_lieu))
    ws = wb[TIEU_DE]
    cot = {o.value: o.column for o in ws[1]}
    dong = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    fn(wb, ws, cot, dong)
    ra = BytesIO()
    wb.save(ra)
    return ra.getvalue()


def _tao(client, token, **kw) -> dict:
    r = client.post("/api/customers", json=kw, headers=_h(token))
    assert r.status_code == 201, r.text
    return r.json()["customer"]


def _chi_tiet(client, token, kh_id: int) -> dict:
    return client.get(f"/api/customers/{kh_id}", headers=_h(token)).json()["customer"]


def _so_nhat_ky() -> int:
    from app.db import SessionLocal
    from app.models.audit import AuditLog

    db = SessionLocal()
    try:
        return db.query(AuditLog).count()
    finally:
        db.close()


def _nhap_svc(du_lieu: bytes, *, scope: str = "all", ghi: bool = True, **quyen):
    """Gọi thẳng service với đúng các cờ router truyền xuống. Mặc định đủ quyền; ca nào thiếu quyền
    thì truyền cờ đó = False."""
    from app.db import SessionLocal
    from app.repositories.audit_repo import AuditLogRepository
    from app.repositories.customer_repo import CustomerRepository
    from app.repositories.user_repo import UserRepository
    from app.services import customer_excel
    from app.services.customer_service import CustomerService

    co = {"co_tai_chinh": True, "co_quyen_tao": True, "co_quyen_sua": True,
          "co_quyen_dieu_chuyen": True, **quyen}
    db = SessionLocal()
    try:
        admin = UserRepository(db).get_by_username("admin")
        svc = CustomerService(CustomerRepository(db), AuditLogRepository(db))
        return customer_excel.nhap(db, svc, du_lieu, actor=admin, scope=scope, ghi=ghi, **co)
    finally:
        db.close()


def _nguoi_sale(username: str, ten: str) -> int:
    """NV Sales phòng Kinh doanh ⇒ đọc được Khách hàng ⇒ đủ tư cách nhận khách (DB test chưa khai
    khối Kinh doanh nên luật lùi về theo quyền)."""
    from app.db import SessionLocal
    from app.repositories.rbac_repo import DepartmentRepository, RoleRepository
    from app.repositories.user_repo import UserRepository
    from app.security import hash_password

    db = SessionLocal()
    try:
        repo = UserRepository(db)
        kd = DepartmentRepository(db).get_by_name("Kinh doanh")
        vai = RoleRepository(db).get_by_name_and_department("NV Sales", kd.id)
        u = repo.create(username=username, name=ten, password_hash=hash_password("x"))
        repo.set_assignment(u, department_id=kd.id, role_id=vai.id, is_active=True)
        return u.id
    finally:
        db.close()


def test_file_xuat_CO_meta_va_dinh_dang_chu_cho_MST_dien_thoai(client):
    token = _token(client)
    _tao(client, token, name="Cty Định Dạng", tax_code="0101234567", phone="0901234567")
    wb = load_workbook(BytesIO(_xuat(client, token)))

    assert wb.sheetnames[0] == TIEU_DE                  # sheet mở ra đầu tiên là dữ liệu
    assert "_meta" in wb.sheetnames and wb["_meta"].sheet_state == "hidden"
    ws = wb[TIEU_DE]
    cot = {o.value: o.column for o in ws[1]}
    r = next(i for i in range(2, ws.max_row + 1) if ws.cell(row=i, column=2).value == "Cty Định Dạng")
    assert ws.cell(row=r, column=cot["Điện thoại"]).number_format == "@"
    assert ws.cell(row=r, column=cot["MST"]).number_format == "@"


def test_xuat_roi_nhap_lai_Y_NGUYEN_khong_doi_gi_va_KHONG_de_nhat_ky(client):
    """Ca quan trọng nhất của bản 2: không sửa gì thì không đụng gì."""
    token = _token(client)
    kh = _tao(client, token, name="Cty Nguyên Vẹn", tax_code="0101111111", phone="0901")
    client.put(f"/api/customers/{kh['id']}/financial", headers=_h(token), json={
        "credit_limit": 7000000, "payment_term_days": 30,
        "discount_min_pct": 1.5, "discount_max_pct": 8, "markup_min_pct": 5, "markup_max_pct": 20,
    })
    _tao(client, token, name="Cty Trơn")
    tong = _tong(client, token)
    nhat_ky = _so_nhat_ky()

    kq = _nhap(client, token, _xuat(client, token), "commit").json()
    assert kq["hop_le"], kq["loi"]
    assert (kq["tao_moi"], kq["cap_nhat"], kq["khong_doi"]) == (0, 0, tong)
    assert kq["thay_doi"] == [] and kq["canh_bao"] == []
    assert _so_nhat_ky() == nhat_ky
    assert _tong(client, token) == tong


def test_file_xuat_DOI_CU_khong_meta_van_nhan(client):
    """File tải trước 17/09/2026 không có `_meta` — chính cái file chủ đã thử và bị chặn."""
    token = _token(client)
    _tao(client, token, name="Cty Đời Cũ")

    def bo_meta(wb, ws, cot, dong):
        del wb["_meta"]

    r = _nhap(client, token, _sua(_xuat(client, token), bo_meta), "preview")
    assert r.status_code == 200, r.text
    assert r.json()["hop_le"] and r.json()["khong_doi"] >= 1


def test_sua_DIEN_THOAI_mot_khach_va_xem_truoc_hien_cu_moi(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Đổi Số", phone="0901000000")
    _tao(client, token, name="Cty Giữ Số", phone="0902000000")

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Điện thoại"], value="0911222333")

    du_lieu = _sua(_xuat(client, token), sua)
    xem = _nhap(client, token, du_lieu, "preview").json()
    assert xem["hop_le"] and xem["cap_nhat"] == 1 and xem["tao_moi"] == 0
    assert xem["thay_doi"] == [{
        "dong": xem["thay_doi"][0]["dong"], "ma": kh["code"], "ten": "Cty Đổi Số",
        "cot": "Điện thoại", "cu": "0901000000", "moi": "0911222333",
    }]
    assert _chi_tiet(client, token, kh["id"])["phone"] == "0901000000"     # xem trước không ghi

    kq = _nhap(client, token, du_lieu, "commit").json()
    assert kq["da_ghi"] and kq["cap_nhat"] == 1
    assert _chi_tiet(client, token, kh["id"])["phone"] == "0911222333"


def test_vua_SUA_vua_THEM_trong_cung_file(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Được Sửa")
    tong = _tong(client, token)

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Địa chỉ"], value="Số 1 Hà Nội")
        moi = ws.max_row + 1
        ws.cell(row=moi, column=cot["Tên khách hàng"], value="Cty Thêm Cuối File")
        ws.cell(row=moi, column=cot["Loại khách"], value="Cá nhân")

    kq = _nhap(client, token, _sua(_xuat(client, token), sua), "commit").json()
    assert kq["hop_le"], kq["loi"]
    assert kq["tao_moi"] == 1 and kq["cap_nhat"] == 1
    assert _tong(client, token) == tong + 1
    assert _chi_tiet(client, token, kh["id"])["address"] == "Số 1 Hà Nội"


def test_MA_KHONG_CO_bao_loi_va_khong_ghi_dong_nao(client):
    token = _token(client)
    a = _tao(client, token, name="Cty Sửa Hợp Lệ", phone="0901")
    b = _tao(client, token, name="Cty Mã Bị Gõ Sai")

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[a["code"]], column=cot["Điện thoại"], value="0999")
        ws.cell(row=dong[b["code"]], column=1, value="KH999")

    kq = _nhap(client, token, _sua(_xuat(client, token), sua), "commit").json()
    assert kq["hop_le"] is False and kq["da_ghi"] is False
    assert len(kq["loi"]) == 1 and kq["loi"][0]["cot"] == "Mã KH"
    assert "KH999" in kq["loi"][0]["ly_do"]
    assert _chi_tiet(client, token, a["id"])["phone"] == "0901"


def test_TRUNG_MA_trong_file_loi_ca_hai_dong(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Bị Chép Dòng")

    def chep(wb, ws, cot, dong):
        goc = dong[kh["code"]]
        moi = ws.max_row + 1
        for c in range(1, ws.max_column + 1):
            ws.cell(row=moi, column=c, value=ws.cell(row=goc, column=c).value)

    kq = _nhap(client, token, _sua(_xuat(client, token), chep), "preview").json()
    assert kq["hop_le"] is False
    assert len(kq["loi"]) == 2 and all(x["cot"] == "Mã KH" for x in kq["loi"])


def test_O_TRONG_la_xoa_THIEU_COT_la_giu(client):
    """Luật 7 bên danh mục: ô trống ở cột có mặt xoá giá trị; thiếu hẳn cột thì giữ nguyên."""
    token = _token(client)
    kh = _tao(client, token, name="Cty Ô Trống", email="a@x.vn", phone="0901")

    def xoa_email(wb, ws, cot, dong):
        # `ws.cell(..., value=None)` KHÔNG xoá ô (openpyxl bỏ qua None) — phải gán `.value`.
        ws.cell(row=dong[kh["code"]], column=cot["Email"]).value = None

    assert _nhap(client, token, _sua(_xuat(client, token), xoa_email), "commit").json()["da_ghi"]
    assert _chi_tiet(client, token, kh["id"])["email"] is None

    def bo_cot_dien_thoai(wb, ws, cot, dong):
        ws.delete_cols(cot["Điện thoại"])

    kq = _nhap(client, token, _sua(_xuat(client, token), bo_cot_dien_thoai), "commit").json()
    assert kq["hop_le"] and kq["cap_nhat"] == 0
    assert _chi_tiet(client, token, kh["id"])["phone"] == "0901"


def test_XOA_TRANG_o_Sale_la_GO_sale_va_co_canh_bao(client):
    """Chủ chốt 17/09/2026: xoá trắng ô Sale thì cho qua và gỡ hẳn người phụ trách."""
    token = _token(client)
    kh = _tao(client, token, name="Cty Bị Gỡ Sale")
    assert _chi_tiet(client, token, kh["id"])["sale_user_id"] is not None

    def xoa_sale(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Sale phụ trách"]).value = None

    kq = _nhap(client, token, _sua(_xuat(client, token), xoa_sale), "commit").json()
    assert kq["hop_le"] and kq["cap_nhat"] == 1, kq["loi"]
    assert any("Gỡ Sale" in c["ly_do"] for c in kq["canh_bao"])
    assert _chi_tiet(client, token, kh["id"])["sale_user_id"] is None


def test_doi_Sale_sang_nguoi_khac(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Đổi Người")
    uid = _nguoi_sale("sale_moi_excel", "Sale Mới Excel")

    def doi(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Sale phụ trách"], value="Sale Mới Excel")

    kq = _nhap(client, token, _sua(_xuat(client, token), doi), "commit").json()
    assert kq["hop_le"] and kq["cap_nhat"] == 1, kq["loi"]
    assert _chi_tiet(client, token, kh["id"])["sale_user_id"] == uid


def test_Sale_TRUNG_HO_TEN_hoac_NGOAI_KHOI_khong_sua_thi_khong_loi(client):
    """File xuất ghi họ tên. Tra thẳng thì hai ca này chặn một file không hề sửa gì: họ tên trùng
    (bảng tra bỏ khoá tên) và người ngoài khối Kinh doanh đang giữ khách cũ."""
    from app.db import SessionLocal
    from app.repositories.audit_repo import AuditLogRepository
    from app.repositories.customer_repo import CustomerRepository
    from app.repositories.user_repo import UserRepository
    from app.security import hash_password
    from app.services.customer_service import CustomerService

    token = _token(client)
    db = SessionLocal()
    try:
        repo = UserRepository(db)
        giu = repo.create(username="trung_ten_1", name="Trùng Họ Tên", password_hash=hash_password("x"))
        repo.create(username="trung_ten_2", name="Trùng Họ Tên", password_hash=hash_password("x"))
        admin = repo.get_by_username("admin")
        # role_id=None ⇒ không đọc được Khách hàng ⇒ NGOÀI khối, nhưng vẫn đang giữ khách này.
        CustomerService(CustomerRepository(db), AuditLogRepository(db)).create_customer(
            name="Cty Của Người Trùng Tên", tax_code=None, phone=None, email=None, address=None,
            contact_name=None, sale_user_id=giu.id, actor=admin,
        )
    finally:
        db.close()

    kq = _nhap(client, token, _xuat(client, token), "preview").json()
    assert kq["hop_le"], kq["loi"]
    assert kq["cap_nhat"] == 0


def test_QUYEN_theo_tung_dong(client):
    """Thiếu `update`: dòng cũ không đổi vẫn qua, chỉ dòng thật sự sửa bị chặn. Thiếu `create`: dòng
    Mã trống bị chặn."""
    token = _token(client)
    kh = _tao(client, token, name="Cty Không Được Sửa")
    goc = _xuat(client, token)

    kq = _nhap_svc(goc, co_quyen_sua=False, ghi=False)
    assert kq.hop_le and kq.khong_doi >= 1

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Địa chỉ"], value="Chỗ mới")

    kq = _nhap_svc(_sua(goc, sua), co_quyen_sua=False, ghi=False)
    assert not kq.hop_le and len(kq.loi) == 1 and kh["code"] in kq.loi[0].ly_do

    def them(wb, ws, cot, dong):
        ws.cell(row=ws.max_row + 1, column=cot["Tên khách hàng"], value="Cty Không Được Thêm")

    kq = _nhap_svc(_sua(goc, them), co_quyen_tao=False, ghi=False)
    assert not kq.hop_le and len(kq.loi) == 1 and "Thêm" in kq.loi[0].ly_do


def test_doi_hoac_go_Sale_can_quyen_DIEU_CHUYEN_va_pham_vi_khac_own(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Không Được Chuyển")
    _nguoi_sale("sale_chuyen_excel", "Sale Chuyển Excel")

    def doi(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Sale phụ trách"], value="Sale Chuyển Excel")

    def go(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Sale phụ trách"]).value = None

    goc = _xuat(client, token)
    for du_lieu in (_sua(goc, doi), _sua(goc, go)):
        kq = _nhap_svc(du_lieu, co_quyen_dieu_chuyen=False, ghi=False)
        assert not kq.hop_le and kq.loi[0].cot == "Sale phụ trách"
        kq = _nhap_svc(du_lieu, scope="own", ghi=False)
        assert not kq.hop_le and "điều chuyển" in kq.loi[0].ly_do


def test_pham_vi_PHONG_chi_chuyen_Sale_trong_phong_minh(client):
    """Y hệt nút Điều chuyển: ở phạm vi phòng, người nhận phải cùng phòng với người nhập."""
    from app.db import SessionLocal
    from app.models.department import Department
    from app.repositories.rbac_repo import DepartmentRepository, RoleRepository
    from app.repositories.user_repo import UserRepository
    from app.security import hash_password

    token = _token(client)
    kh = _tao(client, token, name="Cty Trong Phòng")
    _nguoi_sale("sale_phong_khac", "Sale Phòng Khác")
    db = SessionLocal()
    try:
        repo = UserRepository(db)
        # Một phòng seed sẵn bất kỳ, KHÁC Kinh doanh (nơi "Sale Phòng Khác" ngồi) — bật cờ khối
        # Kinh doanh để người ngồi đó đủ tư cách nhận khách, còn lại chỉ khác nhau ở PHÒNG.
        phong = db.query(Department).filter(Department.name != "Kinh doanh").first()
        DepartmentRepository(db).set_la_kinh_doanh(phong, True)
        admin = repo.get_by_username("admin")
        repo.set_assignment(admin, department_id=phong.id, role_id=admin.role_id, is_active=True)
        kd = DepartmentRepository(db).get_by_name("Kinh doanh")
        vai = RoleRepository(db).get_by_name_and_department("NV Sales", kd.id)
        u = repo.create(username="sale_cung_phong", name="Sale Cùng Phòng",
                        password_hash=hash_password("x"))
        repo.set_assignment(u, department_id=phong.id, role_id=vai.id, is_active=True)
    finally:
        db.close()

    def doi(ten):
        def _f(wb, ws, cot, dong):
            ws.cell(row=dong[kh["code"]], column=cot["Sale phụ trách"], value=ten)
        return _f

    goc = _xuat(client, token)
    kq = _nhap_svc(_sua(goc, doi("Sale Phòng Khác")), scope="department", ghi=False)
    assert not kq.hop_le and "không thuộc phòng của bạn" in kq.loi[0].ly_do
    kq = _nhap_svc(_sua(goc, doi("Sale Cùng Phòng")), scope="department", ghi=False)
    assert kq.hop_le and kq.cap_nhat == 1, kq.loi


def test_khong_quyen_TAI_CHINH_chi_noi_ra_khi_thuc_su_sua(client):
    """File xuất có cột tài chính cho mọi người xem. Không sửa ô nào mà vẫn báo "đã bỏ qua" thì Sale
    nào nhập lại file xuất cũng ăn một câu vô nghĩa."""
    token = _token(client)
    kh = _tao(client, token, name="Cty Hạn Mức Cũ")
    goc = _xuat(client, token)

    kq = _nhap_svc(goc, co_tai_chinh=False, ghi=False)
    assert kq.hop_le and kq.bo_qua_tai_chinh is False

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Hạn mức công nợ (đ)"], value=99000000)

    kq = _nhap_svc(_sua(goc, sua), co_tai_chinh=False, ghi=True)
    assert kq.hop_le and kq.bo_qua_tai_chinh is True and kq.cap_nhat == 0
    assert _chi_tiet(client, token, kh["id"])["credit_limit"] == 0


def test_co_quyen_sua_TAI_CHINH_qua_file_xuat(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Nâng Hạn Mức")

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Hạn mức công nợ (đ)"], value=12000000)
        ws.cell(row=dong[kh["code"]], column=cot["Markup tối đa (%)"], value=25)

    kq = _nhap(client, token, _sua(_xuat(client, token), sua), "commit").json()
    assert kq["hop_le"] and kq["cap_nhat"] == 1, kq["loi"]
    assert {x["cot"] for x in kq["thay_doi"]} == {"Hạn mức công nợ (đ)", "Markup tối đa (%)"}
    ct = _chi_tiet(client, token, kh["id"])
    assert ct["credit_limit"] == 12000000 and ct["markup_max_pct"] == 25


def test_doi_CA_TEN_LAN_MST_canh_bao_tro_nham_nhung_van_ghi(client):
    token = _token(client)
    kh = _tao(client, token, name="Cty Tên Gốc", tax_code="0102222222")

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[kh["code"]], column=cot["Tên khách hàng"], value="Cty Tên Khác Hẳn")
        ws.cell(row=dong[kh["code"]], column=cot["MST"], value="0103333333")

    kq = _nhap(client, token, _sua(_xuat(client, token), sua), "commit").json()
    assert kq["da_ghi"] and kq["cap_nhat"] == 1
    assert any("trỏ nhầm" in c["ly_do"] for c in kq["canh_bao"])


def test_trung_MST_CU_khong_canh_bao_khi_dong_khong_sua(client):
    token = _token(client)
    _tao(client, token, name="Cty Trùng Cũ A", tax_code="0104444444")
    _tao(client, token, name="Cty Trùng Cũ B", tax_code="0104444444")

    kq = _nhap(client, token, _xuat(client, token), "preview").json()
    assert kq["hop_le"] and kq["canh_bao"] == []


def test_loi_o_dong_CUOI_luc_ghi_rollback_ca_dong_sua_phia_tren(client):
    """MST sai dạng chỉ lộ ra ở service — lúc đó dòng sửa phía trên đã flush. Phải rollback hết."""
    token = _token(client)
    a = _tao(client, token, name="Cty Sửa Trước", phone="0901")
    b = _tao(client, token, name="Cty MST Hỏng")

    def sua(wb, ws, cot, dong):
        ws.cell(row=dong[a["code"]], column=cot["Điện thoại"], value="0988")
        ws.cell(row=dong[b["code"]], column=cot["MST"], value="123")

    kq = _nhap(client, token, _sua(_xuat(client, token), sua), "commit").json()
    assert kq["hop_le"] is False and kq["da_ghi"] is False and kq["cap_nhat"] == 0
    assert _chi_tiet(client, token, a["id"])["phone"] == "0901"
