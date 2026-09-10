"""Xuất / Nhập Excel cho màn Hồ sơ nhân sự.

Mục đích của tính năng: NẠP DỮ LIỆU BAN ĐẦU (khai một lượt vài trăm hồ sơ trên Excel rồi đổ vào),
và sửa hàng loạt lúc đang dọn dữ liệu. Vì thế:

* XUẤT phải ra ĐỦ ô của hồ sơ — file xuất ra chính là file nạp ngược lại được.
* NHẬP là upsert theo `Mã`: có mã thì sửa người đó, để trống mã thì tạo mới (máy tự cấp mã).
* Cả file là MỘT giao dịch: còn một dòng sai thì không ghi gì cả.
* `mode=preview` chạy y hệt `commit` rồi rollback ⇒ con số xem trước là con số THẬT.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.db import SessionLocal
from app.repositories.employee_repo import EmployeeRepository
from app.repositories.rbac_repo import DepartmentRepository, RoleRepository
from app.repositories.user_repo import UserRepository
from app.security import create_access_token, hash_password
from app.services.employee_excel import MAU_TIEU_DE, MAU_TIEU_DE_NHAY_CAM

ADMIN = {"username": "admin", "password": "admin123"}
SHEET = "Nhan su"


def _token(client) -> str:
    return client.post("/api/auth/login", json=ADMIN).json()["access_token"]


def _h(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _dept_id(name: str) -> int:
    db = SessionLocal()
    try:
        return DepartmentRepository(db).get_by_name(name).id
    finally:
        db.close()


def _ten_bac() -> str:
    db = SessionLocal()
    try:
        return EmployeeRepository(db).list_job_grades()[0].name
    finally:
        db.close()


def _sales_token() -> str:
    """Tài khoản KHÔNG có quyền nhan_su — dùng để kiểm hàng rào quyền."""
    db = SessionLocal()
    try:
        users = UserRepository(db)
        existing = users.get_by_username("sales-excel")
        if existing is not None:
            return create_access_token(str(existing.id))
        kd = DepartmentRepository(db).get_by_name("Kinh doanh")
        role = RoleRepository(db).get_by_name_and_department("NV Sales", kd.id)
        u = users.create(username="sales-excel", name="S", password_hash=hash_password("x"))
        users.set_assignment(u, department_id=kd.id, role_id=role.id, is_active=True)
        return create_access_token(str(u.id))
    finally:
        db.close()


def _tao_nv(client, token, **over):
    body = {
        "full_name": over.pop("full_name", "Trần Văn A"),
        "department_id": over.pop("department_id", _dept_id("Hành chính nhân sự")),
        "hire_date": over.pop("hire_date", "2024-01-15"),
        "probation_end_date": over.pop("probation_end_date", "2025-12-31"),
    }
    body.update(over)
    r = client.post("/api/employees", json=body, headers=_h(token))
    assert r.status_code == 201, r.text
    return r.json()["employee"]


def _tai_xuat(client, token) -> Workbook:
    r = client.get("/api/employees/export.xlsx", headers=_h(token))
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))


def _bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nhap(client, token, wb: Workbook, mode: str = "preview"):
    r = client.post(
        f"/api/employees/import-excel?mode={mode}",
        headers=_h(token),
        files={"file": ("nhan-su.xlsx", _bytes(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    return r


def _cot(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value}


def _dat(ws, dong: int, cot: dict[str, int], nhan: str, gia_tri) -> None:
    ws.cell(row=dong, column=cot[nhan]).value = gia_tri


# --- XUẤT -------------------------------------------------------------------


def test_xuat_ra_du_o_cua_ho_so(client):
    """Trước 10/09/2026 file xuất chỉ có 8 cột — nạp ngược lại thì mất sạch phần còn lại."""
    token = _token(client)
    _tao_nv(client, token, full_name="Nguyễn Thị B", phone="0900000001",
            national_id="0123456789", bank_name="Vietcombank", dependents_count=2)

    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    for nhan in ("Mã", "Họ tên", "Phòng/Tổ", "Chức danh", "Bậc tay nghề", "Trạng thái",
                 "Ngày vào", "Ngày hết thử việc", "Ngày sinh", "Giới tính", "CCCD",
                 "Ngày cấp CCCD", "Nơi cấp CCCD", "Điện thoại", "Email",
                 "Hộ khẩu thường trú", "Chỗ ở hiện tại", "Người liên hệ khẩn",
                 "SĐT liên hệ khẩn", "Số sổ BHXH", "MST cá nhân", "Số người phụ thuộc",
                 "Cách tính thuế TNCN", "Số tài khoản NH", "Ngân hàng", "Nhóm lương",
                 "Thâm niên trước (tháng)"):
        assert nhan in cot, f"thiếu cột {nhan!r}"
    # Bỏ 10/09/2026 — ca nền gán ở Khai ca, hai cột nghỉ việc do luồng trên màn sinh ra,
    # hai cột tài khoản là dữ liệu của RBAC (file không tạo/gán tài khoản được).
    for nhan in ("Ca mặc định", "Ghi chú", "Ngày nghỉ việc (chỉ xem)",
                 "Lý do nghỉ việc (chỉ xem)", "Tài khoản (chỉ xem)",
                 "Vai trò tài khoản (chỉ xem)"):
        assert nhan not in cot, f"còn cột {nhan!r}"

    dong = next(r for r in ws.iter_rows(min_row=2, values_only=False)
                if r[cot["Họ tên"] - 1].value == "Nguyễn Thị B")
    assert dong[cot["Điện thoại"] - 1].value == "0900000001"
    assert dong[cot["Ngân hàng"] - 1].value == "Vietcombank"
    assert dong[cot["Số người phụ thuộc"] - 1].value == 2


def test_xuat_khong_kem_cot_nhay_cam_khi_thieu_quyen_luong(client):
    """Ô lương/BHXH/ngân hàng che khi ĐỌC hồ sơ thì cũng phải che trong file xuất."""
    token = _token(client)
    _tao_nv(client, token, full_name="Ẩn Lương", bank_account="0011002233")

    db = SessionLocal()
    try:
        users = UserRepository(db)
        hcns = DepartmentRepository(db).get_by_name("Hành chính nhân sự")
        role = RoleRepository(db).list_by_department(hcns.id)[0]
        u = users.create(username="hr-no-salary", name="H", password_hash=hash_password("x"))
        users.set_assignment(u, department_id=hcns.id, role_id=role.id, is_active=True)
        rid = role.id
    finally:
        db.close()
    # Vai này có nhan_su nhưng KHÔNG có edit_salary/view_salary.
    r = client.put(
        f"/api/roles/{rid}/permissions",
        json={"permissions": [{"module_key": "nhan_su", "scope": "all", "can_read": True,
                               "can_export": True}]},
        headers=_h(token),
    )
    assert r.status_code in (200, 204), r.text

    wb = load_workbook(BytesIO(
        client.get("/api/employees/export.xlsx",
                   headers=_h(create_access_token(str(u.id)))).content))
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(r for r in ws.iter_rows(min_row=2)
                if r[cot["Họ tên"] - 1].value == "Ẩn Lương")
    assert dong[cot["Số tài khoản NH"] - 1].value in (None, "")


# --- MẪU --------------------------------------------------------------------


def test_mau_nhap_co_tieu_de_va_sheet_huong_dan(client):
    token = _token(client)
    r = client.get("/api/employees/mau-nhap.xlsx", headers=_h(token))
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert SHEET in wb.sheetnames
    assert "Huong dan" in wb.sheetnames
    ws = wb[SHEET]
    assert "Mã" in _cot(ws)
    # Kẻ sẵn mấy chục dòng trống cho ra hình tờ khai, nhưng KHÔNG kèm dữ liệu của ai.
    assert all(o.value in (None, "") for hang in ws.iter_rows(min_row=2) for o in hang)
    # Sheet Hướng dẫn phải liệt kê giá trị hợp lệ (tên phòng/tổ, bậc, trạng thái).
    chu = "\n".join(str(c.value or "") for r_ in wb["Huong dan"].iter_rows() for c in r_)
    assert "Hành chính nhân sự" in chu
    assert "Thử việc" in chu



def test_file_xuat_duoc_ke_bang_va_co_o_chon(client):
    """Người khai ngồi trên file này cả buổi: phải khoá được dòng tiêu đề, lọc được, và cột
    phải-gõ-đúng-tên thì chọn trong danh sách chứ không gõ chay rồi ăn lỗi lúc nhập."""
    token = _token(client)
    _tao_nv(client, token, full_name="Có Kẻ Bảng")
    ws = _tai_xuat(client, token)[SHEET]

    assert ws.freeze_panes == "C2"          # giữ Mã + Họ tên khi kéo ngang 27 cột
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb == MAU_TIEU_DE
    cot = _cot(ws)
    o_nhay = ws.cell(row=1, column=cot["Số sổ BHXH"])
    assert o_nhay.fill.fgColor.rgb == MAU_TIEU_DE_NHAY_CAM, "khối lương/BHXH phải tô khác"

    vung = " ".join(str(o.sqref) for o in ws.data_validations.dataValidation)
    for nhan in ("Trạng thái", "Giới tính", "Cách tính thuế TNCN", "Phòng/Tổ", "Bậc tay nghề"):
        chu = ws.cell(row=1, column=cot[nhan]).column_letter
        assert f"{chu}2:" in vung, f"cột {nhan!r} chưa có ô chọn"

# --- NHẬP: tạo mới ----------------------------------------------------------


def test_nhap_tao_moi_khi_bo_trong_ma(client):
    token = _token(client)
    wb = load_workbook(BytesIO(client.get("/api/employees/mau-nhap.xlsx",
                                          headers=_h(token)).content))
    del wb["Huong dan"]
    ws = wb[SHEET]
    cot = _cot(ws)
    _dat(ws, 2, cot, "Họ tên", "Lê Văn Mới")
    _dat(ws, 2, cot, "Phòng/Tổ", "Hành chính nhân sự")
    _dat(ws, 2, cot, "Trạng thái", "Chính thức")
    _dat(ws, 2, cot, "Ngày vào", "01/03/2024")
    _dat(ws, 2, cot, "Điện thoại", "0911222333")
    _dat(ws, 2, cot, "Giới tính", "Nam")
    _dat(ws, 2, cot, "Số người phụ thuộc", 1)

    xem = _nhap(client, token, wb, "preview")
    assert xem.status_code == 200, xem.text
    kq = xem.json()
    assert kq["hop_le"] is True
    assert (kq["tao_moi"], kq["cap_nhat"], kq["da_ghi"]) == (1, 0, False)
    # XEM TRƯỚC KHÔNG ĐƯỢC GHI GÌ.
    ds = client.get("/api/employees?q=Lê Văn Mới", headers=_h(token)).json()
    assert ds["total"] == 0

    chot = _nhap(client, token, wb, "commit")
    assert chot.status_code == 200, chot.text
    assert chot.json()["da_ghi"] is True

    ds = client.get("/api/employees?q=Lê Văn Mới", headers=_h(token)).json()
    assert ds["total"] == 1
    row = ds["items"][0]
    assert row["code"].startswith("NV")           # máy tự cấp mã
    assert row["department_name"] == "Hành chính nhân sự"
    assert row["status"] == "active"
    chi_tiet = client.get(f"/api/employees/{row['id']}", headers=_h(token)).json()
    assert chi_tiet["phone"] == "0911222333"
    assert chi_tiet["gender"] == "male"
    assert chi_tiet["dependents_count"] == 1
    # Tạo hồ sơ phải để lại mốc "Vào làm" trên Quá trình công tác.
    su_kien = client.get(f"/api/employees/{row['id']}/events", headers=_h(token)).json()
    assert any(e["event_type"] == "hired" for e in su_kien["items"])


# --- NHẬP: cập nhật theo mã --------------------------------------------------


def test_nhap_lai_chinh_file_vua_xuat_thi_khong_doi_gi(client):
    """Vòng tròn xuất → nhập phải KHÉP KÍN: không sửa gì thì không dòng nào bị coi là đổi."""
    token = _token(client)
    _tao_nv(client, token, full_name="Vòng Tròn", phone="0900000009")
    wb = _tai_xuat(client, token)

    kq = _nhap(client, token, wb, "preview").json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["tao_moi"] == 0
    assert kq["cap_nhat"] == 0
    assert kq["khong_doi"] == kq["tong_dong"]


def test_nhap_cap_nhat_theo_ma(client):
    token = _token(client)
    nv = _tao_nv(client, token, full_name="Sửa Tôi", phone="0900000002")
    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(r[0].row for r in ws.iter_rows(min_row=2)
                if r[cot["Mã"] - 1].value == nv["code"])
    _dat(ws, dong, cot, "Điện thoại", "0999888777")
    _dat(ws, dong, cot, "Email", "sua@svn.vn")

    kq = _nhap(client, token, wb, "commit").json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["cap_nhat"] == 1
    ct = client.get(f"/api/employees/{nv['id']}", headers=_h(token)).json()
    assert ct["phone"] == "0999888777"
    assert ct["email"] == "sua@svn.vn"


def test_o_trong_o_cot_khong_cho_rong_thi_ve_mac_dinh(client):
    """Ô trống ở cột NOT NULL phải về MẶC ĐỊNH, không ghi NULL.

    Bắt được lúc thao tác thật trên dev-browser 10/09/2026: dựng file sửa từ 'Tải file mẫu'
    (mọi ô trống trừ ô cần sửa) rồi nhập ⇒ `pit_mode=None` xuống Postgres ⇒ NotNullViolation
    ⇒ 500 trắng, trình duyệt chỉ hiện 'không tới được máy chủ' vì 500 sinh ngoài lớp CORS.
    """
    token = _token(client)
    nv = _tao_nv(client, token, full_name="Ô Trống", pit_mode="khau_tru_10",
                 dependents_count=3, prior_seniority_months=12)
    assert nv["prior_seniority_months"] == 12  # có số để mà reset, không thì test rỗng
    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(r[0].row for r in ws.iter_rows(min_row=2)
                if r[cot["Mã"] - 1].value == nv["code"])
    for nhan in ("Cách tính thuế TNCN", "Số người phụ thuộc", "Thâm niên trước (tháng)"):
        _dat(ws, dong, cot, nhan, None)
    _dat(ws, dong, cot, "Điện thoại", "0912345678")

    r = _nhap(client, token, wb, "commit")
    assert r.status_code == 200, r.text
    kq = r.json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["cap_nhat"] == 1
    ct = client.get(f"/api/employees/{nv['id']}", headers=_h(token)).json()
    assert ct["pit_mode"] == "luy_tien"
    assert ct["dependents_count"] == 0
    assert ct["prior_seniority_months"] == 0
    assert ct["phone"] == "0912345678"


def test_nhap_doi_phong_to_thi_de_lai_moc_qua_trinh_cong_tac(client):
    token = _token(client)
    nv = _tao_nv(client, token, full_name="Chuyển Tổ")
    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(r[0].row for r in ws.iter_rows(min_row=2)
                if r[cot["Mã"] - 1].value == nv["code"])
    _dat(ws, dong, cot, "Phòng/Tổ", "Kinh doanh")

    kq = _nhap(client, token, wb, "commit").json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["cap_nhat"] == 1
    ct = client.get(f"/api/employees/{nv['id']}", headers=_h(token)).json()
    assert ct["department_name"] == "Kinh doanh"
    su_kien = client.get(f"/api/employees/{nv['id']}/events", headers=_h(token)).json()
    assert any(e["event_type"] == "transferred" for e in su_kien["items"])


def test_nhap_khong_doi_duoc_trang_thai_cua_nguoi_da_co(client):
    """Nghỉ việc / đình chỉ còn KHOÁ TÀI KHOẢN + cắt phiên đang sống — Excel không làm thay được."""
    token = _token(client)
    nv = _tao_nv(client, token, full_name="Giữ Trạng Thái")
    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(r[0].row for r in ws.iter_rows(min_row=2)
                if r[cot["Mã"] - 1].value == nv["code"])
    _dat(ws, dong, cot, "Trạng thái", "Đã nghỉ")

    kq = _nhap(client, token, wb, "preview").json()
    assert kq["hop_le"] is False
    assert any(l["cot"] == "Trạng thái" for l in kq["loi"])
    ct = client.get(f"/api/employees/{nv['id']}", headers=_h(token)).json()
    assert ct["status"] != "resigned"


# --- NHẬP: lỗi ---------------------------------------------------------------


def test_mot_dong_sai_thi_ca_file_khong_ghi(client):
    token = _token(client)
    wb = load_workbook(BytesIO(client.get("/api/employees/mau-nhap.xlsx",
                                          headers=_h(token)).content))
    del wb["Huong dan"]
    ws = wb[SHEET]
    cot = _cot(ws)
    _dat(ws, 2, cot, "Họ tên", "Dòng Đúng")
    _dat(ws, 2, cot, "Phòng/Tổ", "Hành chính nhân sự")
    _dat(ws, 2, cot, "Trạng thái", "Chính thức")
    _dat(ws, 3, cot, "Họ tên", "Dòng Sai")
    _dat(ws, 3, cot, "Phòng/Tổ", "Phòng Không Có Thật")
    _dat(ws, 3, cot, "Trạng thái", "Chính thức")

    kq = _nhap(client, token, wb, "commit").json()
    assert kq["hop_le"] is False
    assert kq["da_ghi"] is False
    assert any(l["dong"] == 3 and l["cot"] == "Phòng/Tổ" for l in kq["loi"]), kq["loi"]
    assert client.get("/api/employees?q=Dòng Đúng", headers=_h(token)).json()["total"] == 0


def test_ma_chua_co_thi_tao_moi_giu_nguyen_ma(client):
    """Chuyển dữ liệu sang máy khác thì mã NV phải theo người, không được cấp lại.

    Bản đầu 10/09/2026 báo lỗi khi mã chưa có (chống gõ nhầm mã). Thử nạp chính file xuất vào
    một DB TRẮNG thì 199/200 dòng đỏ hết vì file mang sẵn NV002…NV200; xoá tay cột Mã thì máy
    cấp lại từ NV001 ⇒ lệch mã của cả công ty. Mã NV nằm trên hợp đồng / thẻ / bảng lương.
    """
    token = _token(client)
    wb = load_workbook(BytesIO(client.get("/api/employees/mau-nhap.xlsx",
                                          headers=_h(token)).content))
    del wb["Huong dan"]
    ws = wb[SHEET]
    cot = _cot(ws)
    _dat(ws, 2, cot, "Mã", "NV999")
    _dat(ws, 2, cot, "Họ tên", "Từ Máy Khác")
    _dat(ws, 2, cot, "Phòng/Tổ", "Hành chính nhân sự")
    _dat(ws, 2, cot, "Trạng thái", "Chính thức")
    _dat(ws, 2, cot, "Ngày vào", "01/03/2024")

    kq = _nhap(client, token, wb, "commit").json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["tao_moi"] == 1
    ds = client.get("/api/employees?q=Từ Máy Khác", headers=_h(token)).json()
    assert [e["code"] for e in ds["items"]] == ["NV999"]

    # Người tiếp theo do máy cấp mã KHÔNG được đụng vào mã vừa nhập.
    tiep = _tao_nv(client, token, full_name="Người Sau Đó")
    assert tiep["code"] != "NV999"


def test_ma_da_co_thi_sua_dung_nguoi_do_chu_khong_de_them_ma_trung(client):
    """Mã đã có người dùng thì KHÔNG được đẻ thêm hồ sơ thứ hai cùng mã — phải sửa đúng người đó."""
    token = _token(client)
    nv = _tao_nv(client, token, full_name="Chủ Mã")
    wb = _tai_xuat(client, token)
    ws = wb[SHEET]
    cot = _cot(ws)
    dong = next(d for d in range(2, ws.max_row + 1)
                if ws.cell(row=d, column=cot["Mã"]).value == nv["code"])
    # Vẫn mã đó nhưng tên khác hẳn ⇒ đi vào nhánh CẬP NHẬT, không đẻ hồ sơ thứ hai cùng mã.
    _dat(ws, dong, cot, "Họ tên", "Tên Khác Hẳn")

    kq = _nhap(client, token, wb, "commit").json()
    assert kq["hop_le"] is True, kq["loi"]
    assert kq["tao_moi"] == 0 and kq["cap_nhat"] == 1
    ds = client.get(f"/api/employees?q={nv['code']}", headers=_h(token)).json()
    assert ds["total"] == 1
    assert ds["items"][0]["full_name"] == "Tên Khác Hẳn"


def test_thieu_ho_ten_khi_tao_moi_thi_bao_dong_nao(client):
    token = _token(client)
    wb = load_workbook(BytesIO(client.get("/api/employees/mau-nhap.xlsx",
                                          headers=_h(token)).content))
    del wb["Huong dan"]
    ws = wb[SHEET]
    cot = _cot(ws)
    _dat(ws, 2, cot, "Phòng/Tổ", "Hành chính nhân sự")
    _dat(ws, 2, cot, "Điện thoại", "0900000003")

    kq = _nhap(client, token, wb, "preview").json()
    assert kq["hop_le"] is False
    assert any(l["dong"] == 2 and l["cot"] == "Họ tên" for l in kq["loi"]), kq["loi"]


def test_khong_nhan_workbook_cua_man_khac(client):
    token = _token(client)
    wb = Workbook()
    wb.active.title = "Cong doan"
    meta = wb.create_sheet("_meta")
    meta["A1"], meta["B1"] = "loai", "cong_doan"
    r = _nhap(client, token, wb, "preview")
    assert r.status_code == 422, r.text


def test_nhap_can_quyen_them_va_sua(client):
    r = client.post(
        "/api/employees/import-excel?mode=preview",
        headers=_h(_sales_token()),
        files={"file": ("x.xlsx", b"x", "application/octet-stream")},
    )
    assert r.status_code == 403
