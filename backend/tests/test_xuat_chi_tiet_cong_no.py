"""File Excel SỔ CHI TIẾT công nợ — mọi đối tượng trong kỳ, một file (24/09/2026).

Kế toán chốt:
* cột "TK công nợ" để TRỐNG;
* cột "TK đối ứng" đi theo PHIẾU THU/CHI: tiền mặt ⇒ "Tiền mặt", chuyển khoản ⇒ tài khoản ngân
  hàng CỦA CÔNG TY mà tiền đi ra/vào. Dòng hàng về / hoá đơn bán không phải phiếu tiền ⇒ trống.

Bất biến giữ nguyên như sổ chi tiết xem trên màn: dòng "Cộng" của mỗi đối tượng = đúng ô dư cuối
kỳ của người đó bên sổ tổng hợp.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from tests.test_bao_cao_cong_no import _bao_cao, _dong, _ngay
from tests.test_payables_api import (
    _da_mua,
    _dong_dau_tien,
    _don,
    _ghi_dot,
    _headers,
    _phieu_chi,
    _supplier,
    _token_vai,
)


def _tai(client, headers, ben: str, *, tu: str, den: str, expect: int = 200):
    r = client.get(
        f"/api/accounting/reports/{ben}-detail.xlsx",
        params={"tu_ngay": tu, "den_ngay": den},
        headers=headers,
    )
    assert r.status_code == expect, r.text
    if expect != 200:
        return None
    assert "chi-tiet-cong-no-" in r.headers["content-disposition"]
    return load_workbook(BytesIO(r.content)).active


def _khoi(ws, ten_chua: str) -> list[tuple]:
    """Các dòng (A..I) thuộc khối của một đối tượng: từ dòng nhóm tới dòng "Cộng"."""
    hang = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=5, max_col=9)]
    dau = next(i for i, r in enumerate(hang) if r[0] and ten_chua in str(r[0]))
    cuoi = next(i for i in range(dau, len(hang)) if hang[i][2] == "Cộng")
    return hang[dau:cuoi + 1]


def _unc(client, headers, purchase_id: int, amount: int, *, delivery_id: int) -> dict:
    tk = client.post(
        "/api/accounting/company-bank-accounts",
        json={"account_holder": "CÔNG TY SAO VIỆT NHẬT", "account_number": "5550001",
              "bank_name": "Vietcombank", "bank_branch": "Hà Nội", "currency": "VND"},
        headers=headers,
    )
    assert tk.status_code == 201, tk.text
    r = client.post(
        "/api/accounting/payment-vouchers",
        json={
            "purchase_request_id": purchase_id, "voucher_type": "bank_transfer",
            "payment_stage": "final", "delivery_id": delivery_id,
            "voucher_date": _ngay(0), "amount": amount, "currency": "VND", "exchange_rate": 1,
            "content": "Chuyển khoản trả tiền giấy",
            "company_bank_account_id": tk.json()["id"],
            "beneficiary_account_holder": "NCC", "beneficiary_account_number": "9990001",
            "beneficiary_bank_name": "BIDV", "beneficiary_bank_branch": "Hà Nội",
            "bank_fee_bearer": "payer",
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


def test_tk_doi_ung_theo_phieu_va_tk_cong_no_de_trong(client):
    headers = _headers(client)
    ncc = _supplier(client, headers, name="NCC XuatCT TKDU")
    don = _don(client, headers, ncc["id"])
    _da_mua(client, headers, don["id"])
    dot = _ghi_dot(
        client, headers, don["id"],
        lines=[{"purchase_request_line_id": _dong_dau_tien(don), "quantity": 400}],
        ngay=_ngay(10),
    )  # 880.000 bên CÓ
    dot_id = dot["deliveries"][0]["id"]
    _phieu_chi(client, headers, don["id"], 300_000, stage="final", delivery_id=dot_id)
    _unc(client, headers, don["id"], 200_000, delivery_id=dot_id)

    ws = _tai(client, headers, "payables", tu=_ngay(30), den=_ngay(0))
    assert ws["A1"].value == "SỔ CHI TIẾT CÔNG NỢ PHẢI TRẢ"
    assert ws["D3"].value == "TK công nợ" and ws["E3"].value == "TK đối ứng"

    khoi = _khoi(ws, "NCC XuatCT TKDU")
    chung_tu = [r for r in khoi[2:-1]]           # bỏ dòng nhóm, dòng đầu kỳ, dòng Cộng
    assert len(chung_tu) == 3
    theo_dg = {r[2]: r for r in chung_tu}
    hang_ve = theo_dg["Hàng đã nhận"]
    tien_mat = theo_dg["Trả tiền giấy"]
    ck = theo_dg["Chuyển khoản trả tiền giấy"]
    assert hang_ve[4] in (None, ""), "hàng về không phải phiếu tiền ⇒ TK đối ứng trống"
    assert tien_mat[4] == "Tiền mặt"
    assert ck[4] == "Vietcombank - 5550001"
    assert all(r[3] in (None, "") for r in chung_tu), "TK công nợ phải để trống"
    assert (tien_mat[5], ck[5], hang_ve[6]) == (300_000, 200_000, 880_000)

    # Dòng "Cộng" khớp sổ tổng hợp.
    dong_th = _dong(_bao_cao(client, headers, "payables", tu=_ngay(30), den=_ngay(0)),
                    "NCC XuatCT TKDU")
    cong = khoi[-1]
    assert (cong[5], cong[6], cong[7] or 0, cong[8] or 0) == (
        dong_th["ps_no"], dong_th["ps_co"], dong_th["cuoi_no"], dong_th["cuoi_co"])
    assert (cong[7] or 0, cong[8]) == (0, 380_000)


def test_phai_thu_xuat_duoc(client):
    ws = _tai(client, _headers(client), "receivables", tu=_ngay(30), den=_ngay(0))
    assert ws["A1"].value == "SỔ CHI TIẾT CÔNG NỢ PHẢI THU"
    assert "Tài khoản: 131" in ws["A2"].value


def test_xuat_chi_tiet_doi_quyen_bao_cao(client):
    thieu = _token_vai("xuatct-khong-quyen", module="cong_no_phai_tra", can_read=True)
    _tai(client, {"Authorization": f"Bearer {thieu}"}, "payables",
         tu=_ngay(30), den=_ngay(0), expect=403)


def test_chon_mot_nguoi_thi_file_chi_co_khoi_cua_nguoi_do(client):
    headers = _headers(client)
    for ten in ("NCC XuatCT MotA", "NCC XuatCT MotB"):
        ncc = _supplier(client, headers, name=ten)
        don = _don(client, headers, ncc["id"])
        _da_mua(client, headers, don["id"])
        _ghi_dot(
            client, headers, don["id"],
            lines=[{"purchase_request_line_id": _dong_dau_tien(don), "quantity": 100}],
            ngay=_ngay(5),
        )
        if ten.endswith("A"):
            ncc_a = ncc

    r = client.get(
        "/api/accounting/reports/payables-detail.xlsx",
        params={"tu_ngay": _ngay(30), "den_ngay": _ngay(0), "doi_tuong_id": ncc_a["id"]},
        headers=headers,
    )
    assert r.status_code == 200, r.text
    ws = load_workbook(BytesIO(r.content)).active
    nhom = [row[0] for row in ws.iter_rows(min_row=5, max_col=1, values_only=True)
            if row[0] and "NCC XuatCT" in str(row[0])]
    assert nhom and all("MotA" in str(x) for x in nhom), nhom
    assert "Tổng cộng (1 đối tượng)" in [c for (c,) in ws.iter_rows(min_col=3, max_col=3, values_only=True)]
    # Tên file mang mã/tên người được chọn để tải nhiều người không đè nhau.
    assert "chi-tiet-cong-no-phai-tra-" in r.headers["content-disposition"]
    assert "-den-" in r.headers["content-disposition"]

    # Người không có chứng từ nào trong hệ ⇒ vẫn ra file hợp lệ (không 500).
    r2 = client.get(
        "/api/accounting/reports/payables-detail.xlsx",
        params={"tu_ngay": _ngay(30), "den_ngay": _ngay(0), "doi_tuong_id": 999999},
        headers=headers,
    )
    assert r2.status_code == 200, r2.text
