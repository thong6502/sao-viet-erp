"""Excel Hồ sơ nhân sự — ba ô lương: Lương cơ bản · Lương trách nhiệm · Mức đóng BHXH (23/09/2026).

Ba ô này nằm trên MỐC LƯƠNG (`employee_salaries`, có lịch sử), không nằm trên hồ sơ:
* Xuất: lấy mốc hiện hành.
* Nhập người MỚI: tạo mốc đầu hiệu lực từ ngày vào (như màn Thêm nhân viên).
* Nhập người ĐÃ CÓ: số khác mốc hiện hành ⇒ thêm mốc MỚI hiệu lực hôm nay, chép nguyên các ô
  lương khác (như Lương → Sửa lương). Ô trống = giữ nguyên.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.db import SessionLocal
from app.repositories.payroll_repo import PayrollRepository
from app.repositories.rbac_repo import DepartmentRepository

ADMIN = {"username": "admin", "password": "admin123"}
SHEET = "Nhan su"


def _h(client) -> dict[str, str]:
    tok = client.post("/api/auth/login", json=ADMIN).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


def _dept_id() -> int:
    db = SessionLocal()
    try:
        return DepartmentRepository(db).get_by_name("Hành chính nhân sự").id
    finally:
        db.close()


def _tao_nv(client, h, ten: str, **luong) -> dict:
    body = {"full_name": ten, "department_id": _dept_id(), "hire_date": "2024-01-15",
            "probation_end_date": "2025-12-31"}
    if luong:
        body["initial_salary"] = {"luong_vi_tri": 5_000_000, "luong_trach_nhiem": 1_000_000,
                                  "insurance_base": 5_500_000, "chuyen_can": 300_000,
                                  "union_member": True, **luong}
    r = client.post("/api/employees", json=body, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["employee"]


def _xuat(client, h) -> Workbook:
    r = client.get("/api/employees/export.xlsx", headers=h)
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))


def _nhap(client, h, wb: Workbook, mode: str = "commit"):
    buf = BytesIO()
    wb.save(buf)
    return client.post(
        f"/api/employees/import-excel?mode={mode}", headers=h,
        files={"file": ("nhan-su.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def _cot(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value}


def _dong_cua(ws, ten: str) -> int:
    cot = _cot(ws)
    for hang in range(2, ws.max_row + 1):
        if ws.cell(row=hang, column=cot["Họ tên"]).value == ten:
            return hang
    raise AssertionError(f"không thấy {ten!r} trong file")


def _moc(emp_id: int) -> list:
    db = SessionLocal()
    try:
        return PayrollRepository(db).list_salaries(emp_id)
    finally:
        db.close()


def test_xuat_co_ba_cot_luong_theo_moc_hien_hanh(client):
    h = _h(client)
    _tao_nv(client, h, "Có Lương", luong_vi_tri=6_000_000)
    _tao_nv(client, h, "Chưa Lương")
    ws = _xuat(client, h)[SHEET]
    cot = _cot(ws)
    for nhan in ("Lương cơ bản", "Lương trách nhiệm", "Mức đóng BHXH"):
        assert nhan in cot, f"thiếu cột {nhan!r}"
    d = _dong_cua(ws, "Có Lương")
    assert ws.cell(row=d, column=cot["Lương cơ bản"]).value == 6_000_000
    assert ws.cell(row=d, column=cot["Lương trách nhiệm"]).value == 1_000_000
    assert ws.cell(row=d, column=cot["Mức đóng BHXH"]).value == 5_500_000
    d = _dong_cua(ws, "Chưa Lương")
    assert ws.cell(row=d, column=cot["Lương cơ bản"]).value in (None, "")


def test_nhap_lai_file_vua_xuat_khong_de_moc_luong(client):
    h = _h(client)
    emp = _tao_nv(client, h, "Giữ Nguyên", luong_vi_tri=5_000_000)
    r = _nhap(client, h, _xuat(client, h))
    assert r.status_code == 200, r.text
    assert r.json()["hop_le"] and r.json()["cap_nhat"] == 0, r.json()
    assert len(_moc(emp["id"])) == 1


def test_sua_luong_nguoi_da_co_them_moc_moi_hom_nay_chep_o_khac(client):
    h = _h(client)
    emp = _tao_nv(client, h, "Tăng Lương", luong_vi_tri=5_000_000)
    wb = _xuat(client, h)
    ws = wb[SHEET]
    cot = _cot(ws)
    d = _dong_cua(ws, "Tăng Lương")
    ws.cell(row=d, column=cot["Lương cơ bản"]).value = "7.000.000"
    ws.cell(row=d, column=cot["Lương trách nhiệm"]).value = None     # trống = giữ nguyên

    xem = _nhap(client, h, wb, mode="preview").json()
    assert xem["hop_le"] and xem["cap_nhat"] == 1, xem
    assert len(_moc(emp["id"])) == 1, "xem trước không được ghi mốc lương"

    r = _nhap(client, h, wb).json()
    assert r["da_ghi"] and r["cap_nhat"] == 1, r
    mocs = _moc(emp["id"])
    assert len(mocs) == 2, "sửa lương phải THÊM mốc, không sửa đè"
    moi, cu = mocs[0], mocs[1]
    assert moi.effective_from == date.today()
    assert float(moi.luong_vi_tri) == 7_000_000
    assert float(moi.luong_trach_nhiem) == 1_000_000
    assert float(moi.insurance_base) == 5_500_000
    # Các ô lương khác chép nguyên từ mốc cũ.
    assert float(moi.chuyen_can) == 300_000 and moi.union_member is True
    assert float(cu.luong_vi_tri) == 5_000_000


def test_nhap_nguoi_moi_kem_luong_tao_moc_tu_ngay_vao(client):
    h = _h(client)
    wb = _xuat(client, h)
    ws = wb[SHEET]
    cot = _cot(ws)
    hang = ws.max_row + 1
    ws.cell(row=hang, column=cot["Họ tên"]).value = "Người Mới Excel"
    ws.cell(row=hang, column=cot["Phòng/Tổ"]).value = "Hành chính nhân sự"
    ws.cell(row=hang, column=cot["Ngày vào"]).value = "01/09/2026"
    ws.cell(row=hang, column=cot["Ngày hết thử việc"]).value = "31/10/2026"
    ws.cell(row=hang, column=cot["Lương cơ bản"]).value = 4_800_000
    ws.cell(row=hang, column=cot["Lương trách nhiệm"]).value = 200_000
    r = _nhap(client, h, wb).json()
    assert r["da_ghi"] and r["tao_moi"] == 1, r

    ws2 = _xuat(client, h)[SHEET]
    d = _dong_cua(ws2, "Người Mới Excel")
    ma = ws2.cell(row=d, column=_cot(ws2)["Mã"]).value
    db = SessionLocal()
    try:
        from app.repositories.employee_repo import EmployeeRepository
        emp_id = EmployeeRepository(db).find_by_code(ma).id
    finally:
        db.close()
    (moc,) = _moc(emp_id)
    assert moc.effective_from == date(2026, 9, 1)
    assert float(moc.luong_vi_tri) == 4_800_000
    # Mức đóng BHXH bỏ trống ⇒ lấy bằng cơ bản + trách nhiệm (như màn Thêm nhân viên).
    assert float(moc.insurance_base) == 5_000_000


def test_luong_sai_thi_bao_dung_o_va_ca_file_khong_ghi(client):
    h = _h(client)
    emp = _tao_nv(client, h, "Sai Lương", luong_vi_tri=5_000_000)
    _tao_nv(client, h, "Chưa Có Mốc")
    wb = _xuat(client, h)
    ws = wb[SHEET]
    cot = _cot(ws)
    ws.cell(row=_dong_cua(ws, "Sai Lương"), column=cot["Mức đóng BHXH"]).value = 0
    ws.cell(row=_dong_cua(ws, "Chưa Có Mốc"), column=cot["Lương trách nhiệm"]).value = 500_000
    ws.cell(row=_dong_cua(ws, "Sai Lương"), column=cot["Chức danh"]).value = "Đổi chức danh"
    r = _nhap(client, h, wb).json()
    assert not r["hop_le"] and not r["da_ghi"], r
    loi = {(x["cot"], x["dong"]) for x in r["loi"]}
    assert ("Mức đóng BHXH", _dong_cua(ws, "Sai Lương")) in loi
    assert ("Lương cơ bản", _dong_cua(ws, "Chưa Có Mốc")) in loi
    assert len(_moc(emp["id"])) == 1


def test_o_tien_go_chu_thi_bao_loi(client):
    h = _h(client)
    _tao_nv(client, h, "Gõ Chữ", luong_vi_tri=5_000_000)
    wb = _xuat(client, h)
    ws = wb[SHEET]
    ws.cell(row=_dong_cua(ws, "Gõ Chữ"), column=_cot(ws)["Lương cơ bản"]).value = "năm triệu"
    r = _nhap(client, h, wb, mode="preview").json()
    assert not r["hop_le"]
    assert any(x["cot"] == "Lương cơ bản" for x in r["loi"]), r
