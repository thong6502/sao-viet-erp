"""Bảng công tháng xuất .xlsx — thay bản .csv (chủ chốt 09/09/2026).

Hai chỗ chủ chê ở bản .csv và phải KHÔNG được quay lại:
  1. ô ngày không ra được công/giờ in chữ "có" (ngày mới xếp ca, nghỉ luân phiên, ngày treo đều
     chung một chữ vô nghĩa) — giờ mỗi loại một ký hiệu, ngày không đi làm để TRỐNG;
  2. tăng ca không hiện ở đâu — giờ ngày có tăng ca đeo "+2h" ngay trong ô, và có cột tổng
     "Tăng ca (giờ)".
"""
from io import BytesIO

from openpyxl import load_workbook

from app.services.bang_cong_excel import (
    gio_tang_ca,
    loc_nhan_vien,
    o_ngay,
    xuat_bang_cong,
)


def _ngay(**kw) -> dict:
    o = {"shift_id": None, "shift_name": None, "first_in": None, "last_out": None, "hours": None,
         "present": False, "cong": None, "late": False, "early": False, "ot_minutes": 0,
         "ot_thieu_cap": False, "night": False, "leave": None, "leave_paid": False,
         "holiday": False, "restday": False, "plain": False, "planned_off": False}
    o.update(kw)
    return o


def _row(days: dict, **kw) -> dict:
    r = {"employee_id": 1, "employee_code": "NV001", "employee_name": "Nguyễn Văn A",
         "department_id": 3, "department_name": "Tổ In", "shift_id": 1, "shift_name": "Ca ngày",
         "days": days, "total_days": len(days), "total_leave": 0, "paid_leave_days": 0,
         "total_hours": 8.0, "total_cong": 1.0}
    r.update(kw)
    return r


# --- ký hiệu ô ngày ---------------------------------------------------------


def test_ngay_khong_co_cong_ca_de_trong_chu_khong_phai_chu_co():
    """Ô do service dựng sẵn (đã xếp ca / chưa tới) KHÔNG được in "có" nữa."""
    assert o_ngay(_ngay(shift_name="Ca ngày")) == ""
    assert o_ngay(None) == ""
    assert o_ngay({}) == ""


def test_ky_hieu_tung_loai_ngay():
    assert o_ngay(_ngay(first_in="07:55", last_out="17:05", cong=1)) == "1"
    assert o_ngay(_ngay(first_in="07:55", last_out="11:30", cong=0.5)) == "0.5"
    # chưa gán ca → số giờ
    assert o_ngay(_ngay(first_in="08:00", last_out="16:00", hours=8)) == "8h"
    # có lượt bấm mà không ra công/giờ = ngày treo (thiếu lượt RA)
    assert o_ngay(_ngay(first_in="08:00")) == "?"
    assert o_ngay(_ngay(leave="Nghỉ phép năm", leave_paid=True)) == "P"
    assert o_ngay(_ngay(leave="Nghỉ việc riêng", leave_paid=False)) == "KL"
    # lễ nghỉ ở nhà: "L", KHÔNG phải "P" — lễ không tiêu ngày phép năm
    assert o_ngay(_ngay(holiday=True, leave="Quốc khánh", leave_paid=True)) == "L"
    assert o_ngay(_ngay(planned_off=True)) == "NL"


def test_ngay_le_di_lam_van_ra_so_cong():
    """Hỏi LƯỢT BẤM trước cờ ngày — hỏi cờ trước là nuốt mất công thật của ngày lễ đi làm."""
    assert o_ngay(_ngay(first_in="08:00", last_out="17:00", cong=1,
                        holiday=True, leave="Quốc khánh")) == "1"


def test_ngay_co_tang_ca_hien_gio_ngay_trong_o():
    assert o_ngay(_ngay(first_in="08:00", last_out="19:00", cong=1, ot_minutes=120)) == "1 +2h"
    assert o_ngay(_ngay(first_in="08:00", last_out="18:30", cong=1, ot_minutes=90)) == "1 +1.5h"
    # phiếu tăng ca đã duyệt mà chưa có cặp bấm ⇒ đang tính 0 giờ, phải thấy để còn chấm bù
    assert o_ngay(_ngay(first_in="08:00", last_out="17:00", cong=1, ot_thieu_cap=True)) == "1 TC?"


def test_gio_tang_ca_cong_ca_thang():
    r = _row({"1": _ngay(ot_minutes=120), "2": _ngay(ot_minutes=90), "3": _ngay()})
    assert gio_tang_ca(r) == 3.5


# --- file .xlsx -------------------------------------------------------------


def _wb(rows, **kw):
    tham_so = {"nam": 2026, "thang": 9, "so_ngay": 30}
    tham_so.update(kw)
    return load_workbook(BytesIO(xuat_bang_cong(rows, **tham_so))).active


def test_khuon_file_va_cot_tang_ca():
    ws = _wb([_row({
        "1": _ngay(first_in="07:55", last_out="17:05", cong=1),
        "2": _ngay(first_in="07:55", last_out="19:05", cong=1, ot_minutes=120),
        "3": _ngay(shift_name="Ca ngày"),          # đã xếp ca, chưa tới
        "5": _ngay(leave="Nghỉ phép năm", leave_paid=True),
    }, total_cong=2.0, total_hours=21.0)], ten_bo_phan="Tổ In", cong_chuan=26)

    assert ws.cell(row=1, column=1).value == "BẢNG CÔNG THÁNG 09/2026"
    assert "Tổ In" in ws.cell(row=2, column=1).value
    assert "Công chuẩn tháng: 26" in ws.cell(row=2, column=1).value
    assert "ô trống = không đi làm" in ws.cell(row=3, column=1).value

    # 4 cột thông tin + 30 cột ngày + 4 cột tổng
    assert [ws.cell(row=5, column=i).value for i in range(1, 5)] == ["Mã", "Họ tên", "Phòng/Tổ", "Ca"]
    assert [ws.cell(row=5, column=i).value for i in (35, 36, 37, 38)] == \
        ["Số công", "Công CN/Lễ", "Tăng ca (giờ)", "Tổng giờ"]
    assert ws.cell(row=6, column=5).value == 1 and ws.cell(row=6, column=34).value == 30
    assert ws.cell(row=5, column=5).value == "T3"   # 01/09/2026 là thứ Ba

    assert ws.cell(row=7, column=1).value == "NV001"
    assert ws.cell(row=7, column=5).value == "1"        # ngày 1
    assert ws.cell(row=7, column=6).value == "1 +2h"    # ngày 2 có tăng ca
    assert ws.cell(row=7, column=7).value is None       # ngày 3 xếp ca mà chưa tới ⇒ TRỐNG
    assert ws.cell(row=7, column=9).value == "P"        # ngày 5 nghỉ phép
    assert ws.cell(row=7, column=35).value == 2.0     # số công
    assert ws.cell(row=7, column=36).value == 0.0     # công CN/Lễ
    assert ws.cell(row=7, column=37).value == 2.0     # tăng ca (giờ)
    assert ws.cell(row=7, column=38).value == 21.0    # tổng giờ

    assert ws.cell(row=8, column=1).value == "TỔNG"
    assert [ws.cell(row=8, column=i).value for i in (35, 36, 37, 38)] == [2.0, 0.0, 2.0, 21.0]


def test_ngay_le_giu_chu_thu_va_goi_ten_le_o_dau_file():
    """Cột ngày lễ chỉ TÔ NỀN; chữ thứ phải còn, tên lễ nói ở dòng 2."""
    ws = _wb([_row({"2": _ngay(holiday=True, leave="Quốc khánh", leave_paid=True)},
                   total_cong=1.0, total_hours=0.0)],
             ngay_le={2: "Quốc khánh"})
    assert "Ngày lễ: 02/09 Quốc khánh" in ws.cell(row=2, column=1).value
    assert ws.cell(row=5, column=6).value == "T4"     # 02/09/2026 là thứ Tư
    assert ws.cell(row=7, column=6).value == "L"
    assert ws.cell(row=7, column=6).fill.fgColor.rgb == "FFFDE7E9"
    # thứ Bảy 05/09 tô nền xám nhạt
    assert ws.cell(row=7, column=9).fill.fgColor.rgb == "FFF1F3F5"


def test_khong_con_chu_co_trong_file():
    ws = _wb([_row({"1": _ngay(shift_name="Ca ngày"), "2": _ngay(planned_off=True)},
                   total_cong=0.0, total_hours=0.0)])
    o_du_lieu = [c.value for h in ws.iter_rows(min_row=7, max_row=7) for c in h]
    assert "có" not in o_du_lieu


def test_khong_co_total_cong_thi_lay_total_days():
    ws = _wb([_row({"1": _ngay(first_in="08:00", last_out="17:00", hours=8)},
                   total_cong=None, total_days=1, total_hours=8.0)])
    assert ws.cell(row=7, column=35).value == 1.0


def test_cot_cong_cn_le_cong_ba_loai_ngay():
    """1 công lễ + 0.94 công Chủ nhật ⇒ cột CN/Lễ nói 1.94 (chủ 09/09/2026)."""
    ws = _wb([_row({}, holiday_cong=1.0, restday_cong=0.94, plain_cong=0.0)])
    assert ws.cell(row=7, column=36).value == 1.94


# --- ô tìm tên / mã NV ------------------------------------------------------


def test_loc_nhan_vien_bo_dau_va_theo_ma():
    rows = [_row({}, employee_code="NV001", employee_name="Nguyễn Văn A"),
            _row({}, employee_code="NV002", employee_name="Cao Minh Quân")]
    assert [r["employee_code"] for r in loc_nhan_vien(rows, "quan")] == ["NV002"]   # gõ không dấu
    assert [r["employee_code"] for r in loc_nhan_vien(rows, "QUÂN")] == ["NV002"]   # hoa + có dấu
    assert [r["employee_code"] for r in loc_nhan_vien(rows, "nv00")] == ["NV001", "NV002"]
    assert loc_nhan_vien(rows, "  ") == rows      # ô tìm trống ⇒ giữ nguyên
    assert loc_nhan_vien(rows, None) == rows
    assert loc_nhan_vien(rows, "không có ai") == []


def test_file_noi_ro_dang_loc_theo_gi():
    ws = _wb([_row({"1": _ngay(first_in="08:00", last_out="17:00", cong=1)})], tim="Quân")
    assert "Lọc theo: Quân" in ws.cell(row=2, column=1).value


def test_bang_rong_van_xuat_duoc():
    ws = _wb([], thang=2, so_ngay=28)
    assert ws.cell(row=1, column=1).value == "BẢNG CÔNG THÁNG 02/2026"
    assert ws.cell(row=7, column=1).value == "TỔNG"
