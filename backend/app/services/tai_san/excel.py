"""Xuất bảng khấu hao / phân bổ một kỳ ra .xlsx.

Bố cục cố ý ĐƠN GIẢN, không bắt chước khuôn MISA như báo cáo công nợ: file này không dán vào bộ
hồ sơ nào cả, nó là bảng để kế toán đọc rồi tự gõ định khoản sang phần mềm kế toán. Cột cuối là
Còn lại — cột ghi chú hạch toán đã gỡ (mg 0278), định khoản ai cần nhớ thì ghi ở ô ghi chú của
chính tài sản.

    A1  BẢNG TRÍCH KHẤU HAO / PHÂN BỔ — kỳ MM/YYYY
    A2  (trống)
    A3  tiêu đề cột
    A4+ dữ liệu
    cuối: dòng TỔNG
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DINH_DANG_TIEN = "#,##0"

TIEU_DE = [
    "Mã", "Tên tài sản", "Bộ phận", "Nguyên giá",
    "Trích kỳ này", "Lũy kế", "Còn lại",
]
#: Khoá trong dict `KyService.bang()` theo đúng thứ tự cột ở trên.
KHOA = ["ma", "ten", "bo_phan_ten", "nguyen_gia", "muc_trich", "luy_ke", "con_lai"]
COT_TIEN = {4, 5, 6, 7}
RONG_COT = {"A": 12, "B": 34, "C": 18, "D": 16, "E": 15, "F": 16, "G": 16}


def xuat_bang_ky(rows: list[dict], *, nam: int, thang: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Khau hao {thang:02d}-{nam}"

    ws.cell(row=1, column=1, value=f"BẢNG TRÍCH KHẤU HAO / PHÂN BỔ — kỳ {thang:02d}/{nam}")
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TIEU_DE))
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    for i, ten in enumerate(TIEU_DE, start=1):
        o = ws.cell(row=3, column=i, value=ten)
        o.font = Font(bold=True)
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    dong = 4
    for r in rows:
        for i, khoa in enumerate(KHOA, start=1):
            gia_tri = r.get(khoa)
            o = ws.cell(row=dong, column=i, value=gia_tri)
            if i in COT_TIEN:
                o.number_format = DINH_DANG_TIEN
        dong += 1

    o = ws.cell(row=dong, column=1, value="TỔNG")
    o.font = Font(bold=True)
    for i in COT_TIEN:
        khoa = KHOA[i - 1]
        o = ws.cell(row=dong, column=i, value=sum(int(r.get(khoa) or 0) for r in rows))
        o.font = Font(bold=True)
        o.number_format = DINH_DANG_TIEN

    for chu, rong in RONG_COT.items():
        ws.column_dimensions[chu].width = rong
    ws.freeze_panes = "A4"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
