"""Xuất BẢNG LƯƠNG THÁNG ra .xlsx theo ĐÚNG khuôn bảng lương công ty đang dùng (chủ chốt 17/09/2026).

Khuôn gốc: sheet `BL CT` của `BẢNG LƯƠNG T05.2026 (duyệt).xlsx`. Chủ đặt cột theo thứ tự:

    Số TT · MNV · Họ và tên · Chức vụ · NCT · CN/Lễ · Tổng NC · Tăng ca · Lương BHXH · Lương cơ bản ·
    Lương trách nhiệm · các khoản phụ cấp · Chuyên cần · Lương thời gian · Tổng lương · (các khoản trừ) ·
    Thực nhận

và để phần trừ cho mình tự xếp. Bản 09/09/2026 (60 cột tự đặt) bị thay bằng khuôn này.

Đọc cho đúng nghĩa từng khối — y như file của công ty:

- **Mức lương tháng** (Lương BHXH · Lương cơ bản · Lương trách nhiệm · từng khoản phụ cấp) là SỐ THÁNG
  để tham chiếu, KHÔNG cộng vào Tổng lương. Tiền thật của chúng nằm trong "Lương thời gian".
- **Lương thời gian** = lương theo công (tháng lấy bù lỗ: trọn bù lỗ, đã gồm tiền khoán / km) + công lễ
  nghỉ + phụ cấp trả theo công + phần thêm làm nguyên ngày CN / lễ + tiền ngày off1x. Bảng công ty:
  `X = (vị trí + trách nhiệm + phụ cấp) ÷ 26 × Tổng NC` — "Tổng NC" ở đây là đúng số công đó.
- **Ngoài giờ/Tăng ca** chỉ là tiền GIỜ tăng ca (`payroll_lines.tien_gio_tang_ca`, mg 0305).

⚠️ LUẬT SỐ HỌC — có test khoá (`test_luong_excel.py`), đừng phá:

    Tổng lương  = Σ cột khối THU (gồm Chuyên cần, Lương thời gian) = tổng thu TRƯỚC phạt
    Thực nhận   = Tổng lương − Σ cột khối TRỪ

Phạt ghi ở cột trừ là số GHI NHẬN; phần engine không trừ (trần 30% Điều 102, hoặc khấu trừ lớn hơn lương)
hiện ở cột "Không trừ được kỳ này" (số âm) — cột này chỉ có khi có người rơi vào. Thêm khoản mới vào
engine thì thêm vào một cột Ở ĐÂY, nếu không hai vế lệch và test đỏ ngay.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

TIEN = "#,##0"
#: `General` chứ không phải `0.##`: dạng đó làm Excel in "26." (thừa dấu chấm) cho số công tròn.
CONG = "General"

NEN_NHOM = {
    "tt": "FFEFF2F5",     # thông tin
    "cong": "FFEAF3EA",   # ngày công
    "muc": "FFF1EEF8",    # mức lương tháng (tham chiếu)
    "thu": "FFFDF6E3",    # các khoản thu nhập
    "tong": "FFFBE3B8",   # tổng lương
    "tru": "FFFBEAEA",    # các khoản trừ
    "cuoi": "FFE3F2FD",   # thực nhận
    "phu": "FFEFF2F5",    # ghi chú sau thực nhận
}
NHAN_NHOM = {
    "tt": "THÔNG TIN NHÂN VIÊN",
    "cong": "NGÀY CÔNG",
    "muc": "MỨC LƯƠNG THÁNG (tham chiếu, không cộng)",
    "thu": "CÁC KHOẢN THU NHẬP",
    "tong": "",
    "tru": "CÁC KHOẢN TRỪ",
    "cuoi": "",
    "phu": "",
}
#: Bảng lương công ty tô VÀNG dòng người thử việc.
NEN_THU_VIEC = "FFFFF4C2"
VIEN = Border(*(Side(style="thin", color="FFCED4DA"),) * 4)


def _f(v) -> float:
    return float(v or 0)


@dataclass(frozen=True)
class Cot:
    ten: str
    #: tt | cong | muc | thu | tong | tru | cuoi | phu
    nhom: str
    #: (dòng lương, thông tin NV, ngữ cảnh) → giá trị; None = điền ở `dong_so`.
    lay: Callable | None
    fmt: str | None = TIEN


# --- đọc số từ dòng lương -----------------------------------------------------------------------


def _lay_bu_lo(ln) -> bool:
    """Tháng này người đó LẤY bù lỗ theo công (tổ khoán / tổ Giao hàng) — xem `PayrollService._compute`."""
    return bool(getattr(ln, "lay_bu_lo", False))


def _luong_cong_hien(ln) -> float:
    """Lương theo công để in: tháng lấy bù lỗ = trọn bù lỗ (khoán + km + phần bù thêm) — hai khoản THAY
    nhau, in cả hai là kế toán đọc thành cộng dồn (chủ chốt 16/09/2026)."""
    if not _lay_bu_lo(ln):
        return _f(getattr(ln, "luong_cong", 0))
    return _f(getattr(ln, "khoan", 0)) + _f(getattr(ln, "khoan_km", 0)) + _f(getattr(ln, "luong_cong", 0))


def _khoan_phat_sinh_thu(ln) -> float:
    """Khoản danh mục PHÁT SINH riêng kỳ này (nguồn `line`, loại thu) — nằm NGOÀI `allowance`."""
    return sum(_f(c.amount) for c in getattr(ln, "components", []) or []
               if getattr(c, "kind", "") != "tru" and getattr(c, "source", "") == "line")


def _khoan_tru(ln) -> float:
    """Khoản danh mục loại TRỪ (mọi nguồn) — engine trừ thẳng vào thực nhận."""
    return sum(_f(c.amount) for c in getattr(ln, "components", []) or []
               if getattr(c, "kind", "") == "tru")


def _tien_gio_tang_ca(ln) -> float:
    """Cột "Ngoài giờ/Tăng ca" = tiền GIỜ tăng ca, không gồm phần thêm ngày CN / lễ.

    Dòng tính trước mg 0305 (NULL) chưa tách được: người chế độ khoán (trừ Giao hàng) thì giờ tăng ca
    vốn 0đ nên `ot_pay` toàn là phần thêm CN / lễ; còn lại để trọn `ot_pay` trừ tiền ngày off1x ở cột
    này. Tổng lương không đổi — chỉ khác số nằm ở cột nào."""
    v = getattr(ln, "tien_gio_tang_ca", None)
    if v is not None:
        return _f(v)
    if getattr(ln, "che_do_khoan", False) and not getattr(ln, "la_giao_hang", False):
        return 0.0
    return max(0.0, _f(getattr(ln, "ot_pay", 0)) - _f(getattr(ln, "off1x_pay", 0)))


def _luong_thoi_gian(ln) -> float:
    """Cột X "Lương thời gian (lễ, chủ nhật)" của bảng công ty — xem docstring module."""
    return (_luong_cong_hien(ln) + _f(getattr(ln, "luong_ngay_le", 0)) + _f(getattr(ln, "allowance", 0))
            + _f(getattr(ln, "ot_pay", 0)) - _tien_gio_tang_ca(ln))


def _tong_nc(ln, ty_le_thu_viec: float) -> float:
    """Số công hưởng lương thời gian — bảng công ty: `Tổng NC = NCT + CN × 2` (người sản lượng: CN + lễ).

    Chính là số công engine đã trả phụ cấp theo công (`cong_phu_cap`: công theo lương + phần thêm CN / lễ +
    ngày off1x; người khoán chỉ còn công lễ nghỉ + phần thêm). Tháng LẤY BÙ LỖ phụ cấp nằm trong số bù lỗ
    nên cộng thêm số công bù lỗ, chia ngược từ tiền: bù lỗ = (mức nền × tỉ lệ + phụ cấp) ÷ công chuẩn × công.
    Kỳ tính trước 15/09/2026 chưa chụp số công này ⇒ dùng tổng công."""
    cpc = getattr(ln, "cong_phu_cap", None)
    if cpc is None:
        return round(_f(getattr(ln, "actual_cong", 0)), 2)
    tong = _f(cpc)
    if getattr(ln, "bu_lo_theo_cong", None) is not None and _lay_bu_lo(ln):
        ty_le = float(ty_le_thu_viec) if getattr(ln, "is_probation", False) else 1.0
        cong_chuan = _f(getattr(ln, "standard_cong", 0)) or 1.0
        don_gia = (_f(getattr(ln, "monthly_salary", 0)) * ty_le
                   + _f(getattr(ln, "phu_cap_thang", 0))) / cong_chuan
        if don_gia > 0:
            tong += _f(ln.bu_lo_theo_cong) / don_gia
    return round(tong, 2)


def _phu_cap_theo_khoan(ln, nv: dict) -> dict[str, float]:
    """Mức THÁNG từng khoản phụ cấp gán ở hồ sơ.

    Dòng lương chụp SỐ TRẢ (= mức tháng × công hưởng ÷ công chuẩn) ⇒ chia ngược ra đúng mức của kỳ đó. Không
    chia được (tháng lấy bù lỗ, cả tháng không công, HCNS đè tay số kỳ này) thì đọc mức đang gán ở hồ sơ.
    Kỳ tính trước 15/09/2026 phụ cấp còn cộng phẳng ⇒ số trả chính là mức tháng."""
    ho_so = dict(nv.get("khoan_ho_so") or {})
    cpc = getattr(ln, "cong_phu_cap", None)
    cong_chuan = _f(getattr(ln, "standard_cong", 0))
    out: dict[str, float] = {}
    for c in getattr(ln, "components", []) or []:
        if getattr(c, "kind", "") == "tru" or getattr(c, "source", "") != "employee":
            continue
        if cpc is None:
            so = _f(c.amount)
        elif _f(cpc) > 0 and cong_chuan > 0 and not getattr(c, "da_de_tay", False):
            chia = _f(c.amount) * cong_chuan / _f(cpc)
            # Số trả làm tròn tới đồng, `cong_phu_cap` tới 0,01 công ⇒ chia ngược lệch vài chục đồng khi công
            # hưởng ít (NV002 kỳ 09/2026: 56.769 × 26 ÷ 7,38 = 199.999). Sát mức hồ sơ trong sai số đó thì lấy
            # đúng mức hồ sơ — số tròn như lúc khai.
            sai_so = chia * 0.005 / _f(cpc) + 0.5 * cong_chuan / _f(cpc) + 1
            so = (ho_so[c.name] if c.name in ho_so and abs(chia - ho_so[c.name]) <= sai_so
                  else round(chia))
        else:
            so = ho_so.get(c.name, _f(c.amount))
        out[c.name] = out.get(c.name, 0.0) + so
    if cpc is not None:
        for ten, so in ho_so.items():
            out.setdefault(ten, so)
    return out


def _muc_nen(ln, nv: dict) -> tuple[float, float]:
    """(Lương cơ bản, Lương trách nhiệm) tháng theo mốc lương của kỳ. Mốc lệch số đã chụp (`monthly_salary`
    — dữ liệu cũ chỉ khai một số tổng, hoặc mốc bị sửa sau khi tính) thì theo số đã chụp: cả vào cơ bản."""
    chup = _f(getattr(ln, "monthly_salary", 0))
    vt, tn = nv.get("luong_vi_tri"), nv.get("luong_trach_nhiem")
    if vt is None or tn is None or abs(_f(vt) + _f(tn) - chup) > 1:
        return chup, 0.0
    return _f(vt), _f(tn)


def _phu_cap_khac_thang(ln, nv: dict, khoan: dict[str, float]) -> float:
    """Mức tháng ô "Phụ cấp khác" của mốc lương. Tổng phụ cấp khai đã chụp (`phu_cap_thang`) thắng khi lệch."""
    pc = nv.get("phu_cap_khac")
    chup = getattr(ln, "phu_cap_thang", None)
    if chup is not None:
        con_lai = _f(chup) - sum(khoan.values())
        if pc is None or abs(con_lai - _f(pc)) > 1:
            return max(0.0, round(con_lai))
    return _f(pc)


def _tam_ung_tru_ky_nay(ln) -> float:
    """Tạm ứng + lương đợt 1 + nợ kỳ trước ĐÃ TRỪ kỳ này = tổng phải trừ − phần chuyển sang kỳ sau."""
    phai_tru = (_f(getattr(ln, "advance_total", 0)) + _f(getattr(ln, "luong_dot_1_total", 0))
                + _f(getattr(ln, "no_ung_ky_truoc", 0)))
    return round(max(0.0, phai_tru - _f(getattr(ln, "no_ung_chuyen_ky_sau", 0))))


# --- khuôn cột ----------------------------------------------------------------------------------

KHONG_TRU = "Không trừ được kỳ này"


def cot_bang_luong(lines, nhan_vien: dict, *, bh_tach, ty_le_thu_viec: float = 1.0) -> list[Cot]:
    """Danh sách cột của sheet Bảng lương. Cột từng khoản phụ cấp và cột "Không trừ được kỳ này" sinh
    theo dữ liệu của chính kỳ đó (không ai có thì không có cột)."""
    lines = list(lines)
    khoan_ten: dict[str, None] = {}
    for ln in lines:
        for ten, so in _phu_cap_theo_khoan(ln, _nv(nhan_vien, ln)).items():
            if so:
                khoan_ten.setdefault(ten, None)

    cot = [
        Cot("Số TT", "tt", None, None),
        Cot("MNV", "tt", lambda ln, nv, ctx: getattr(ln, "employee_code", "") or "", None),
        Cot("Họ và tên", "tt", lambda ln, nv, ctx: getattr(ln, "employee_name", "") or "", None),
        Cot("Chức vụ", "tt", lambda ln, nv, ctx: nv.get("chuc_vu") or "", None),

        # NCT = ngày công thực tế (ngày thường + phép + lễ nghỉ); CN/Lễ = công ngày CN / lễ có đi làm.
        Cot("NCT", "cong", lambda ln, nv, ctx: round(
            max(0.0, _f(ln.actual_cong) - _f(getattr(ln, "special_cong", 0))), 2), CONG),
        Cot("CN/Lễ", "cong", lambda ln, nv, ctx: round(_f(getattr(ln, "special_cong", 0)), 2), CONG),
        Cot("Tổng NC", "cong", lambda ln, nv, ctx: _tong_nc(ln, ty_le_thu_viec), CONG),
        Cot("Tăng ca", "cong", lambda ln, nv, ctx: round(int(getattr(ln, "ot_minutes", 0) or 0) / 60, 2),
            CONG),

        Cot("Lương BHXH", "muc", lambda ln, nv, ctx: _f(getattr(ln, "insurance_base", 0))),
        Cot("Lương cơ bản", "muc", lambda ln, nv, ctx: _muc_nen(ln, nv)[0]),
        Cot("Lương trách nhiệm", "muc", lambda ln, nv, ctx: _muc_nen(ln, nv)[1]),
        Cot("Phụ cấp khác", "muc", lambda ln, nv, ctx: _phu_cap_khac_thang(ln, nv, ctx["khoan"])),
    ]
    for ten in khoan_ten:
        cot.append(Cot(ten, "muc", lambda ln, nv, ctx, ten=ten: _f(ctx["khoan"].get(ten))))

    cot += [
        Cot("Phép năm", "thu", lambda ln, nv, ctx: _f(getattr(ln, "phep_nam", 0))),
        Cot("Ngoài giờ/Tăng ca", "thu", lambda ln, nv, ctx: _tien_gio_tang_ca(ln)),
        # Bảng công ty gộp lương kinh doanh (hoa hồng) và sản lượng vào một cột. Tháng lấy bù lỗ thì tiền
        # khoán / km đã nằm trong "Lương thời gian".
        Cot("Lương kinh doanh/Sản lượng", "thu", lambda ln, nv, ctx: _f(getattr(ln, "hoa_hong", 0)) + (
            0.0 if _lay_bu_lo(ln) else _f(getattr(ln, "khoan", 0)) + _f(getattr(ln, "khoan_km", 0)))),
        Cot("Cơm/Phụ cấp ca đêm", "thu", lambda ln, nv, ctx: (
            _f(getattr(ln, "meal_allowance_pay", 0)) + _f(getattr(ln, "com_tang_ca_pay", 0))
            + _f(getattr(ln, "shift_allowance_pay", 0)) + _f(getattr(ln, "night_premium_pay", 0))
            + _f(getattr(ln, "night_pay", 0)))),
        # Thưởng khai qua danh mục (khoản phát sinh kỳ này) + các ô thưởng cũ của dòng lương.
        Cot("Thưởng/khoản phát sinh", "thu", lambda ln, nv, ctx: (
            _khoan_phat_sinh_thu(ln) + _f(getattr(ln, "other_bonus", 0))
            + _f(getattr(ln, "thuong_thanh_tich", 0)) + _f(getattr(ln, "thuong_doanh_so", 0))
            + _f(getattr(ln, "thuong_5s", 0)) + _f(getattr(ln, "tra_dong_phuc", 0)))),
        Cot("Điều chỉnh lương", "thu", lambda ln, nv, ctx: _f(getattr(ln, "dieu_chinh_luong", 0))),
        Cot("Chuyên cần", "thu", lambda ln, nv, ctx: _f(getattr(ln, "chuyen_can", 0))),
        Cot("Lương thời gian", "thu", lambda ln, nv, ctx: _luong_thoi_gian(ln)),
        Cot("Tổng lương", "tong", None),

        # BHXH / BHYT / BHTN tách bằng ĐÚNG hàm phiếu lương dùng — ba cột luôn cộng đúng tổng đã đóng băng.
        Cot("BHXH", "tru", lambda ln, nv, ctx: round(ctx["bh3"][0])),
        Cot("BHYT", "tru", lambda ln, nv, ctx: round(ctx["bh3"][1])),
        Cot("BHTN", "tru", lambda ln, nv, ctx: round(ctx["bh3"][2])),
        Cot("Công đoàn", "tru", lambda ln, nv, ctx: _f(getattr(ln, "cong_doan", 0))),
        Cot("Tạm ứng/Lương đợt 1", "tru", lambda ln, nv, ctx: _tam_ung_tru_ky_nay(ln)),
        Cot("Đi trễ/về sớm", "tru", lambda ln, nv, ctx: _f(getattr(ln, "di_tre", 0))),
        Cot("Thuế TNCN", "tru", lambda ln, nv, ctx: _f(getattr(ln, "pit", 0))),
        Cot("Phạt biên bản vi phạm", "tru", lambda ln, nv, ctx: (
            _f(getattr(ln, "phat_bien_ban", 0)) + _f(getattr(ln, "vi_pham", 0)))),
        Cot("ĐT vượt trội", "tru", lambda ln, nv, ctx: _f(getattr(ln, "dt_vuot_troi", 0))),
        Cot("Đồng phục, phạt 5S", "tru", lambda ln, nv, ctx: _f(getattr(ln, "phat_5s_dong_phuc", 0))),
        Cot("Khoản trừ khác", "tru", lambda ln, nv, ctx: _khoan_tru(ln)),
    ]
    tam = cot + [Cot("Thực nhận", "cuoi", lambda ln, nv, ctx: _f(getattr(ln, "net_pay", 0)))]
    if any(_khong_tru(tam, ln, _nv(nhan_vien, ln), bh_tach(ln), ty_le_thu_viec) for ln in lines):
        cot.append(Cot(KHONG_TRU, "tru", None))
    cot.append(Cot("Thực nhận", "cuoi", lambda ln, nv, ctx: _f(getattr(ln, "net_pay", 0))))
    # Sau Thực nhận là cột phụ, không cộng trừ gì — bảng công ty cũng để số ca đêm ở khu này.
    cot.append(Cot("Ngày ca đêm", "phu", lambda ln, nv, ctx: int(getattr(ln, "night_days", 0) or 0), CONG))
    if any(_f(getattr(ln, "no_ung_chuyen_ky_sau", 0)) for ln in lines):
        cot.append(Cot("Nợ ứng chuyển kỳ sau", "phu",
                       lambda ln, nv, ctx: _f(getattr(ln, "no_ung_chuyen_ky_sau", 0))))
    cot.append(Cot("Ghi chú", "phu", lambda ln, nv, ctx: getattr(ln, "note", None) or "", None))
    return cot


def _nv(nhan_vien: dict, ln) -> dict:
    return nhan_vien.get(getattr(ln, "employee_id", None), {}) or {}


#: Lệch làm tròn tối đa giữa tổng các cột thu và `gross` — mỗi khoản trên dòng lương làm tròn riêng tới đồng, còn
#: `gross` làm tròn MỘT lần trên tổng chưa tròn (~8 khoản tính ra, mỗi khoản ±0,5đ).
LECH_LAM_TRON = 10


def _gia_tri_tho(cot: list[Cot], ln, nv: dict, bh3, ty_le_thu_viec: float) -> list:
    ctx = {"bh3": tuple(bh3), "khoan": _phu_cap_theo_khoan(ln, nv), "ty_le": ty_le_thu_viec}
    gia_tri = [c.lay(ln, nv, ctx) if c.lay is not None else None for c in cot]
    # Dồn lệch LÀM TRÒN vào "Lương thời gian" (cột tính gộp, như file công ty) để Tổng lương đúng bằng tổng thu
    # engine đã dùng = gross + phạt đã ghi. Lệch lớn hơn nghĩa là phạt bị trần cắt ⇒ để nguyên, cột "Không trừ
    # được kỳ này" nói phần đó. Không dồn thì 1đ làm tròn cũng đẻ ra cột đó.
    thu = sum(_f(gia_tri[i]) for i, c in enumerate(cot) if c.nhom == "thu")
    dich = (_f(getattr(ln, "gross", 0)) + _f(getattr(ln, "vi_pham", 0)) + _f(getattr(ln, "di_tre", 0))
            + _f(getattr(ln, "phat_bien_ban", 0)) + _f(getattr(ln, "dt_vuot_troi", 0))
            + _f(getattr(ln, "phat_5s_dong_phuc", 0)))
    lech = round(dich - thu)
    if lech and abs(lech) <= LECH_LAM_TRON:
        i = [c.ten for c in cot].index("Lương thời gian")
        gia_tri[i] = round(_f(gia_tri[i]) + lech)
    return gia_tri


def _khong_tru(cot: list[Cot], ln, nv: dict, bh3, ty_le_thu_viec: float) -> float:
    """Phần khấu trừ GHI NHẬN mà engine không trừ kỳ này — số âm (0 khi không có)."""
    gia_tri = _gia_tri_tho(cot, ln, nv, bh3, ty_le_thu_viec)
    tong = sum(_f(gia_tri[i]) for i, c in enumerate(cot) if c.nhom == "thu")
    tru = sum(_f(gia_tri[i]) for i, c in enumerate(cot) if c.nhom == "tru" and c.lay is not None)
    return round(tong - tru - _f(getattr(ln, "net_pay", 0)))


def dong_so(cot: list[Cot], ln, nv: dict, bh3, *, ty_le_thu_viec: float = 1.0) -> list:
    """Một dòng lương → list giá trị theo đúng `cot`. Tách ra để test đối chiếu số học."""
    gia_tri = _gia_tri_tho(cot, ln, nv, bh3, ty_le_thu_viec)
    ten = [c.ten for c in cot]
    gia_tri[ten.index("Tổng lương")] = round(
        sum(_f(gia_tri[i]) for i, c in enumerate(cot) if c.nhom == "thu"))
    if KHONG_TRU in ten:
        gia_tri[ten.index(KHONG_TRU)] = _khong_tru(cot, ln, nv, bh3, ty_le_thu_viec)
    return gia_tri


# --- sheet ---------------------------------------------------------------------------------------


def _to_dam(ws, hang: int, cot: int, gia_tri, fmt: str | None, nen: str | None = None,
            dam: bool = False):
    o = ws.cell(row=hang, column=cot, value=gia_tri)
    if fmt:
        o.number_format = fmt
    if nen:
        o.fill = PatternFill("solid", fgColor=nen)
    if dam:
        o.font = Font(bold=True)
    o.border = VIEN
    return o


NOI_BAT = ("Lương thời gian", "Tổng lương", "Thực nhận")


def _sheet_bang_luong(wb, lines, nam, thang, nhan_vien, bh_tach, ty_le_thu_viec):
    ws = wb.active
    ws.title = f"Bang luong {thang:02d}-{nam}"
    cot = cot_bang_luong(lines, nhan_vien, bh_tach=bh_tach, ty_le_thu_viec=ty_le_thu_viec)
    so_cot = len(cot)
    rong_tieu_de = min(so_cot, 16)

    o = ws.cell(row=1, column=1, value="BẢNG THANH TOÁN LƯƠNG NHÂN VIÊN")
    o.font = Font(bold=True, size=14)
    o.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=rong_tieu_de)
    cong_chuan = _f(getattr(lines[0], "standard_cong", 0)) if lines else 0
    o = ws.cell(row=2, column=1, value=(
        f"Tháng {thang:02d}/{nam}" + (f" · Công chuẩn {cong_chuan:g}" if cong_chuan else "")))
    o.font = Font(italic=True)
    o.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=rong_tieu_de)

    # Hàng 3: dải tên khối, gộp ô theo từng khối liền nhau.
    dau = 1
    for i in range(1, so_cot + 1):
        het_khoi = i == so_cot or cot[i].nhom != cot[i - 1].nhom
        if not het_khoi:
            continue
        nhom = cot[i - 1].nhom
        o = ws.cell(row=3, column=dau, value=NHAN_NHOM[nhom] or None)
        o.font = Font(bold=True, size=9)
        o.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(dau, i + 1):
            ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor=NEN_NHOM[nhom])
            ws.cell(row=3, column=c).border = VIEN
        if i > dau:
            ws.merge_cells(start_row=3, start_column=dau, end_row=3, end_column=i)
        dau = i + 1

    for i, c in enumerate(cot, start=1):
        o = ws.cell(row=4, column=i, value=c.ten)
        o.font = Font(bold=True, size=10)
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.fill = PatternFill("solid", fgColor=NEN_NHOM[c.nhom])
        o.border = VIEN
    ws.row_dimensions[4].height = 42

    hang = 5
    co_thu_viec = False
    for stt, ln in enumerate(lines, start=1):
        gia_tri = dong_so(cot, ln, _nv(nhan_vien, ln), bh_tach(ln), ty_le_thu_viec=ty_le_thu_viec)
        gia_tri[0] = stt
        thu_viec = bool(getattr(ln, "is_probation", False))
        co_thu_viec = co_thu_viec or thu_viec
        for i, c in enumerate(cot, start=1):
            noi_bat = c.ten in NOI_BAT
            nen = NEN_NHOM[c.nhom] if noi_bat else (NEN_THU_VIEC if thu_viec else None)
            _to_dam(ws, hang, i, gia_tri[i - 1], c.fmt, nen=nen, dam=noi_bat)
        hang += 1

    if lines:
        _to_dam(ws, hang, 3, "Tổng cộng", None, dam=True)
        for i, c in enumerate(cot, start=1):
            if c.fmt not in (TIEN, CONG):
                continue
            chu = get_column_letter(i)
            _to_dam(ws, hang, i, f"=SUM({chu}5:{chu}{hang - 1})", c.fmt, dam=True)
        hang += 2
        ghi_chu = [
            "Mức lương tháng (Lương BHXH, cơ bản, trách nhiệm, phụ cấp) chỉ để tham chiếu — tiền thật nằm "
            "trong Lương thời gian.",
            "Lương thời gian = lương theo công (tháng lấy bù lỗ: trọn bù lỗ, đã gồm tiền khoán) + công lễ "
            "nghỉ + phụ cấp theo công + phần thêm ngày Chủ nhật / lễ ≈ ((cơ bản + trách nhiệm) × tỉ lệ + "
            "phụ cấp) ÷ công chuẩn × Tổng NC.",
            "Ngoài giờ/Tăng ca chỉ là tiền giờ tăng ca. Tổng lương = các khoản thu nhập + chuyên cần + lương "
            "thời gian. Thực nhận = Tổng lương − các khoản trừ.",
        ]
        if co_thu_viec:
            ghi_chu.append(f"Nền vàng: đang thử việc — lương cơ bản + trách nhiệm tính "
                           f"{float(ty_le_thu_viec) * 100:g}%, không đóng bảo hiểm.")
        if any(c.ten == KHONG_TRU for c in cot):
            ghi_chu.append(f"{KHONG_TRU}: số âm là phần phạt vượt trần khấu trừ (Điều 102) hoặc phần "
                           "khấu trừ lớn hơn lương còn lại — không trừ vào kỳ này.")
        for dong in ghi_chu:
            o = ws.cell(row=hang, column=3, value=dong)
            o.font = Font(italic=True, size=9)
            hang += 1

    rong = {"Số TT": 6, "MNV": 10, "Họ và tên": 24, "Chức vụ": 14, "Ghi chú": 24}
    for i, c in enumerate(cot, start=1):
        ws.column_dimensions[get_column_letter(i)].width = rong.get(
            c.ten, 9 if c.nhom == "cong" else 13)
    ws.freeze_panes = "E5"
    return ws


def _sheet_ky_nhan(wb, lines, nam, thang, nhan_vien):
    ws = wb.create_sheet("Ky nhan")
    o = ws.cell(row=1, column=1, value=f"BẢNG KÝ NHẬN LƯƠNG THÁNG {thang:02d}/{nam}")
    o.font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    tieu_de = ["STT", "Mã NV", "Họ và tên", "Chức vụ", "Bộ phận", "Thực nhận", "Thu lại", "Thực chi", "Ký nhận"]
    for i, ten in enumerate(tieu_de, start=1):
        o = ws.cell(row=3, column=i, value=ten)
        o.font = Font(bold=True)
        o.alignment = Alignment(horizontal="center")
        o.fill = PatternFill("solid", fgColor=NEN_NHOM["tt"])
        o.border = VIEN
    hang = 4
    for stt, ln in enumerate(lines, start=1):
        nv = _nv(nhan_vien, ln)
        net = _f(ln.net_pay)
        # "Thu lại" để TRỐNG cho kế toán ghi tay (thu lại tiền đã ứng thừa ngoài hệ thống);
        # "Thực chi" là công thức để sửa ô Thu lại là số tự chạy.
        for i, v in enumerate([stt, getattr(ln, "employee_code", "") or "",
                               getattr(ln, "employee_name", "") or "", nv.get("chuc_vu") or "",
                               getattr(ln, "department_name", "") or "", net, None,
                               # `N()` biến ô Thu lại bỏ trống thành 0 ⇒ Thực chi tự chạy khi kế
                               # toán ghi tay số thu lại.
                               f"=F{hang}-N(G{hang})", None], start=1):
            _to_dam(ws, hang, i, v, TIEN if i in (6, 7, 8) else None)
        hang += 1
    for chu, w in {"A": 5, "B": 10, "C": 24, "D": 16, "E": 16, "F": 15, "G": 13, "H": 15, "I": 18}.items():
        ws.column_dimensions[chu].width = w
    ws.freeze_panes = "A4"
    return ws


def _sheet_tam_ung(wb, lines, nam, thang, nhan_vien, tam_ung):
    """Bảng tạm ứng xếp theo NGÀY CHI, đúng cách kế toán đang làm (mỗi ngày một cột)."""
    ws = wb.create_sheet("Tam ung")
    ngay = sorted({t["ngay"] for t in tam_ung if t.get("ngay")})
    o = ws.cell(row=1, column=1, value=f"BẢNG TẠM ỨNG LƯƠNG THÁNG {thang:02d}/{nam}")
    o.font = Font(bold=True, size=13)
    tieu_de = ["STT", "Mã NV", "Họ và tên", "Chức vụ"] + [d.strftime("%d/%m") for d in ngay] + \
        ["Tổng ứng", "Lương đợt 1", "Nợ ứng kỳ trước", "Đã trừ kỳ này", "Còn nợ kỳ sau"]
    for i, ten in enumerate(tieu_de, start=1):
        o = ws.cell(row=3, column=i, value=ten)
        o.font = Font(bold=True)
        o.alignment = Alignment(horizontal="center")
        o.fill = PatternFill("solid", fgColor=NEN_NHOM["tru"])
        o.border = VIEN
    theo_nv: dict[int, dict] = {}
    for t in tam_ung:
        theo_nv.setdefault(t["employee_id"], {}).setdefault(t.get("ngay"), 0.0)
        theo_nv[t["employee_id"]][t.get("ngay")] += _f(t.get("so_tien"))
    hang = 4
    for stt, ln in enumerate(lines, start=1):
        eid = getattr(ln, "employee_id", None)
        nv = nhan_vien.get(eid, {}) or {}
        cua_nv = theo_nv.get(eid, {})
        no_sau = _f(getattr(ln, "no_ung_chuyen_ky_sau", 0))
        cot = [stt, getattr(ln, "employee_code", "") or "", getattr(ln, "employee_name", "") or "",
               nv.get("chuc_vu") or ""]
        cot += [cua_nv.get(d) or None for d in ngay]
        cot += [_f(ln.advance_total), _f(getattr(ln, "luong_dot_1_total", 0)),
                _f(getattr(ln, "no_ung_ky_truoc", 0)), _tam_ung_tru_ky_nay(ln), no_sau]
        for i, v in enumerate(cot, start=1):
            _to_dam(ws, hang, i, v, TIEN if i > 4 else None)
        hang += 1
    for i in range(1, len(tieu_de) + 1):
        ws.column_dimensions[get_column_letter(i)].width = {1: 5, 2: 10, 3: 24, 4: 16}.get(i, 14)
    ws.freeze_panes = "E4"
    return ws


def xuat_bang_luong(lines, *, nam: int, thang: int, nhan_vien: dict, bh_tach,
                    tam_ung: list[dict] | None = None, ty_le_thu_viec: float = 1.0) -> bytes:
    """Ba sheet: Bảng lương (khuôn `BL CT` của công ty) · Ký nhận · Tạm ứng.

    `nhan_vien` = {employee_id: {chuc_vu, ngay_vao_lam, nguoi_phu_thuoc, luong_vi_tri, luong_trach_nhiem,
    phu_cap_khac, khoan_ho_so}} — thứ dòng lương không giữ mà bảng của kế toán có (mức tháng lấy từ
    `PayrollService.muc_luong_thang_cho_file`). `bh_tach(line)` trả `(BHXH, BHYT, BHTN)` đã tách từ tổng
    đã đóng băng (router dùng lại `_insurance_lines`, phần dư dồn vào BHTN nên luôn cộng đúng tổng).
    `ty_le_thu_viec` = `payroll_params.probation_ratio` — chỉ để suy "Tổng NC" tháng lấy bù lỗ của người
    thử việc và ghi chú cuối bảng.
    """
    wb = Workbook()
    lines = list(lines)
    _sheet_bang_luong(wb, lines, nam, thang, nhan_vien, bh_tach, ty_le_thu_viec)
    _sheet_ky_nhan(wb, lines, nam, thang, nhan_vien)
    _sheet_tam_ung(wb, lines, nam, thang, nhan_vien, tam_ung or [])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
