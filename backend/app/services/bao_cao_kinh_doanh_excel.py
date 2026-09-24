"""Xuất BÁO CÁO KINH DOANH theo khách hàng ra .xlsx (24/09/2026).

    dòng 1-2  tiêu đề + kỳ
    dòng 4    tiêu đề cột (14 cột, A→N)
    rồi từng khách:
        [KHÁCH]  Mã - Tên (n đơn)                       … tổng tiền + cọc của khách   (tô nền)
        [ĐƠN]    Số đơn · Ngày chốt · Sale · PO          … tổng đơn + % cọc + cọc      (đậm)
        [DÒNG]            sản phẩm · SL · ĐVT · đơn giá · VAT · thành tiền
    chân: Tổng cộng

Một cột mang MỘT nghĩa cho cả ba loại dòng (Thành tiền của dòng sản phẩm = cột I; tổng chưa VAT
của đơn/khách cũng nằm cột I) — nên cộng dọc cột I theo dòng ĐƠN là ra đúng tổng của khách.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DINH_DANG_TIEN = "#,##0"

COT = (
    ("A", "Số đơn", 16), ("B", "Ngày chốt", 11), ("C", "Sale", 18),
    ("D", "Sản phẩm", 44), ("E", "SL", 9), ("F", "ĐVT", 8),
    ("G", "Đơn giá", 13), ("H", "VAT %", 7), ("I", "Thành tiền (chưa VAT)", 16),
    ("J", "Tổng có VAT", 16), ("K", "% cọc", 7), ("L", "Cọc phải thu", 15),
    ("M", "Cọc đã nhận", 15), ("N", "Cọc còn thiếu", 15),
)
_COT_TIEN = ("G", "I", "J", "L", "M", "N")


def _dmy(d: date | None) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year}" if d else ""


def xuat_xlsx(bc: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo kinh doanh"
    mong = Side(style="thin", color="FFBFC5CD")
    vien = Border(left=mong, right=mong, top=mong, bottom=mong)
    nen_khach = PatternFill("solid", fgColor="FFE8EDF4")
    nen_tieu_de = PatternFill("solid", fgColor="FF1B2A41")

    ws.merge_cells("A1:N1")
    ws["A1"] = "BÁO CÁO KINH DOANH THEO KHÁCH HÀNG"
    ws["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:N2")
    ws["A2"] = (f"Đơn đã chốt từ ngày {_dmy(bc['tu_ngay'])} đến ngày {_dmy(bc['den_ngay'])}"
                f" · {bc['tong']['so_khach']} khách · {bc['tong']['so_don']} đơn")
    ws["A2"].font = Font(name="Times New Roman", size=11, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    for cot, nhan, rong in COT:
        o = ws[f"{cot}4"]
        o.value = nhan
        o.font = Font(bold=True, color="FFFFFFFF", size=10)
        o.fill = nen_tieu_de
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.border = vien
        ws.column_dimensions[cot].width = rong
    ws.row_dimensions[4].height = 30

    def _ke(hang: int, *, dam=False, nen=None) -> None:
        for cot, _, _ in COT:
            o = ws[f"{cot}{hang}"]
            o.border = vien
            o.font = Font(size=10, bold=dam)
            if nen is not None:
                o.fill = nen
            if cot in _COT_TIEN:
                o.number_format = DINH_DANG_TIEN
                o.alignment = Alignment(horizontal="right", vertical="center")
            elif cot in ("E", "H", "K", "B", "F"):
                o.alignment = Alignment(horizontal="center", vertical="center")
            else:
                o.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _tien_tong(hang: int, m: dict) -> None:
        ws[f"I{hang}"] = m["tong"]
        ws[f"J{hang}"] = m["tong_vat"]
        ws[f"L{hang}"] = m["coc_phai_thu"]
        ws[f"M{hang}"] = m["coc_da_nhan"]
        ws[f"N{hang}"] = m["coc_con_thieu"]

    hang = 5
    for k in bc["khach"]:
        ws.merge_cells(f"A{hang}:H{hang}")
        ws[f"A{hang}"] = (" - ".join(x for x in (k["ma"], k["ten"]) if x)
                          + f"  ({k['so_don']} đơn)")
        _tien_tong(hang, k)
        _ke(hang, dam=True, nen=nen_khach)
        hang += 1
        for d in k["don"]:
            ws[f"A{hang}"] = d["order_no"]
            ws[f"B{hang}"] = _dmy(d["ngay_chot"])
            ws[f"C{hang}"] = d["sale"] or ""
            ws[f"D{hang}"] = f"PO khách: {d['po_khach']}" if d.get("po_khach") else ""
            _tien_tong(hang, d)
            ws[f"K{hang}"] = f"{d['coc_pct']:g}%" if d["coc_pct"] else ""
            _ke(hang, dam=True)
            hang += 1
            for ln in d["dong"]:
                ws[f"D{hang}"] = ln["ten"]
                ws[f"E{hang}"] = ln["so_luong"]
                ws[f"F{hang}"] = ln["dvt"] or ""
                ws[f"G{hang}"] = ln["don_gia"]
                ws[f"H{hang}"] = f"{ln['vat_pct']}%" if ln["vat_pct"] else ""
                ws[f"I{hang}"] = ln["thanh_tien"]
                _ke(hang)
                ws[f"E{hang}"].number_format = "#,##0"
                hang += 1

    ws.merge_cells(f"A{hang}:H{hang}")
    ws[f"A{hang}"] = (f"TỔNG CỘNG ({bc['tong']['so_khach']} khách · "
                      f"{bc['tong']['so_don']} đơn)")
    _tien_tong(hang, bc["tong"])
    _ke(hang, dam=True, nen=nen_khach)

    ws.freeze_panes = "A5"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ten_file(bc: dict, *, ma_khach: str | None = None) -> str:
    """`bao-cao-kinh-doanh-2026-09-01-den-2026-09-24.xlsx`; một khách thì chèn mã khách (ASCII —
    tên file nằm trong header `Content-Disposition`)."""
    import re
    import unicodedata

    ai = ""
    if ma_khach:
        tho = unicodedata.normalize("NFKD", ma_khach.replace("đ", "d").replace("Đ", "D"))
        tho = re.sub(r"[^A-Za-z0-9]+", "-", tho.encode("ascii", "ignore").decode()).strip("-")
        ai = f"{tho[:40]}-" if tho else ""
    return (f"bao-cao-kinh-doanh-{ai}{bc['tu_ngay']:%Y-%m-%d}"
            f"-den-{bc['den_ngay']:%Y-%m-%d}.xlsx")
