"""Xuất BẢNG CÔNG THÁNG ra .xlsx (thay bản .csv cũ — chủ chốt 09/09/2026).

Bản .csv trước đây có hai chỗ làm kế toán đọc không hiểu:

1. Ô ngày KHÔNG có công, KHÔNG có ca vẫn in chữ **"có"** — vì service dựng ô ngày cho cả ngày mới
   XẾP CA (chưa tới), ngày nghỉ luân phiên, ngày chấm thiếu lượt… nên nhánh cuối `else: "có"` nuốt
   hết mấy loại đó vào một chữ vô nghĩa. Giờ mỗi loại một ký hiệu, có bảng chú thích ngay đầu file;
   ngày không đi làm để **TRỐNG**.
2. TĂNG CA không hiện ở đâu cả — bảng công là chỗ kế toán soi giờ làm thêm, mà file xuất ra lại
   giấu mất. Giờ ngày có tăng ca hiện thêm ``+2h`` ngay trong ô, và có cột **Tăng ca (giờ)** ở cuối.

Khuôn:

    A1   BẢNG CÔNG THÁNG MM/YYYY
    A2   Phòng/tổ · công chuẩn tháng · ngày xuất
    A3   Chú thích ký hiệu
    A5   tiêu đề (thứ)  ┐ hai dòng; cột thông tin + cột tổng gộp ô
    A6   tiêu đề (ngày) ┘
    A7+  dữ liệu — mỗi NV một dòng
    cuối dòng TỔNG (số công · tăng ca · tổng giờ)
"""
from __future__ import annotations

import unicodedata
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

THU_VN = ["T2", "T3", "T4", "T5", "T6", "T7", "CN"]

COT_THONG_TIN = ["Mã", "Họ tên", "Phòng/Tổ", "Ca"]
COT_TONG = ["Số công", "Công CN/Lễ", "Tăng ca (giờ)", "Tổng giờ"]

CHU_THICH = (
    "Ký hiệu: 1 / 0.5 = số công · 8h = số giờ (ngày chưa gán ca) · +2h = giờ tăng ca của ngày đó · "
    "P = nghỉ phép có lương · KL = nghỉ không lương · L = nghỉ lễ (vẫn hưởng lương) · "
    "NL = nghỉ theo lịch phân ca · ? = có chấm nhưng thiếu lượt (ngày treo) · "
    "TC? = có phiếu tăng ca nhưng chưa có cặp bấm (đang tính 0 giờ) · ô trống = không đi làm. "
    "Nền hồng = ngày lễ · nền xám = thứ Bảy / Chủ nhật."
)

NEN_CUOI_TUAN = PatternFill("solid", fgColor="FFF1F3F5")
NEN_LE = PatternFill("solid", fgColor="FFFDE7E9")
NEN_TIEU_DE = PatternFill("solid", fgColor="FFE9ECEF")
VIEN = Border(*(Side(style="thin", color="FFCED4DA"),) * 4)


def _so(v: float) -> str:
    """1 chứ không phải 1.0; 0.5 vẫn là 0.5."""
    return f"{round(float(v), 2):g}"


def o_ngay(d: dict | None) -> str:
    """Một ô ngày → ký hiệu đọc được.

    TRẬT TỰ hỏi GIỮ ĐÚNG như lịch trên màn (`docONgay` bên FE): hỏi LƯỢT BẤM trước, vì ngày lễ /
    Chủ nhật ĐI LÀM vừa có giờ vào-ra vừa mang cờ ngày — hỏi cờ trước là nuốt mất công thật.
    """
    if not d:
        return ""
    if d.get("first_in") or d.get("last_out"):
        if d.get("cong") is not None:
            nhan = _so(d["cong"])
        elif d.get("hours") is not None:
            nhan = f"{_so(d['hours'])}h"
        else:
            nhan = "?"          # có lượt bấm mà không ra được công/giờ ⇒ ngày treo
    elif d.get("holiday"):
        nhan = "L"
    elif d.get("leave"):
        nhan = "P" if d.get("leave_paid") else "KL"
    elif d.get("planned_off"):
        nhan = "NL"
    else:
        nhan = ""               # đã xếp ca mà chưa tới / không đi làm ⇒ TRỐNG, đừng in "có"
    phut = int(d.get("ot_minutes") or 0)
    if phut > 0:
        nhan = f"{nhan} +{_so(phut / 60)}h".strip()
    elif d.get("ot_thieu_cap"):
        nhan = f"{nhan} TC?".strip()
    return nhan


def _khong_dau(s: str) -> str:
    """Bỏ dấu + thường hoá — gõ "quan" phải ra "Quân" (cùng luật với ô tìm trên màn)."""
    s = unicodedata.normalize("NFD", (s or "").strip().lower())
    return "".join(ch for ch in s if not unicodedata.combining(ch)).replace("đ", "d")


def loc_nhan_vien(rows: list[dict], tim: str | None) -> list[dict]:
    """Lọc theo TÊN hoặc MÃ nhân viên — để file xuất ra đúng thứ đang thấy trên màn.

    Cùng luật với ô tìm bên FE (bỏ dấu, không phân biệt hoa thường, khớp một đoạn)."""
    q = _khong_dau(tim or "")
    if not q:
        return rows
    return [r for r in rows
            if q in _khong_dau(r.get("employee_name") or "")
            or q in _khong_dau(r.get("employee_code") or "")]


def gio_tang_ca(row: dict) -> float:
    """Tổng giờ tăng ca cả tháng của một NV (2 chữ số)."""
    phut = sum(int((d or {}).get("ot_minutes") or 0) for d in (row.get("days") or {}).values())
    return round(phut / 60, 2)


def cong_cn_le(row: dict) -> float:
    """Công của ngày nghỉ tuần + ngày lễ + ngày công ty cho nghỉ (công THẬT, chưa nhân hệ số).

    Cùng con số với cột "CN/Lễ" trên màn (chủ 09/09/2026). Không cộng sau khi quy đổi: hệ số lễ khác
    hệ số Chủ nhật nên số quy đổi gộp không đối chiếu được với bảng chấm công."""
    return round(sum(float(row.get(k) or 0)
                     for k in ("restday_cong", "holiday_cong", "plain_cong")), 2)


def _cong_thang(row: dict) -> float:
    cong = row.get("total_cong")
    if cong is None:
        cong = row.get("total_days") or 0
    return round(float(cong), 2)


def xuat_bang_cong(rows: list[dict], *, nam: int, thang: int, so_ngay: int,
                   ten_bo_phan: str | None = None, cong_chuan: int | None = None,
                   ngay_le: dict[int, str] | None = None, tim: str | None = None) -> bytes:
    ngay_le = ngay_le or {}
    wb = Workbook()
    ws = wb.active
    ws.title = f"Bang cong {thang:02d}-{nam}"

    so_cot = len(COT_THONG_TIN) + so_ngay + len(COT_TONG)
    cot_ngay_dau = len(COT_THONG_TIN) + 1
    cot_tong_dau = cot_ngay_dau + so_ngay

    o = ws.cell(row=1, column=1, value=f"BẢNG CÔNG THÁNG {thang:02d}/{nam}")
    o.font = Font(bold=True, size=13)
    o.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=so_cot)

    phu = [f"Phòng/tổ: {ten_bo_phan or 'Tất cả'}"]
    if tim and tim.strip():
        # Nói rõ file đang lọc theo ô tìm — nếu không, người nhận file thấy thiếu người mà không
        # hiểu vì sao.
        phu.append(f"Lọc theo: {tim.strip()}")
    if cong_chuan is not None:
        phu.append(f"Công chuẩn tháng: {cong_chuan}")
    if ngay_le:
        # Gọi TÊN ngày lễ ngay đầu file — cột ngày chỉ tô nền được, mà kế toán cần biết vì sao
        # ngày đó ăn hệ số.
        phu.append("Ngày lễ: " + ", ".join(
            f"{d:02d}/{thang:02d} {ten}" for d, ten in sorted(ngay_le.items())))
    phu.append(f"Xuất ngày {date.today():%d/%m/%Y}")
    o = ws.cell(row=2, column=1, value=" · ".join(phu))
    o.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=so_cot)

    o = ws.cell(row=3, column=1, value=CHU_THICH)
    o.font = Font(italic=True, size=9)
    o.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=so_cot)
    ws.row_dimensions[3].height = 28

    for i, ten in enumerate(COT_THONG_TIN, start=1):
        ws.cell(row=5, column=i, value=ten)
        ws.merge_cells(start_row=5, start_column=i, end_row=6, end_column=i)
    for i, ten in enumerate(COT_TONG, start=cot_tong_dau):
        ws.cell(row=5, column=i, value=ten)
        ws.merge_cells(start_row=5, start_column=i, end_row=6, end_column=i)
    for d in range(1, so_ngay + 1):
        cot = cot_ngay_dau + d - 1
        ws.cell(row=5, column=cot, value=THU_VN[date(nam, thang, d).weekday()])
        ws.cell(row=6, column=cot, value=d)
        ws.column_dimensions[get_column_letter(cot)].width = 6.5
    for hang_td in (5, 6):
        for cot in range(1, so_cot + 1):
            o = ws.cell(row=hang_td, column=cot)
            o.font = Font(bold=True, size=10)
            o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            o.fill = NEN_TIEU_DE
            o.border = VIEN

    # Nền cột ngày lễ / cuối tuần: kế toán soi bảng là soi mấy cột này trước (ngày lễ, Chủ nhật đi
    # làm ăn hệ số). GIỮ NGUYÊN chữ thứ ở tiêu đề — thay bằng "LỄ" là mất luôn thông tin ngày đó
    # rơi vào thứ mấy; tên lễ đã nói ở dòng 2, màu nền đã nói ở dòng chú thích.
    nen_cot: dict[int, PatternFill] = {}
    for d in range(1, so_ngay + 1):
        cot = cot_ngay_dau + d - 1
        if d in ngay_le:
            nen_cot[cot] = NEN_LE
        elif date(nam, thang, d).weekday() >= 5:
            nen_cot[cot] = NEN_CUOI_TUAN
        if cot in nen_cot:
            ws.cell(row=5, column=cot).fill = nen_cot[cot]
            ws.cell(row=6, column=cot).fill = nen_cot[cot]

    hang = 7
    for r in rows:
        days = r.get("days") or {}
        thong_tin = [r.get("employee_code"), r.get("employee_name"),
                     r.get("department_name") or "", r.get("shift_name") or ""]
        for i, v in enumerate(thong_tin, start=1):
            o = ws.cell(row=hang, column=i, value=v)
            o.border = VIEN
            o.alignment = Alignment(vertical="center")
        for d in range(1, so_ngay + 1):
            cot = cot_ngay_dau + d - 1
            o = ws.cell(row=hang, column=cot, value=o_ngay(days.get(str(d))) or None)
            o.alignment = Alignment(horizontal="center", vertical="center")
            o.border = VIEN
            if cot in nen_cot:
                o.fill = nen_cot[cot]
        so_tong = [_cong_thang(r), cong_cn_le(r), gio_tang_ca(r),
                   round(float(r.get("total_hours") or 0), 2)]
        for i, v in enumerate(so_tong, start=cot_tong_dau):
            o = ws.cell(row=hang, column=i, value=v)
            o.font = Font(bold=True)
            o.alignment = Alignment(horizontal="center", vertical="center")
            o.border = VIEN
            o.number_format = "0.##"
        hang += 1

    o = ws.cell(row=hang, column=1, value="TỔNG")
    o.font = Font(bold=True)
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=cot_tong_dau - 1)
    tong = [
        sum(_cong_thang(r) for r in rows),
        sum(cong_cn_le(r) for r in rows),
        sum(gio_tang_ca(r) for r in rows),
        sum(float(r.get("total_hours") or 0) for r in rows),
    ]
    for i, v in enumerate(tong, start=cot_tong_dau):
        o = ws.cell(row=hang, column=i, value=round(v, 2))
        o.font = Font(bold=True)
        o.alignment = Alignment(horizontal="center")
        o.number_format = "0.##"

    for chu, rong in {"A": 12, "B": 26, "C": 18, "D": 14}.items():
        ws.column_dimensions[chu].width = rong
    for i in range(cot_tong_dau, so_cot + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = f"{get_column_letter(cot_ngay_dau)}7"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
