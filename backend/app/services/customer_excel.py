"""Nhập / xuất Excel — KHÁCH HÀNG.

- **Bản 1 (11/09/2026)** — tải mẫu rỗng → điền khách mới → xem trước → ghi.
- **Bản 2 (17/09/2026)** — Xuất Excel → sửa → nhập LẠI chính file đó: dòng có `Mã KH` là SỬA khách
  đó, dòng Mã trống là THÊM MỚI. Thiết kế + các chỗ đã chốt: `docs/spec-nhap-excel-khach-hang.md` §9.

Vì sao KHÔNG dùng `services/catalog_excel.py` (chủ chốt 11/09/2026: *"không muốn động sâu, các màn
khác đang ổn"*): cơ chế chung **bắt buộc có cột Mã** và **từ chối dòng Mã trống**
(`catalog_excel.py` — `'Sheet chính thiếu cột "Mã"'` / `"Thiếu mã."`). Mà mã khách hàng là mã hệ tự
cấp (`KH001`…), người dùng KHÔNG gõ — model ghi thẳng *"never entered by the user"*. Dạy cơ chế
chung hiểu "mã trống = dòng mới" thì sửa được cho cả ba màn đang vướng (Công việc khoán · Kho hàng ·
Khuôn bế), nhưng đó là code dùng chung của 15 màn ⇒ để lại, ghi trong
`docs/spec-nhap-excel-khach-hang.md`. Bản 2 dùng lại LUẬT của cơ chế chung (upsert theo mã, ô trống
xoá, thiếu cột giữ, dòng không đổi không ghi) chứ không dùng lại code.

Dùng lại mà KHÔNG sửa gì: ba hằng số `SHEET_META` / `PHIEN_BAN` / `ExcelSaiMan` của cơ chế chung —
nhờ đó file của màn này mang đúng dấu `_meta` như mọi màn khác, và nhập nhầm file của màn khác vẫn
bị chặn bằng đúng một cơ chế. Import hằng số là ĐỌC, không phải sửa.

LUẬT XƯƠNG SỐNG:

1. **`Mã KH` là con trỏ, không phải chỗ đặt mã.** Có mã ⇒ sửa đúng khách đó (phải trong phạm vi của
   người nhập). Trống ⇒ khách mới, mã do hệ cấp. Mã không tìm thấy thì báo lỗi, KHÔNG tạo khách
   mang mã đó.
2. **Một giao dịch cho cả file** — còn một dòng lỗi thì KHÔNG ghi gì cả. `ghi=False` (xem trước)
   chạy y hệt lượt thật rồi `rollback`, nên con số xem trước là con số THẬT.
3. **Dòng không đổi gì thì không đụng tới** — không ghi, không đẻ nhật ký. Xuất ra nhập lại y nguyên
   phải ra "0 thêm · 0 sửa · N không đổi".
4. **Trùng MST / tên / email là CẢNH BÁO, không phải lỗi.** Model nói rõ `tax_code` được index
   nhưng KHÔNG unique, trùng là cảnh báo mềm (§34, §41) — form nhập tay cũng chỉ cảnh báo. Chặn ở
   Excel là chặt hơn cả form, tức là sai.
"""
from __future__ import annotations

from dataclasses import dataclass, field as dc_field
from io import BytesIO
from typing import Any, NamedTuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models.customer import KIND_CA_NHAN, KIND_CONG_TY, Customer
from ..models.role import SCOPE_DEPARTMENT, SCOPE_OWN
from ..models.user import User
from .catalog_excel import PHIEN_BAN, SHEET_META, ExcelSaiMan
from .customer_service import (
    CustomerError,
    CustomerService,
    nguoi_du_tu_cach_nhan_khach,
)

#: Dấu nhận diện màn, ghi vào sheet `_meta` ẩn. Nhập file của màn khác ⇒ chặn ngay ở cửa.
LOAI = "khach_hang"
TIEU_DE = "Khách hàng"
#: Sheet ẨN chứa danh sách Sale — nguồn cho ô chọn xổ xuống ở cột "Sale phụ trách".
SHEET_SALE = "_sale"
#: Cột mã — CHỈ có ở file xuất. Mẫu rỗng cố ý không có (mọi dòng của mẫu là khách mới).
NHAN_MA = "Mã KH"
KHOA_MA = "code"

# Nhãn tiếng Việt của LOẠI KHÁCH. Ô Excel là chữ người đọc, không phải mã máy — người khai gõ
# "Công ty", không gõ "cong_ty".
NHAN_LOAI = {"công ty": KIND_CONG_TY, "cá nhân": KIND_CA_NHAN}
NHAN_LOAI_NGUOC = {KIND_CONG_TY: "Công ty", KIND_CA_NHAN: "Cá nhân"}

# Tên tiêu chí trùng, để câu cảnh báo đọc được thành lời.
NHAN_TRUNG = {"tax_code": "MST", "name": "tên", "email": "email"}

#: Cột phải để định dạng CHỮ: gõ `0901234567` vào ô kiểu Số là Excel nuốt số 0 đầu. MST còn có rào
#: 10/13 số nên sẽ báo lỗi, nhưng Điện thoại không có rào ⇒ mất số 0 âm thầm.
KHOA_CHU = ("tax_code", "phone")


@dataclass(frozen=True)
class Cot:
    """Một cột trên sheet chính.

    `tai_chinh=True` ⇒ chỉ xuất ra mẫu cho người có quyền `set_credit_terms`, và khi nhập thì
    người không có quyền sẽ bị BỎ QUA cột đó (kèm lời nói rõ), chứ không ghi lén và cũng không
    chặn cả file — xem `nhap()`.
    """

    nhan: str
    khoa: str
    kieu: str            # chu | so_nguyen | phan_tram | loai_khach | sale
    bat_buoc: bool = False
    rong: int = 18
    tai_chinh: bool = False


COT: tuple[Cot, ...] = (
    Cot("Tên khách hàng", "name", "chu", bat_buoc=True, rong=32),
    Cot("Loại khách", "customer_kind", "loai_khach", bat_buoc=True, rong=14),
    Cot("MST", "tax_code", "chu", rong=16),
    Cot("Điện thoại", "phone", "chu", rong=16),
    Cot("Email", "email", "chu", rong=26),
    Cot("Địa chỉ", "address", "chu", rong=40),
    Cot("Người liên hệ", "contact_name", "chu", rong=22),
    Cot("Sale phụ trách", "sale_user_id", "sale", rong=22),
    Cot("Hạn mức công nợ (đ)", "credit_limit", "so_nguyen", rong=20, tai_chinh=True),
    Cot("Số ngày nợ tối đa", "payment_term_days", "so_nguyen", rong=18, tai_chinh=True),
    Cot("Chiết khấu tối thiểu (%)", "discount_min_pct", "phan_tram", rong=20, tai_chinh=True),
    Cot("Chiết khấu tối đa (%)", "discount_max_pct", "phan_tram", rong=20, tai_chinh=True),
    Cot("Markup tối thiểu (%)", "markup_min_pct", "phan_tram", rong=20, tai_chinh=True),
    Cot("Markup tối đa (%)", "markup_max_pct", "phan_tram", rong=20, tai_chinh=True),
)

KHOA_TAI_CHINH = tuple(c.khoa for c in COT if c.tai_chinh)
KHOA_DINH_DANH = tuple(c.khoa for c in COT if not c.tai_chinh)
NHAN_THEO_KHOA = {c.khoa: c.nhan for c in COT}


# ======================= kết quả =======================


@dataclass
class LoiDong:
    dong: int
    cot: str
    ly_do: str


@dataclass
class CanhBaoDong:
    dong: int
    ly_do: str


@dataclass
class ThayDoi:
    """Một ô sẽ đổi trên một khách đã có — để xem trước đọc được "cũ → mới" trước khi bấm Ghi."""

    dong: int
    ma: str
    ten: str
    cot: str
    cu: str
    moi: str


@dataclass
class KetQua:
    tong_dong: int = 0
    tao_moi: int = 0
    cap_nhat: int = 0
    khong_doi: int = 0
    da_ghi: bool = False
    #: Có ô tài chính người nhập ĐÃ ĐIỀN / ĐÃ SỬA nhưng không có quyền ⇒ đã bỏ qua đúng mấy cột đó.
    bo_qua_tai_chinh: bool = False
    loi: list[LoiDong] = dc_field(default_factory=list)
    canh_bao: list[CanhBaoDong] = dc_field(default_factory=list)
    thay_doi: list[ThayDoi] = dc_field(default_factory=list)

    @property
    def hop_le(self) -> bool:
        return not self.loi


# ======================= đọc ô =======================


def _chuoi(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _so_nguyen(v: Any, nhan: str) -> tuple[int | None, str | None]:
    """(giá trị, lỗi). Ô trống ⇒ (None, None) — "chưa đặt", không phải 0."""
    s = _chuoi(v)
    if s is None:
        return None, None
    # Excel trả số về dạng float (2000000.0) — ép về int khi phần lẻ bằng 0.
    try:
        f = float(str(s).replace(",", "").replace(" ", ""))
    except ValueError:
        return None, f"{nhan}: phải là số."
    if f != int(f):
        return None, f"{nhan}: phải là số nguyên."
    if f < 0:
        return None, f"{nhan}: không được âm."
    return int(f), None


def _phan_tram(v: Any, nhan: str) -> tuple[float | None, str | None]:
    s = _chuoi(v)
    if s is None:
        return None, None
    try:
        f = float(str(s).replace("%", "").replace(",", ".").strip())
    except ValueError:
        return None, f"{nhan}: phải là số phần trăm."
    if not 0 <= f <= 100:
        return None, f"{nhan}: phải trong khoảng 0–100."
    return f, None


DOC_TAI_CHINH = {
    "credit_limit": ("Hạn mức công nợ", _so_nguyen),
    "payment_term_days": ("Số ngày nợ tối đa", _so_nguyen),
    "discount_min_pct": ("Chiết khấu tối thiểu", _phan_tram),
    "discount_max_pct": ("Chiết khấu tối đa", _phan_tram),
    "markup_min_pct": ("Markup tối thiểu", _phan_tram),
    "markup_max_pct": ("Markup tối đa", _phan_tram),
}


class _Nguoi(NamedTuple):
    username: str | None
    ten: str | None
    department_id: int | None


def _nguoi(db: Session) -> dict[int, _Nguoi]:
    return {
        int(i): _Nguoi(u, t, d)
        for i, u, t, d in db.execute(
            select(User.id, User.username, User.name, User.department_id)
        ).all()
    }


def _ten_hien(nguoi: dict[int, _Nguoi], uid: int | None) -> str:
    n = nguoi.get(uid) if uid else None
    if n is None:
        return ""
    return (n.ten or n.username or "").strip()


def _map_sale(db: Session) -> tuple[dict[str, int], set[int], list[str]]:
    """({khoá tra: user_id}, {id đủ tư cách}, [tên hiển thị trong ô chọn]).

    Tra được bằng CẢ tên đăng nhập lẫn họ tên (hạ chữ, gom khoảng trắng) — người khai gõ cái nào
    cũng ra. Họ tên TRÙNG NHAU thì bỏ khoá tên đi, chỉ còn tra bằng tên đăng nhập: đoán bừa một
    trong hai là gán khách cho nhầm người, mà đây là cột quyết định ai được thấy khách đó.

    Bảng tra giữ MỌI người, nhưng ô chọn và cửa kiểm chỉ nhận người **đủ tư cách nhận khách**
    (`nguoi_du_tu_cach_nhan_khach` — cùng luật với ô "NV phụ trách" trên màn). Giữ cả người ngoài
    khối trong bảng tra là có chủ ý: gõ tên Thủ kho thì phải báo *"không thuộc khối Kinh doanh"*
    chứ không phải *"không tìm thấy người nào tên đó"* — hai câu dẫn người dùng đi hai hướng khác
    hẳn nhau.
    """
    rows = list(db.execute(select(User.id, User.username, User.name)).all())
    du_tu_cach = nguoi_du_tu_cach_nhan_khach(db)

    dem_ten: dict[str, int] = {}
    for _, _, ten in rows:
        k = _khoa(ten)
        if k:
            dem_ten[k] = dem_ten.get(k, 0) + 1

    tra: dict[str, int] = {}
    hien: list[str] = []
    for uid, username, ten in rows:
        ku = _khoa(username)
        if ku:
            tra[ku] = int(uid)
        kt = _khoa(ten)
        if kt and dem_ten.get(kt, 0) == 1:
            tra[kt] = int(uid)
        if int(uid) in du_tu_cach:
            hien.append((ten or username or "").strip() or str(username))
    return tra, du_tu_cach, sorted(set(h for h in hien if h))


def _khoa(s: str | None) -> str:
    return " ".join((s or "").strip().lower().split())


# ======================= dựng file =======================


def _ghi_meta(ws_meta) -> None:
    ws_meta.title = SHEET_META
    ws_meta.append(["khoa", "gia_tri"])
    ws_meta.append(["loai", LOAI])
    ws_meta.append(["phien_ban", PHIEN_BAN])
    ws_meta.sheet_state = "hidden"


def _tieu_de(ws, nhan: list[str], rong: list[int]) -> None:
    ws.append(nhan)
    dam = Font(bold=True, color="FFFFFF")
    nen = PatternFill("solid", fgColor="1F3864")
    for i, r in enumerate(rong, start=1):
        o = ws.cell(row=1, column=i)
        o.font, o.fill = dam, nen
        ws.column_dimensions[get_column_letter(i)].width = r
    ws.freeze_panes = "A2"      # tiêu đề dính khi cuộn — file này người ta gõ hàng trăm dòng


def _o_chon_va_dinh_dang(db: Session, wb, ws, khoa_cot: list[str], den_dong: int) -> None:
    """Ô chọn xổ xuống (Loại khách, Sale) + định dạng CHỮ cho MST/Điện thoại — dùng chung cho mẫu
    rỗng lẫn file xuất, để hai file cư xử y hệt nhau trong Excel.

    Định dạng chữ đặt ở CẢ CỘT (ô người dùng gõ thêm sau này ăn theo) — không tạo sẵn ô trống, vì
    ô trống có định dạng làm `max_row` phình ra và dòng thêm bằng `append` rơi xuống tận dưới."""
    for khoa in KHOA_CHU:
        if khoa in khoa_cot:
            ws.column_dimensions[get_column_letter(khoa_cot.index(khoa) + 1)].number_format = "@"

    _o_chon(ws, khoa_cot, "customer_kind", '"Công ty,Cá nhân"', den_dong)
    _, _, ten_sale = _map_sale(db)
    if ten_sale:
        ws_sale = wb.create_sheet(SHEET_SALE)
        for t in ten_sale:
            ws_sale.append([t])
        ws_sale.sheet_state = "hidden"
        _o_chon(ws, khoa_cot, "sale_user_id", f"={SHEET_SALE}!$A$1:$A${len(ten_sale)}", den_dong)


def _o_chon(ws, khoa_cot: list[str], khoa: str, cong_thuc: str, den_dong: int) -> None:
    """Gắn ô chọn cho một cột, phủ tới `den_dong` — quá số đó thì gõ tay, vẫn nhập được."""
    if khoa not in khoa_cot:
        return
    dv = DataValidation(type="list", formula1=cong_thuc, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    chu = get_column_letter(khoa_cot.index(khoa) + 1)
    dv.add(f"{chu}2:{chu}{den_dong}")


def tao_mau(db: Session, *, co_tai_chinh: bool) -> bytes:
    """Workbook RỖNG — chỉ dòng tiêu đề (chốt 11/09/2026), để THÊM MỚI hàng loạt.

    Không kèm khách đang có và không có cột Mã: mọi dòng của mẫu là khách mới. Sửa khách đã có thì
    dùng file Xuất Excel (có cột `Mã KH`) — xem `xuat`.

    Người không có quyền `set_credit_terms` thì 6 cột tài chính KHÔNG xuất ra — đưa ra một cột họ
    không được ghi chỉ tổ mời họ điền vào chỗ sẽ bị bỏ qua.
    """
    cot = [c for c in COT if co_tai_chinh or not c.tai_chinh]

    wb = Workbook()
    _ghi_meta(wb.active)
    ws = wb.create_sheet(TIEU_DE)
    _tieu_de(ws, [c.nhan for c in cot], [c.rong for c in cot])
    _o_chon_va_dinh_dang(db, wb, ws, [c.khoa for c in cot], 501)

    ra = BytesIO()
    wb.save(ra)
    return ra.getvalue()


#: Cột của file XUẤT = `Mã KH` + đủ cột của mẫu, cùng nhãn. Sửa trong Excel rồi nhập lại được
#: (bản 2, 17/09/2026): dòng có mã là sửa, dòng thêm vào để trống mã là khách mới.
COT_XUAT = (NHAN_MA, *(c.nhan for c in COT))


def xuat(db: Session, khach: list) -> bytes:
    """Xuất danh bạ ĐANG THẤY ra .xlsx — phạm vi do người gọi lọc sẵn (router đã soi scope).

    Mang dấu `_meta` như mẫu nên nhập ngược lại được. Sheet chính đứng ĐẦU (là sheet mở ra khi bấm
    vào file); `_meta` và `_sale` là sheet ẩn phía sau.
    """
    nguoi = _nguoi(db)

    wb = Workbook()
    ws = wb.active
    ws.title = TIEU_DE
    _tieu_de(ws, list(COT_XUAT), [14, *(c.rong for c in COT)])

    for k in khach:
        ws.append([
            k.code,
            k.name,
            NHAN_LOAI_NGUOC.get(k.customer_kind, ""),
            k.tax_code or "",
            k.phone or "",
            k.email or "",
            k.address or "",
            k.contact_name or "",
            _ten_hien(nguoi, k.sale_user_id),
            k.credit_limit,
            k.payment_term_days,
            k.discount_min_pct,
            k.discount_max_pct,
            k.markup_min_pct,
            k.markup_max_pct,
        ])

    khoa_cot = [KHOA_MA, *(c.khoa for c in COT)]
    # Ô đã ghi mang style riêng nên không ăn định dạng cột — đặt chữ cho từng ô đã có dữ liệu.
    for khoa in KHOA_CHU:
        i = khoa_cot.index(khoa) + 1
        for r in range(2, len(khach) + 2):
            ws.cell(row=r, column=i).number_format = "@"
    _o_chon_va_dinh_dang(db, wb, ws, khoa_cot, len(khach) + 501)
    _ghi_meta(wb.create_sheet())

    ra = BytesIO()
    wb.save(ra)
    return ra.getvalue()


# ======================= đọc file =======================


def _mo_workbook(du_lieu: bytes):
    try:
        return load_workbook(BytesIO(du_lieu), data_only=True)
    except Exception:  # noqa: BLE001 — mọi kiểu hỏng của file đều cùng một câu trả lời
        raise ExcelSaiMan("Không đọc được file — phải là .xlsx đúng mẫu tải từ hệ thống.") from None


def _kiem_meta(wb) -> None:
    if SHEET_META not in wb.sheetnames:
        # File XUẤT đời trước 17/09/2026 không mang `_meta`. Nhận theo chế độ tương thích: sheet
        # "Khách hàng" + ô A1 là "Mã KH" là đủ đặc trưng, không nhận nhầm file màn khác.
        if TIEU_DE in wb.sheetnames and _chuoi(wb[TIEU_DE].cell(row=1, column=1).value) == NHAN_MA:
            return
        raise ExcelSaiMan(
            'File không phải file của màn Khách hàng (thiếu dấu "_meta"). Dùng file tải từ nút '
            '"Tải file mẫu" hoặc "Xuất Excel" của màn này.'
        )
    gia: dict[str, Any] = {}
    for hang in wb[SHEET_META].iter_rows(min_row=2, values_only=True):
        if hang and hang[0]:
            gia[str(hang[0]).strip()] = hang[1]
    if str(gia.get("loai") or "").strip() != LOAI:
        raise ExcelSaiMan(
            f'File này là mẫu của màn khác ("{gia.get("loai")}"), không phải Khách hàng.'
        )


# ======================= nhập =======================


@dataclass
class _Buoc:
    """Kế hoạch cho MỘT dòng hợp lệ. `kh is None` ⇒ thêm mới."""

    dong: int
    kh: Customer | None
    thong: dict
    #: Rỗng ⇒ không đụng chính sách tài chính. Dòng sửa: đủ 6 khoá (cột vắng lấy giá trị cũ).
    tai_chinh: dict
    bo_sale: bool = False
    doi: set[str] = dc_field(default_factory=set)


def nhap(db: Session, svc: CustomerService, du_lieu: bytes, *, actor, scope: str,
         co_tai_chinh: bool, co_quyen_tao: bool, co_quyen_sua: bool,
         co_quyen_dieu_chuyen: bool, ghi: bool) -> KetQua:
    """Đọc file → dựng kế hoạch → chạy trong MỘT giao dịch.

    `ghi=False` là XEM TRƯỚC: chạy y hệt lượt thật rồi `rollback`. Cố ý không làm một bản kiểm
    "sơ bộ" nhẹ hơn — nếu xem trước dễ dãi hơn lúc ghi thì người dùng bấm Xác nhận xong mới ăn lỗi,
    đúng thứ mà cái nút xem trước sinh ra để tránh.

    Quyền xét theo TỪNG DÒNG (chốt 17/09/2026): thêm cần `create`, sửa cần `update`, đổi/gỡ Sale cần
    thêm `reassign`. Không bắt cả hai như danh mục — người chỉ có `create` vẫn đang nhập khách mới được.
    """
    wb = _mo_workbook(du_lieu)
    _kiem_meta(wb)
    if TIEU_DE not in wb.sheetnames:
        raise ExcelSaiMan(f'File thiếu sheet "{TIEU_DE}".')

    hang = list(wb[TIEU_DE].iter_rows(values_only=True))
    if not hang:
        raise ExcelSaiMan(f'Sheet "{TIEU_DE}" trống — không có cả dòng tiêu đề.')

    tieu_de = [(_chuoi(o) or "") for o in hang[0]]
    vi_tri = {c.khoa: tieu_de.index(c.nhan) for c in COT if c.nhan in tieu_de}
    thieu = [c.nhan for c in COT if c.bat_buoc and c.khoa not in vi_tri]
    if thieu:
        raise ExcelSaiMan("File thiếu cột bắt buộc: " + " · ".join(thieu) + ". Tải lại file mẫu.")
    co_cot_ma = NHAN_MA in tieu_de
    if co_cot_ma:
        vi_tri[KHOA_MA] = tieu_de.index(NHAN_MA)

    def _o(dong) -> dict:
        return {k: (dong[i] if i < len(dong) else None) for k, i in vi_tri.items()}

    kq = KetQua()
    bo_doc = _BoDoc(
        kq=kq, actor=actor, scope=scope, co_tai_chinh=co_tai_chinh,
        co_quyen_dieu_chuyen=co_quyen_dieu_chuyen, db=db,
    )

    # Chỉ khách TRONG PHẠM VI mới trỏ tới được. Ngoài phạm vi thì coi như không có — cùng một câu
    # với mã sai, để không dò ra được khách của người khác (như `PUT /{id}` trả 404).
    theo_ma: dict[str, Customer] = {}
    dong_cua_ma: dict[str, list[int]] = {}
    if co_cot_ma:
        theo_ma = {(c.code or "").upper(): c for c in svc.list_scoped_all(scope=scope, actor=actor)}
        for so_dong, dong in enumerate(hang[1:], start=2):
            ma = (_chuoi(_o(dong).get(KHOA_MA)) or "").upper()
            if ma:
                dong_cua_ma.setdefault(ma, []).append(so_dong)

    ke_hoach: list[_Buoc] = []
    for so_dong, dong in enumerate(hang[1:], start=2):
        o = _o(dong)
        # Dòng TRẮNG HOÀN TOÀN: bỏ qua, không tính là lỗi — ai cũng để thừa vài dòng cuối file.
        if all(_chuoi(v) is None for v in o.values()):
            continue
        kq.tong_dong += 1

        ma = (_chuoi(o.pop(KHOA_MA, None)) or "").upper()
        kh: Customer | None = None
        if ma:
            if len(dong_cua_ma[ma]) > 1:
                kq.loi.append(LoiDong(so_dong, NHAN_MA, (
                    f'Mã "{ma}" nằm ở nhiều dòng ('
                    + ", ".join(str(d) for d in dong_cua_ma[ma])
                    + ") — mỗi khách chỉ được một dòng.")))
                continue
            kh = theo_ma.get(ma)
            if kh is None:
                kq.loi.append(LoiDong(so_dong, NHAN_MA, (
                    f'Không có khách mã "{ma}", hoặc bạn không phụ trách khách này. '
                    "Khách mới thì để trống Mã KH.")))
                continue
        elif not co_quyen_tao:
            kq.loi.append(LoiDong(so_dong, NHAN_MA,
                                  "Mã KH trống là thêm khách mới — bạn không có quyền Thêm "
                                  "khách hàng."))
            continue

        buoc = bo_doc.doc_dong(o, so_dong, kh)
        if buoc is None:
            continue
        if kh is not None and not buoc.doi:
            kq.khong_doi += 1
            continue
        if kh is not None and not co_quyen_sua:
            # Xét SAU khi so cũ/mới: người chỉ có `create` xuất file ra, thêm vài dòng cuối rồi nhập
            # lại thì các dòng cũ không đổi phải đi qua êm, chỉ dòng thật sự sửa mới bị chặn.
            kq.loi.append(LoiDong(so_dong, NHAN_MA,
                                  f"Dòng này sửa khách {kh.code} — bạn không có quyền Sửa khách "
                                  "hàng."))
            continue
        ke_hoach.append(buoc)

    # Còn lỗi thì DỪNG TRƯỚC KHI chạm DB — không cần rollback cái chưa từng ghi.
    if kq.loi:
        kq.canh_bao.sort(key=lambda c: c.dong)
        return kq

    try:
        for b in ke_hoach:
            so_dong = b.dong
            if b.kh is None:
                kh, trung = svc.create_customer(actor=actor, commit=False, **b.thong)
                kq.tao_moi += 1
                if b.tai_chinh:
                    svc.update_financial(customer_id=kh.id, scope=scope, actor=actor,
                                         commit=False, **b.tai_chinh)
            else:
                trung = []
                if b.doi & set(KHOA_DINH_DANH):
                    _, trung = svc.update_customer(
                        customer_id=b.kh.id, scope=scope, actor=actor,
                        allow_reassign=co_quyen_dieu_chuyen, bo_sale=b.bo_sale,
                        commit=False, **b.thong,
                    )
                    # Chỉ báo trùng theo ô VỪA SỬA — nếu không, nhập lại file xuất sẽ rải cảnh báo
                    # cho mọi cặp trùng đã có từ lâu.
                    trung = [(tc, khac) for tc, khac in trung if tc in b.doi]
                if b.tai_chinh:
                    svc.update_financial(customer_id=b.kh.id, scope=scope, actor=actor,
                                         commit=False, **b.tai_chinh)
                kq.cap_nhat += 1
            for tieu_chi, khac in trung:
                kq.canh_bao.append(CanhBaoDong(
                    so_dong,
                    f"Trùng {NHAN_TRUNG.get(tieu_chi, tieu_chi)} với {khac.code} · {khac.name}",
                ))
        if ghi:
            db.commit()
            kq.da_ghi = True
        else:
            db.rollback()
    except CustomerError as e:
        # Lỗi nghiệp vụ lộ ra ở tầng service (vd MST sai dạng, rào min > max, ngoài phạm vi quyền)
        # — trả về dưới dạng lỗi DÒNG chứ không 500, và KHÔNG ghi gì cả.
        db.rollback()
        kq.loi.append(LoiDong(so_dong, "", str(e)))
        kq.tao_moi = 0
        kq.cap_nhat = 0
    except Exception:
        db.rollback()
        raise
    # Cảnh báo lúc dựng kế hoạch (gỡ Sale, trỏ nhầm) và lúc chạy (trùng) đến theo hai lượt — xếp
    # lại theo dòng cho người đọc dò từ trên xuống.
    kq.canh_bao.sort(key=lambda c: c.dong)
    return kq


class _BoDoc:
    """Đọc MỘT dòng thành `_Buoc`. Gom những thứ tra một lần cho cả file (bảng Sale, người dùng)
    để không hỏi DB mỗi dòng."""

    def __init__(self, *, kq: KetQua, actor, scope: str, co_tai_chinh: bool,
                 co_quyen_dieu_chuyen: bool, db: Session) -> None:
        self.kq = kq
        self.actor = actor
        self.scope = scope
        self.co_tai_chinh = co_tai_chinh
        self.co_quyen_dieu_chuyen = co_quyen_dieu_chuyen
        self.tra_sale, self.du_tu_cach, _ = _map_sale(db)
        self.nguoi = _nguoi(db)

    def _loi(self, dong: int, cot: str, ly_do: str) -> None:
        self.kq.loi.append(LoiDong(dong, cot, ly_do))

    # ---------------------------------------------------------------- dòng

    def doc_dong(self, o: dict, so_dong: int, kh: Customer | None) -> _Buoc | None:
        """Lỗi ⇒ None (lỗi đã ghi vào `kq`). Dòng sửa: cột VẮNG MẶT giữ giá trị cũ."""
        hong = len(self.kq.loi)

        def _gia(khoa: str) -> str | None:
            if khoa not in o and kh is not None:
                return getattr(kh, khoa)
            return _chuoi(o.get(khoa))

        ten = _gia("name")
        if not ten:
            self._loi(so_dong, "Tên khách hàng", "Chưa điền tên khách hàng.")

        loai = kh.customer_kind if kh is not None and "customer_kind" not in o else KIND_CONG_TY
        loai_raw = _chuoi(o.get("customer_kind"))
        if loai_raw:
            loai = NHAN_LOAI.get(_khoa(loai_raw), "")
            if not loai:
                self._loi(so_dong, "Loại khách",
                          f'"{loai_raw}" không hợp lệ — điền "Công ty" hoặc "Cá nhân".')
        elif kh is not None and "customer_kind" in o:
            # Khách mới để trống ⇒ Công ty (§4.1). Khách ĐÃ CÓ mà xoá trắng thì không đoán: lặng lẽ
            # đổi một khách Cá nhân thành Công ty là sai mà không ai thấy.
            self._loi(so_dong, "Loại khách", 'Chưa điền loại khách — "Công ty" hoặc "Cá nhân".')

        sale_id, bo_sale = self._sale(o, so_dong, kh)
        tai_chinh = self._tai_chinh(o, so_dong, kh)

        if len(self.kq.loi) > hong:
            return None

        thong = {
            "name": ten,
            "customer_kind": loai,
            "tax_code": _gia("tax_code"),
            "phone": _gia("phone"),
            "email": _gia("email"),
            "address": _gia("address"),
            "contact_name": _gia("contact_name"),
            "sale_user_id": sale_id,
        }
        buoc = _Buoc(so_dong, kh, thong, tai_chinh, bo_sale=bo_sale)
        if kh is not None:
            self._so_sanh(buoc)
        return buoc

    # ---------------------------------------------------------------- Sale

    def _sale(self, o: dict, so_dong: int, kh: Customer | None) -> tuple[int | None, bool]:
        """(sale_user_id, bo_sale). Khách mới: None ⇒ service lấy người nhập.

        Dòng sửa so với người ĐANG phụ trách TRƯỚC khi tra bảng: file xuất ghi họ tên, mà tra thẳng
        thì gãy ở hai ca có thật — hai người trùng họ tên (bảng tra bỏ khoá tên), và người đã ra
        khỏi khối Kinh doanh nhưng còn giữ khách cũ. Cả hai đều chặn một file không hề sửa gì.
        """
        if kh is not None and "sale_user_id" not in o:
            return kh.sale_user_id, False

        raw = _chuoi(o.get("sale_user_id"))
        if kh is not None:
            if raw is None:
                if kh.sale_user_id is None:
                    return None, False
                # Xoá trắng ô Sale ⇒ GỠ người phụ trách (chủ chốt 17/09/2026). Vẫn là đổi Sale.
                if self._duoc_dieu_chuyen(so_dong, None):
                    self.kq.canh_bao.append(CanhBaoDong(so_dong, (
                        f"Gỡ Sale phụ trách (đang là {_ten_hien(self.nguoi, kh.sale_user_id)}) — "
                        "khách sẽ không còn ai phụ trách.")))
                return None, True
            if kh.sale_user_id is not None and _khoa(raw) in self._khoa_nguoi(kh.sale_user_id):
                return kh.sale_user_id, False

        if raw is None:
            return None, False
        sale_id = self.tra_sale.get(_khoa(raw))
        if sale_id is None:
            self._loi(so_dong, "Sale phụ trách", f'Không tìm thấy người nào tên "{raw}".')
            return None, False
        if kh is not None and sale_id == kh.sale_user_id:
            return sale_id, False       # gõ tên đăng nhập của chính người đang phụ trách
        if sale_id not in self.du_tu_cach:
            # CÓ người này, nhưng không thuộc khối Kinh doanh. Nói đúng lý do — cùng luật với ô
            # "NV phụ trách" trên màn, để đi đường Excel không gán được khách cho Thủ kho.
            self._loi(
                so_dong, "Sale phụ trách",
                f'"{raw}" không thuộc khối Kinh doanh nên không nhận khách được. '
                "Bật cờ khối Kinh doanh cho phòng của họ ở Phòng ban, hoặc chọn người khác.")
            return None, False
        if kh is not None:
            self._duoc_dieu_chuyen(so_dong, sale_id, raw)
        return sale_id, False

    def _khoa_nguoi(self, uid: int) -> set[str]:
        n = self.nguoi.get(uid)
        if n is None:
            return set()
        return {k for k in (_khoa(n.ten), _khoa(n.username)) if k}

    def _duoc_dieu_chuyen(self, so_dong: int, den: int | None, raw: str = "") -> bool:
        """Đổi/gỡ Sale của khách ĐÃ CÓ — y hệt luật nút Điều chuyển (`routers/customers.py`
        `reassign_customers`): cần quyền `reassign`, phạm vi khác `own`, và ở phạm vi phòng thì
        người nhận phải cùng phòng với người nhập."""
        if not self.co_quyen_dieu_chuyen or self.scope == SCOPE_OWN:
            self._loi(so_dong, "Sale phụ trách",
                      "Bạn không có quyền điều chuyển người phụ trách khách hàng.")
            return False
        if self.scope == SCOPE_DEPARTMENT and den is not None:
            n = self.nguoi.get(den)
            if (self.actor.department_id is None or n is None
                    or n.department_id != self.actor.department_id):
                self._loi(so_dong, "Sale phụ trách",
                          f'"{raw}" không thuộc phòng của bạn — chỉ điều chuyển được trong phòng.')
                return False
        return True

    # ---------------------------------------------------------------- tài chính

    def _tai_chinh(self, o: dict, so_dong: int, kh: Customer | None) -> dict:
        """Khách mới: chỉ trả dict khi người khai THẬT SỰ điền gì đó. Dòng sửa: đủ 6 khoá (cột vắng
        lấy giá trị cũ) — `_so_sanh` quyết có đổi hay không.

        Không có quyền `set_credit_terms`: bỏ qua, và chỉ bật cờ "đã bỏ qua" khi có ô ĐÃ ĐIỀN (khách
        mới) hoặc ĐÃ SỬA (khách cũ). File xuất có cột tài chính cho mọi người xem, nên bật cờ chỉ vì
        cột có mặt thì Sale nào nhập lại file xuất cũng ăn một câu cảnh báo vô nghĩa."""
        co_mat = [k for k in KHOA_TAI_CHINH if k in o]
        if not co_mat:
            return {}

        gia: dict = {}
        hong = False
        for khoa in co_mat:
            nhan, doc = DOC_TAI_CHINH[khoa]
            v, loi = doc(o.get(khoa), nhan)
            if loi:
                hong = True
                if self.co_tai_chinh:
                    self._loi(so_dong, nhan, loi)
            gia[khoa] = v

        if not self.co_tai_chinh:
            if kh is None:
                co_dien = any(_chuoi(o.get(k)) is not None for k in co_mat)
            else:
                co_dien = hong or any(not _bang_tai_chinh(k, getattr(kh, k), gia[k]) for k in co_mat)
            if co_dien:
                self.kq.bo_qua_tai_chinh = True
            return {}

        if kh is None:
            # Gọi `update_financial` với toàn None là ghi đè "chưa đặt" lên chính sách mặc định, và
            # đẻ một dòng nhật ký cho việc không xảy ra.
            return {} if all(v is None for v in gia.values()) else gia
        return {k: (gia[k] if k in gia else getattr(kh, k)) for k in KHOA_TAI_CHINH}

    # ---------------------------------------------------------------- so cũ/mới

    def _so_sanh(self, b: _Buoc) -> None:
        kh = b.kh
        for khoa in KHOA_DINH_DANH:
            cu, moi = getattr(kh, khoa), b.thong[khoa]
            if khoa == "sale_user_id":
                khac = cu != moi
            else:
                khac = (cu or "").strip() != (moi or "").strip()
            if khac:
                b.doi.add(khoa)
                self._ghi_thay_doi(b, khoa, cu, moi)

        if b.tai_chinh:
            doi_tc = [k for k in KHOA_TAI_CHINH
                      if not _bang_tai_chinh(k, getattr(kh, k), b.tai_chinh[k])]
            if doi_tc:
                for k in doi_tc:
                    b.doi.add(k)
                    self._ghi_thay_doi(b, k, getattr(kh, k), b.tai_chinh[k])
            else:
                b.tai_chinh = {}

        # Rủi ro lớn nhất của đường sửa là TRỎ NHẦM khách (xuất máy này nhập máy kia, hoặc sắp xếp
        # một cột làm lệch dòng). Không đoán được ý, nhưng đổi cùng lúc cả tên lẫn MST đang có là
        # dấu hiệu đáng nói ra. Cảnh báo, không chặn: đổi tên + sửa MST gõ sai là việc có thật.
        if "name" in b.doi and "tax_code" in b.doi and kh.tax_code:
            self.kq.canh_bao.append(CanhBaoDong(
                b.dong, f"Đổi cả tên lẫn MST của {kh.code} — kiểm lại dòng này có trỏ nhầm khách không."))

    def _ghi_thay_doi(self, b: _Buoc, khoa: str, cu: Any, moi: Any) -> None:
        self.kq.thay_doi.append(ThayDoi(
            dong=b.dong, ma=b.kh.code, ten=b.kh.name, cot=NHAN_THEO_KHOA[khoa],
            cu=self._hien(khoa, cu), moi=self._hien(khoa, moi),
        ))

    def _hien(self, khoa: str, v: Any) -> str:
        if khoa == "credit_limit":
            return f"{int(v or 0):,}".replace(",", ".")     # trống = 0, đúng thứ service ghi
        if v is None or v == "":
            return ""
        if khoa == "customer_kind":
            return NHAN_LOAI_NGUOC.get(v, str(v))
        if khoa == "sale_user_id":
            return _ten_hien(self.nguoi, v)
        if isinstance(v, float):
            return f"{v:g}"
        return str(v)


def _bang_tai_chinh(khoa: str, cu: Any, moi: Any) -> bool:
    """So một ô tài chính cũ/mới. Hạn mức trống = 0 (service ép vậy). Phần trăm làm tròn 4 chữ số:
    Excel trả `8.000000001` thì dòng không sửa cũng bị tính là "sửa"."""
    if khoa == "credit_limit":
        return (cu or 0) == (moi or 0)
    if cu is None or moi is None:
        return cu is None and moi is None
    if khoa == "payment_term_days":
        return int(cu) == int(moi)
    return round(float(cu), 4) == round(float(moi), 4)
