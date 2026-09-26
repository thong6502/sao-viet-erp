"""File CHUYỂN KHOẢN tạm ứng / lương đợt 1 theo khuôn lô lương BIZ MBBank (25/09/2026).

Chủ gửi mẫu `CK LUONG ỨNG T8.2026.xlsx` (sheet `eMB_BulkPayment` — file tải lên BIZ MBBank để
chuyển lương một lô): *"tiền mặt thì ghi tiền mặt, còn chuyển khoản thì ghi số tài khoản, tên
ngân hàng của nhân viên"*. Giữ ĐÚNG 6 cột, tiêu đề hai thứ tiếng và vị trí dòng của mẫu (dòng 1 tiêu
đề, dòng 2 đầu cột, dữ liệu từ dòng 3) để kế toán tải lên ngân hàng không phải sửa tay.

Luật của ngân hàng (sheet "Hướng dẫn nhập liệu" trong mẫu) mà file tự làm sẵn:
  * Số tài khoản: không dấu cách ở đầu / giữa / cuối — hồ sơ gõ tay hay dính dấu cách.
  * Chi tiết thanh toán = LÝ DO của phiếu (chủ 25/09/2026: "chi tiết thanh toán là cái lý do"),
    tối đa 140 ký tự; phiếu bỏ trống lý do mới dùng câu mặc định kiểu mẫu.
  * Tên thụ hưởng + chi tiết thanh toán: chữ có dấu tiếng Việt bị NGÂN HÀNG XOÁ khi tải lên
    ("ký tự khác [A-Za-z0-9] → rỗng") ⇒ in HOA, bỏ dấu sẵn ở đây ("Đoàn Thị" → "DOAN THI").
  * Số tiền: số nguyên VND.

Dòng TIỀN MẶT ghi chữ "TIỀN MẶT" vào chỗ số tài khoản / ngân hàng và xếp XUỐNG CUỐI — khối chuyển
khoản liền một mạch ở trên; tải lên ngân hàng thì xoá khối tiền mặt đi.
"""
from __future__ import annotations

import re
import unicodedata
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

#: Tên công ty in ở cột "Chi tiết thanh toán" — đúng chữ của mẫu kế toán gửi.
TEN_CONG_TY = "CTY SAO VIET NHAT"
TIEN_MAT = "TIỀN MẶT"

_DAU_COT = [
    "STT\n(Ord. No.)\n(1)",
    "Số tài khoản\n(Account No.)\n(2)",
    "Tên đơn vị thụ hưởng\n(Beneficiary Organization)\n(3)",
    "Ngân hàng thụ hưởng/Chi nhánh\n(Beneficiary Bank)\n(4)",
    "Số tiền\n(Amount)\n(5)",
    "Chi tiết thanh toán\n(Payment Detail)\n(6)",
]
_RONG_COT = {"A": 11.7, "B": 27.4, "C": 32.7, "D": 36.9, "E": 22.9, "F": 44.3}
_DINH_DANG_TIEN = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'


def khong_dau_hoa(s: str | None) -> str:
    """"Đoàn Thị Chúc Phấn" → "DOAN THI CHUC PHAN" (luật ký tự của ngân hàng, xem đầu file)."""
    s = (s or "").replace("Đ", "D").replace("đ", "d")
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().upper()


def chi_tiet_thanh_toan(kind: str, nam: int, thang: int, ly_do: str | None = None) -> str:
    """Lý do của phiếu (IN HOA không dấu, ≤ 140 ký tự — luật ngân hàng). Trống thì câu mặc định
    theo chữ của mẫu: "CTY SAO VIET NHAT CHI LUONG T8.2026 DOT 1"."""
    ly_do = khong_dau_hoa(ly_do)
    if ly_do:
        return ly_do[:140]
    if kind == "luong_dot_1":
        return f"{TEN_CONG_TY} CHI LUONG T{thang}.{nam} DOT 1"
    return f"{TEN_CONG_TY} TAM UNG LUONG T{thang}.{nam}"


def xuat_file_chuyen_khoan(dong: list[dict], *, nam: int, thang: int) -> bytes:
    """`dong` đã xếp sẵn (chuyển khoản trước, tiền mặt sau) — mỗi phần tử:
    `{ten, so_tai_khoan, ngan_hang, so_tien, kind, ly_do, chuyen_khoan}`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "eMB_BulkPayment"

    ws.merge_cells("B1:F1")
    ws["B1"] = "DANH SÁCH GIAO DỊCH\n(LIST OF TRANSACTIONS)"
    ws["B1"].font = Font(name="Cambria", size=16, bold=True)
    ws["B1"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.row_dimensions[1].height = 43.5

    vien = Side(style="thin")
    for i, chu in enumerate(_DAU_COT, start=1):
        o = ws.cell(row=2, column=i, value=chu)
        o.font = Font(name="Times New Roman", size=12, bold=True)
        o.fill = PatternFill("solid", fgColor="FFD8DAD9")
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.border = Border(left=vien, right=vien, top=vien, bottom=vien)
    ws.row_dimensions[2].height = 47.25
    for cot, rong in _RONG_COT.items():
        ws.column_dimensions[cot].width = rong

    for stt, d in enumerate(dong, start=1):
        r = stt + 2
        if d["chuyen_khoan"]:
            so_tk = re.sub(r"\s+", "", d.get("so_tai_khoan") or "")
            ngan_hang = (d.get("ngan_hang") or "").strip()
        else:
            so_tk = ngan_hang = TIEN_MAT
        ws.cell(row=r, column=1, value=stt)
        # Số tài khoản là CHỮ: để Excel hiểu là số thì mất số 0 đầu ("0908872122" → 908872122).
        o = ws.cell(row=r, column=2, value=so_tk)
        o.number_format = "@"
        ws.cell(row=r, column=3, value=khong_dau_hoa(d.get("ten")))
        ws.cell(row=r, column=4, value=ngan_hang)
        o = ws.cell(row=r, column=5, value=int(round(float(d.get("so_tien") or 0))))
        o.number_format = _DINH_DANG_TIEN
        o = ws.cell(row=r, column=6, value=chi_tiet_thanh_toan(d.get("kind") or "", nam, thang,
                                                                 d.get("ly_do")))
        o.number_format = "@"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
